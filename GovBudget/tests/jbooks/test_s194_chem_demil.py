"""scripts/era/s194_chem_demil.py (families piece-1 follow-ups design §4.4,
#194; owner Q1 (a), Q2, Q8, Q24 (b)).

The database starts in production's pre-write shape: PB2024 and PB2025 hold
0390D's two Chem Demilitarization lines under the printed codes 'O&M' and
'RDT&E' (the stale keys the 2026-07-05 guard never cleaned), PB2026 holds
neither (that cleanup deleted them), and each edition's R-1 carries 0708081D
and 0708083D, which the RDT&E line restates. The values are the live ones
(budget_lines ids 349016-349025, 47733-47738, 1031-1036; PB2026 P-1 rows
303-304), two amount columns per edition.
"""
import dataclasses
import hashlib
import importlib.util
import json
import os
import stat
from pathlib import Path
from urllib.parse import urlsplit

import psycopg
import pytest
from openpyxl import Workbook

from govbudget.jbooks.p1_loader import parse_p1_rollup, write_p1_rows

ROOT = Path(__file__).resolve().parents[2]
_spec = importlib.util.spec_from_file_location(
    "s194_chem_demil", ROOT / "scripts" / "era" / "s194_chem_demil.py")
s194 = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(s194)

ID_HEADERS = ["Account", "Account Title", "Organization", "Budget Activity",
              "Budget Activity Title", "Line Number", "BSA", "Budget SubActivity (BSA) Title",
              "Budget Line Item", "Budget Line Item (BLI) Title", "Cost Type",
              "Cost Type Title", "Add/Non-Add"]
# edition -> the two FY amount headers its P-1 prints (PB2025's CR column
# keeps its footnote '*', as the real workbook does)
FY_HEADERS = {
    2024: ("FY 2022 Actuals", "FY 2024 Request"),
    2025: ("FY 2023 Actuals", "FY 2024 PB Request with CR Adjustments"),
    2026: ("FY 2024 Actuals", "FY 2026 Total"),
}
# edition -> (ordinary BLI, O&M, RDT&E) amounts in the two columns
P1_VALUES = {
    2024: ((100, 200), (66587, 89284), (1026665, 1002560)),
    2025: ((110, 210), (84612, 89284), (975206, 1002560)),
    2026: ((120, 220), (87404, 3243), (1002560, 210039)),
}
# edition -> [(pe, amount_type, amount)] of the R-1 0390D rows
R1_ROWS = {
    2024: [("0708081D", "fy_2022_actuals", 6120), ("0708081D", "fy_2024_request", 2093),
           ("0708083D", "fy_2022_actuals", 1020545), ("0708083D", "fy_2024_request", 1000467)],
    2025: [("0708081D", "fy_2023_actuals", 3464),
           ("0708081D", "fy_2024_pb_request_with_cr_amounts", 2093),
           ("0708083D", "fy_2023_actuals", 971742),
           ("0708083D", "fy_2024_pb_request_with_cr_amounts", 1000467)],
    2026: [("0708081D", "fy_2024_actuals", 2093), ("0708083D", "fy_2024_actuals", 1000467),
           ("0708083D", "fy_2026_total", 210039)],
}
CHEM = ("0390D", "Chemical Agents and Munitions Destruction, Defense", "A")


def _workbook(path: Path, fy: int) -> Path:
    a, b = FY_HEADERS[fy]
    star = "*" if fy == 2025 else ""
    ordinary, om, rdte = P1_VALUES[fy]
    wb = Workbook()
    ws = wb.active
    ws.title = "Exhibit P-1"
    ws.append(["Total of Displayed Rows"])
    ws.append(ID_HEADERS + [f"{a} Quantity", f"{a} Amount", f"{b} Quantity", f"{b} Amount{star}"])
    ws.append(["0300D", "Procurement, Defense-Wide", "CBDP", "03", "Chem/Bio Defense", "120",
               "1", "CBDP", "7001SA1000", "CB Situational Awareness", "A", "Weapon System Cost",
               "Add", "", ordinary[0], "", ordinary[1]])
    ws.append([*CHEM, "01", "Operation and maintenance", "1", "", "", "O&M",
               "Chem Demilitarization - O&M", "A", "Weapon System Cost", "Add", "", om[0], "", om[1]])
    ws.append([*CHEM, "02", "Research, Development, Test, and Evaluation", "2", "", "", "RDT&E",
               "Chem Demilitarization - RDT&E", "A", "Weapon System Cost", "Add",
               "", rdte[0], "", rdte[1]])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _register(con, fy, title, path=None):
    sha = hashlib.sha256(path.read_bytes()).hexdigest() if path else None
    return con.execute(
        "insert into jbook_documents (org, exhibit_family, fiscal_year, title, source_url,"
        " file_path, sha256, status) values ('DoD', 'rollup', %s, %s, %s, %s, %s,"
        " 'downloaded') returning id",
        (fy, title, f"https://example.test/fy{fy}/{title}", str(path or f"/x/{fy}/{title}"),
         sha)).fetchone()[0]


@pytest.fixture()
def prod(pg_dsn, tmp_path):
    """Production's pre-write state; returns {edition: (doc_id, workbook)}."""
    docs = {}
    with psycopg.connect(pg_dsn) as con:
        for fy in (2024, 2025, 2026):
            p = _workbook(tmp_path / f"fy{fy}" / "p1_display.xlsx", fy)
            docs[fy] = (_register(con, fy, "p1_display.xlsx", p), p)
            r1_doc = _register(con, fy, "r1_display.xlsx")
            for pe, at, amt in R1_ROWS[fy]:
                con.execute(
                    "insert into budget_lines (exhibit, fiscal_year, account, account_title,"
                    " organization, budget_activity, pe_bli, title, amount_type,"
                    " amount_thousands, source_document_id, source_sheet, source_cells)"
                    " values ('R-1', %s, '0390D', %s, 'A', '02', %s, 'x', %s, %s, %s,"
                    " 'Exhibit R-1', '{J9}')", (fy, CHEM[1], pe, at, amt, r1_doc))
        for fy, (doc_id, p) in docs.items():
            parsed = parse_p1_rollup(p, exhibit="P-1", fiscal_year=fy)
            rows = [r for r in parsed.rows if r.pe_bli != "0390DOM"]
            if fy != 2026:  # the stale keys; PB2026's were deleted on 2026-07-05
                rows += [dataclasses.replace(r, pe_bli="O&M")
                         for r in parsed.rows if r.pe_bli == "0390DOM"]
                rows += list(parsed.restated)
            write_p1_rows(con, rows, source_document_id=doc_id)
    return docs


def _snapshot(pg_dsn):
    with psycopg.connect(pg_dsn) as con:
        return con.execute(
            f"select {', '.join(s194.ROW_COLUMNS)} from budget_lines order by id").fetchall()


def _run(pg_dsn, *args):
    return s194.main([*map(str, args), "--dsn", pg_dsn])


def test_check_sorts_every_difference_into_194s_three_kinds(pg_dsn, prod, capsys):
    before = _snapshot(pg_dsn)
    assert _run(pg_dsn, "--check") == 0
    out = capsys.readouterr().out
    assert ("check PB2024 P-1 (doc %d): parse 4 | db 6 | re-key 2 | insert 0 | delete 2 |"
            " restated rows checked against the R-1 2 | other differences 0" % prod[2024][0]) in out
    assert "check PB2025 P-1 (doc %d): parse 4 | db 6 | re-key 2 | insert 0 | delete 2" % (
        prod[2025][0]) in out
    assert "check PB2026 P-1 (doc %d): parse 4 | db 2 | re-key 0 | insert 2 | delete 0" % (
        prod[2026][0]) in out
    assert "check: CLEAN, the only differences are #194's (re-key 4, insert 2, delete 4)" in out
    assert _snapshot(pg_dsn) == before


@pytest.mark.parametrize("tamper, expect", [
    ("update budget_lines set amount_thousands = 101 where pe_bli = '7001SA1000'"
     " and fiscal_year = 2024 and amount_type = 'fy_2022_actuals'", "diff ("),
    ("update budget_lines set title = 'Chem Demil O&M' where pe_bli = 'O&M'"
     " and fiscal_year = 2025 and amount_type = 'fy_2023_actuals'", "stale ("),
    ("insert into budget_lines (exhibit, fiscal_year, account, organization, budget_activity,"
     " pe_bli, title, amount_type, amount_thousands, source_document_id)"
     " select 'P-1', 2024, '0300D', 'CBDP', '03', 'A/B', 'x', 'fy_2022_actuals', 1, id"
     " from jbook_documents where fiscal_year = 2024 and title = 'p1_display.xlsx'",
     "extra in db: ('0300D', 'CBDP', '03', 'A/B', 'fy_2022_actuals')"),
])
def test_check_fails_on_any_other_difference(pg_dsn, prod, capsys, tamper, expect):
    with psycopg.connect(pg_dsn) as con:
        con.execute(tamper)
    assert _run(pg_dsn, "--check") == 1
    out = capsys.readouterr().out
    assert expect in out and "check: DIFFERENCES" in out


def test_apply_writes_exactly_194_and_records_the_pre_write_rows(pg_dsn, prod, tmp_path, capsys):
    before = {r[0]: r for r in _snapshot(pg_dsn)}
    stale = {i: r for i, r in before.items() if r[9] in ("O&M", "RDT&E")}
    revert_file = tmp_path / "revert.json"
    assert _run(pg_dsn, "--apply", "--revert-file", revert_file) == 0
    out = capsys.readouterr().out
    assert ("apply: re-keyed 4, deleted 4, inserted 2; untouched rows %d (digest equal);"
            " route-unsafe pe_bli left 0" % (len(before) - 8)) in out
    assert f"apply: COMMITTED (revert file {revert_file})" in out
    after = {r[0]: r for r in _snapshot(pg_dsn)}
    for i, r in stale.items():
        if r[9] == "O&M":       # re-keyed in place: same id, only pe_bli moves
            assert after[i] == r[:9] + ("0390DOM",) + r[10:]
        else:                   # the restatement is gone
            assert i not in after
    new = [r for i, r in after.items() if i not in before]
    assert sorted((r[2], r[9], r[11], r[12], r[16]) for r in new) == [
        (2026, "0390DOM", "fy_2024_actuals", 87404, "O&M"),
        (2026, "0390DOM", "fy_2026_total", 3243, "O&M")]
    assert {i: r for i, r in after.items() if i in before and i not in stale} == {
        i: r for i, r in before.items() if i not in stale}
    rec = json.loads(revert_file.read_text())
    assert rec["database"] == pg_dsn.rsplit("/", 1)[1]
    assert sorted(r[0] for r in rec["before"]) == sorted(stale)
    assert (len(rec["rekeyed"]), len(rec["deleted_ids"]), len(rec["inserted"])) == (4, 4, 2)
    assert rec["rows_sha256"] == s194.record_sha256(rec)
    assert _run(pg_dsn, "--check") == 0
    assert "check: CLEAN, nothing left to write" in capsys.readouterr().out


def test_a_second_apply_writes_nothing(pg_dsn, prod, tmp_path, capsys):
    assert _run(pg_dsn, "--apply", "--revert-file", tmp_path / "one.json") == 0
    snap = _snapshot(pg_dsn)
    assert _run(pg_dsn, "--apply", "--revert-file", tmp_path / "two.json") == 0
    assert "apply: nothing to write (#194 is already applied)" in capsys.readouterr().out
    assert _snapshot(pg_dsn) == snap and not (tmp_path / "two.json").exists()


def test_dry_run_changes_nothing(pg_dsn, prod, capsys):
    before = _snapshot(pg_dsn)
    assert _run(pg_dsn, "--dry-run") == 0
    assert "apply: dry run, every check passed: ROLLED BACK" in capsys.readouterr().out
    assert _snapshot(pg_dsn) == before


def test_apply_needs_a_new_revert_file(pg_dsn, prod, tmp_path, capsys):
    before = _snapshot(pg_dsn)
    assert _run(pg_dsn, "--apply") == 2
    taken = tmp_path / "taken.json"
    taken.write_text("{}")
    assert _run(pg_dsn, "--apply", "--revert-file", taken) == 2
    out = capsys.readouterr().out
    assert "--revert-file is required" in out and "refusing to overwrite a revert file" in out
    assert _snapshot(pg_dsn) == before


def test_apply_refuses_when_the_check_is_not_clean(pg_dsn, prod, tmp_path, capsys):
    with psycopg.connect(pg_dsn) as con:
        con.execute("update budget_lines set amount_thousands = 101 where pe_bli = '7001SA1000'"
                    " and fiscal_year = 2026 and amount_type = 'fy_2024_actuals'")
    before = _snapshot(pg_dsn)
    assert _run(pg_dsn, "--apply", "--revert-file", tmp_path / "r.json") == 1
    assert "apply: check not clean, nothing written (ROLLED BACK)" in capsys.readouterr().out
    assert _snapshot(pg_dsn) == before and not (tmp_path / "r.json").exists()


def test_apply_rolls_back_when_a_write_touches_another_row(
        pg_dsn, prod, tmp_path, capsys, monkeypatch):
    real = s194.write_p1_rows

    def stray(con, rows, *, source_document_id):
        n = real(con, rows, source_document_id=source_document_id)
        con.execute("update budget_lines set title = 'moved' where pe_bli = '7001SA1000'")
        return n

    monkeypatch.setattr(s194, "write_p1_rows", stray)
    before = _snapshot(pg_dsn)
    assert _run(pg_dsn, "--apply", "--revert-file", tmp_path / "r.json") == 1
    out = capsys.readouterr().out
    assert "PROBLEM rows the writes did not target changed" in out
    assert "problem(s): ROLLED BACK" in out
    assert _snapshot(pg_dsn) == before and not (tmp_path / "r.json").exists()


def test_a_drifted_restatement_stops_before_any_write(pg_dsn, prod, tmp_path, capsys):
    with psycopg.connect(pg_dsn) as con:
        con.execute("update budget_lines set amount_thousands = 1000466 where pe_bli = '0708083D'"
                    " and amount_type = 'fy_2024_pb_request_with_cr_amounts'")
    before = _snapshot(pg_dsn)
    for args in (["--check"], ["--apply", "--revert-file", tmp_path / "r.json"]):
        assert _run(pg_dsn, *args) == 1
        out = capsys.readouterr().out
        assert "s194: REFUSED — P1RestatementMismatch: PB2025 P-1 0390D 'RDT&E'" in out
        # the R-1 side is the sum of both program elements: 2,093 + 1,000,466
        assert "fy_2024_pb_request_with_cr_amounts = 1002559" in out
    assert _snapshot(pg_dsn) == before and not (tmp_path / "r.json").exists()


def test_a_workbook_the_loader_refuses_stops_before_any_write(pg_dsn, prod, tmp_path, capsys):
    """An unknown route-unsafe code (Task 6's UnsafeBudgetLineCode) refuses the
    whole run; so would a P1KeyConflict (Task 3)."""
    doc_id, path = prod[2025]
    _workbook(path, 2025)
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb["Exhibit P-1"]
    ws.append(["0300D", "Procurement, Defense-Wide", "CBDP", "03", "Chem/Bio Defense", "121",
               "1", "CBDP", "A/B", "Stray", "A", "Weapon System Cost", "Add", "", 1, "", 2])
    wb.save(path)
    with psycopg.connect(pg_dsn) as con:
        con.execute("update jbook_documents set sha256 = %s where id = %s",
                    (hashlib.sha256(path.read_bytes()).hexdigest(), doc_id))
    before = _snapshot(pg_dsn)
    assert _run(pg_dsn, "--apply", "--revert-file", tmp_path / "r.json") == 1
    assert "s194: REFUSED — UnsafeBudgetLineCode: p1_display.xlsx (P-1 PB2025) row 6: 0300D prints 'A/B'" in (
        capsys.readouterr().out)
    assert _snapshot(pg_dsn) == before and not (tmp_path / "r.json").exists()


def test_revert_restores_the_exact_pre_write_rows(pg_dsn, prod, tmp_path, capsys):
    before = _snapshot(pg_dsn)
    revert_file = tmp_path / "revert.json"
    assert _run(pg_dsn, "--apply", "--revert-file", revert_file) == 0
    assert _run(pg_dsn, "--revert", "--revert-file", revert_file) == 0
    out = capsys.readouterr().out
    assert ("revert: restored 8 pre-write row(s) (re-keyed 4 back, re-inserted 4,"
            " removed 2); COMMITTED") in out
    assert _snapshot(pg_dsn) == before
    assert _run(pg_dsn, "--revert", "--revert-file", revert_file) == 0
    assert ("revert: the database already holds the pre-write rows; nothing to revert"
            in capsys.readouterr().out)
    assert _snapshot(pg_dsn) == before


def test_revert_refuses_a_database_that_moved(pg_dsn, prod, tmp_path, capsys):
    revert_file = tmp_path / "revert.json"
    assert _run(pg_dsn, "--apply", "--revert-file", revert_file) == 0
    with psycopg.connect(pg_dsn) as con:
        con.execute("update budget_lines set title = 'moved' where pe_bli = '7001SA1000'"
                    " and fiscal_year = 2024 and amount_type = 'fy_2022_actuals'")
    moved = _snapshot(pg_dsn)
    assert _run(pg_dsn, "--revert", "--revert-file", revert_file) == 1
    out = capsys.readouterr().out
    assert "PROBLEM rows --apply did not touch changed since" in out
    assert "revert: the database moved since --apply (1 problem(s)); nothing written" in out
    assert _snapshot(pg_dsn) == moved


def test_revert_refuses_another_database(pg_dsn, prod, tmp_path, capsys):
    revert_file = tmp_path / "revert.json"
    assert _run(pg_dsn, "--apply", "--revert-file", revert_file) == 0
    rec = json.loads(revert_file.read_text())
    rec["database"] = "govbudget"
    rec["rows_sha256"] = s194.record_sha256(rec)   # a genuine file of another database
    revert_file.write_text(json.dumps(rec))
    applied = _snapshot(pg_dsn)
    assert _run(pg_dsn, "--revert", "--revert-file", revert_file) == 1
    assert "was written on database 'govbudget'" in capsys.readouterr().out
    assert _snapshot(pg_dsn) == applied


def test_revert_refuses_an_edited_revert_file(pg_dsn, prod, tmp_path, capsys):
    revert_file = tmp_path / "revert.json"
    assert _run(pg_dsn, "--apply", "--revert-file", revert_file) == 0
    rec = json.loads(revert_file.read_text())
    rec["before"][0][s194.ROW_COLUMNS.index("amount_thousands")] = "999999999"
    revert_file.write_text(json.dumps(rec))
    applied = _snapshot(pg_dsn)
    assert _run(pg_dsn, "--revert", "--revert-file", revert_file) == 1
    assert (f"revert: {revert_file} sha256 does not match its contents; nothing written"
            in capsys.readouterr().out)
    assert _snapshot(pg_dsn) == applied


def test_revert_refuses_a_changed_in_scope_row(pg_dsn, prod, tmp_path, capsys):
    revert_file = tmp_path / "revert.json"
    assert _run(pg_dsn, "--apply", "--revert-file", revert_file) == 0
    with psycopg.connect(pg_dsn) as con:
        rid = con.execute("select id from budget_lines where pe_bli = '0390DOM'"
                          " and fiscal_year = 2024 and amount_type = 'fy_2022_actuals'").fetchone()[0]
        con.execute("update budget_lines set amount_thousands = 1 where id = %s", (rid,))
    moved = _snapshot(pg_dsn)
    assert _run(pg_dsn, "--revert", "--revert-file", revert_file) == 1
    out = capsys.readouterr().out
    assert f"PROBLEM re-keyed row {rid} is not as --apply wrote it" in out
    assert "revert: the database moved since --apply (1 problem(s)); nothing written" in out
    assert _snapshot(pg_dsn) == moved


def test_revert_refuses_a_file_naming_rows_outside_194s_scope(pg_dsn, prod, tmp_path, capsys):
    revert_file = tmp_path / "revert.json"
    assert _run(pg_dsn, "--apply", "--revert-file", revert_file) == 0
    rec = json.loads(revert_file.read_text())
    gone = next(r for r in rec["before"] if r[0] in rec["deleted_ids"])
    gone[s194.ROW_COLUMNS.index("pe_bli")] = "7001SA1000"
    rec["rows_sha256"] = s194.record_sha256(rec)   # consistent, but not #194's rows
    revert_file.write_text(json.dumps(rec))
    applied = _snapshot(pg_dsn)
    assert _run(pg_dsn, "--revert", "--revert-file", revert_file) == 1
    out = capsys.readouterr().out
    assert (f"PROBLEM row {gone[0]} ('0390D', '7001SA1000') is neither re-keyed to its"
            " page key nor a restatement") in out
    assert f"revert: {revert_file} names 1 row(s) outside #194's scope; nothing written" in out
    assert _snapshot(pg_dsn) == applied


# Task 7 fix round 1 (review, 2026-10-14): the refusals that guard Task 22b's
# live write. None of these may reach a database it was not pointed at.


class _NoConnect(Exception):
    pass


def _forbid_connect(monkeypatch) -> list:
    """Make every psycopg.connect raise; returns the DSNs it was asked for."""
    asked = []

    def refuse(dsn, *args, **kwargs):
        asked.append(dsn)
        raise _NoConnect(dsn)

    monkeypatch.setattr(s194.psycopg, "connect", refuse)
    return asked


@pytest.mark.parametrize("argv", [
    ["--check", "--dsn", ""],
    ["--dry-run", "--dsn", " "],
    ["--apply", "--dsn", "  ", "--revert-file", "{tmp}/r.json"],
    ["--revert", "--dsn", "\t", "--revert-file", "{tmp}/r.json"],
])
def test_an_empty_dsn_is_refused_before_any_connection(monkeypatch, tmp_path, capsys, argv):
    """An empty --dsn (say, an unset "$POST_DSN") must never fall back to
    GOVBUDGET_PG_DSN, which is the live database once env.sh is sourced."""
    asked = _forbid_connect(monkeypatch)
    monkeypatch.setattr(s194.config, "PG_DSN", "postgresql://sentinel.invalid/never")
    assert s194.main([a.replace("{tmp}", str(tmp_path)) for a in argv]) == 2
    captured = capsys.readouterr()
    assert "usage:" in captured.err and "s194: --dsn is empty" in captured.err
    assert captured.out == "" and asked == [] and not (tmp_path / "r.json").exists()


def test_an_empty_govbudget_pg_dsn_is_refused_too(monkeypatch, capsys):
    asked = _forbid_connect(monkeypatch)
    monkeypatch.setattr(s194.config, "PG_DSN", "  ")
    assert s194.main(["--check"]) == 2
    assert "s194: GOVBUDGET_PG_DSN is empty" in capsys.readouterr().err
    assert asked == []


def test_the_dsn_falls_back_to_govbudget_pg_dsn_only_when_omitted(monkeypatch):
    asked = _forbid_connect(monkeypatch)
    monkeypatch.setattr(s194.config, "PG_DSN", "postgresql://sentinel.invalid/never")
    for argv in (["--check"], ["--check", "--dsn", "postgresql://other.invalid/x"]):
        with pytest.raises(_NoConnect):
            s194.main(argv)
    assert asked == ["postgresql://sentinel.invalid/never", "postgresql://other.invalid/x"]


def test_every_mode_first_names_the_database_and_host(pg_dsn, prod, tmp_path, capsys):
    url = urlsplit(pg_dsn)
    where = f"database {url.path[1:]!r} on host {url.hostname}:{url.port}"
    revert_file = tmp_path / "revert.json"
    for args in (["--check"], ["--dry-run"], ["--apply", "--revert-file", revert_file],
                 ["--revert", "--revert-file", revert_file]):
        assert _run(pg_dsn, *args) == 0
        assert capsys.readouterr().out.splitlines()[0] == f"s194 {args[0]}: {where}"


def test_a_p1_document_count_other_than_one_is_a_precondition_error(
        pg_dsn, prod, tmp_path, capsys):
    with psycopg.connect(pg_dsn) as con:
        _register(con, 2025, "p1_display_copy.xlsx")
    before = _snapshot(pg_dsn)
    for args in (["--check"], ["--dry-run"], ["--apply", "--revert-file", tmp_path / "r.json"]):
        assert _run(pg_dsn, *args) == 2
        assert ("s194: expected one P-1 rollup document per edition, got {2025: 2}"
                in capsys.readouterr().out)
    assert _snapshot(pg_dsn) == before and not (tmp_path / "r.json").exists()


def test_check_refuses_a_workbook_that_is_not_the_registered_file(pg_dsn, prod, tmp_path, capsys):
    doc_id, path = prod[2024]
    actual = hashlib.sha256(path.read_bytes()).hexdigest()
    with psycopg.connect(pg_dsn) as con:
        con.execute("update jbook_documents set sha256 = %s where id = %s", ("0" * 64, doc_id))
    before = _snapshot(pg_dsn)
    assert _run(pg_dsn, "--check") == 1
    out = capsys.readouterr().out
    assert f"check PB2024: workbook sha256 {actual} != jbook_documents {'0' * 64}" in out
    assert "check: DIFFERENCES" in out
    assert _run(pg_dsn, "--apply", "--revert-file", tmp_path / "r.json") == 1
    assert "apply: check not clean, nothing written (ROLLED BACK)" in capsys.readouterr().out
    assert _snapshot(pg_dsn) == before and not (tmp_path / "r.json").exists()


def test_revert_refuses_a_changed_inserted_row(pg_dsn, prod, tmp_path, capsys):
    revert_file = tmp_path / "revert.json"
    assert _run(pg_dsn, "--apply", "--revert-file", revert_file) == 0
    rid = json.loads(revert_file.read_text())["inserted"][0][0]
    with psycopg.connect(pg_dsn) as con:
        con.execute("update budget_lines set amount_thousands = 1 where id = %s", (rid,))
    moved = _snapshot(pg_dsn)
    capsys.readouterr()
    assert _run(pg_dsn, "--revert", "--revert-file", revert_file) == 1
    out = capsys.readouterr().out
    assert f"PROBLEM inserted row {rid} is not as --apply wrote it" in out
    assert "revert: the database moved since --apply (1 problem(s)); nothing written" in out
    assert _snapshot(pg_dsn) == moved


def test_revert_refuses_when_a_deleted_row_is_back(pg_dsn, prod, tmp_path, capsys):
    revert_file = tmp_path / "revert.json"
    assert _run(pg_dsn, "--apply", "--revert-file", revert_file) == 0
    rec = json.loads(revert_file.read_text())
    row = next(r for r in rec["before"] if r[0] in rec["deleted_ids"])
    with psycopg.connect(pg_dsn) as con:
        con.execute(f"insert into budget_lines ({', '.join(s194.ROW_COLUMNS)}) values"
                    f" ({', '.join(['%s'] * len(s194.ROW_COLUMNS))})", row)
    moved = _snapshot(pg_dsn)
    capsys.readouterr()
    assert _run(pg_dsn, "--revert", "--revert-file", revert_file) == 1
    out = capsys.readouterr().out
    assert f"PROBLEM deleted row {row[0]} is back" in out
    # the row is also one --apply did not leave in place, so that digest moved too
    assert "PROBLEM rows --apply did not touch changed since" in out
    assert "revert: the database moved since --apply (2 problem(s)); nothing written" in out
    assert _snapshot(pg_dsn) == moved


@pytest.mark.parametrize("column, value, why", [
    ("exhibit", "R-1", "exhibit 'R-1', not 'P-1'"),
    ("pe_bli", "7001SA1000", "('0390D', '7001SA1000') is not a route-safe page key"),
    ("fiscal_year", 2019, "fiscal year 2019 is not among the editions [2024, 2025, 2026]"),
    ("id", "a before id", "its id is also a `before` row"),
])
def test_revert_refuses_inserted_rows_outside_194s_scope(
        pg_dsn, prod, tmp_path, capsys, column, value, why):
    revert_file = tmp_path / "revert.json"
    assert _run(pg_dsn, "--apply", "--revert-file", revert_file) == 0
    rec = json.loads(revert_file.read_text())
    row = rec["inserted"][0]
    if value == "a before id":
        value = rec["before"][0][0]
    row[s194.ROW_COLUMNS.index(column)] = value
    rec["rows_sha256"] = s194.record_sha256(rec)   # consistent, but not #194's rows
    revert_file.write_text(json.dumps(rec))
    applied = _snapshot(pg_dsn)
    capsys.readouterr()
    assert _run(pg_dsn, "--revert", "--revert-file", revert_file) == 1
    out = capsys.readouterr().out
    assert f"PROBLEM inserted row {row[0]}: {why}" in out
    assert f"revert: {revert_file} names 1 row(s) outside #194's scope; nothing written" in out
    assert _snapshot(pg_dsn) == applied


@pytest.fixture()
def widen_budget_lines(pg_dsn):
    """Adds a column to budget_lines on demand; drops it again afterwards."""
    def widen():
        with psycopg.connect(pg_dsn) as con:
            con.execute("alter table budget_lines add column s194_extra text")
    yield widen
    with psycopg.connect(pg_dsn) as con:
        con.execute("alter table budget_lines drop column if exists s194_extra")


def test_apply_and_revert_refuse_a_budget_lines_of_another_shape(
        pg_dsn, prod, tmp_path, capsys, widen_budget_lines):
    revert_file = tmp_path / "revert.json"
    assert _run(pg_dsn, "--apply", "--revert-file", revert_file) == 0
    widen_budget_lines()
    applied = _snapshot(pg_dsn)
    capsys.readouterr()
    for args in (["--revert", "--revert-file", revert_file], ["--dry-run"],
                 ["--apply", "--revert-file", tmp_path / "two.json"]):
        assert _run(pg_dsn, *args) == 2
        out = capsys.readouterr().out
        assert "s194: budget_lines columns" in out and "s194_extra" in out
    assert _snapshot(pg_dsn) == applied and not (tmp_path / "two.json").exists()


def _truncated(text, rec):
    return text[: len(text) // 2]


def _without(field):
    def damage(text, rec):
        del rec[field]
        rec["rows_sha256"] = s194.record_sha256(rec)
        return json.dumps(rec)
    return damage


def _short_before_row(text, rec):
    rec["before"][0] = rec["before"][0][:5]
    rec["rows_sha256"] = s194.record_sha256(rec)
    return json.dumps(rec)


@pytest.mark.parametrize("damage", [
    _truncated, lambda text, rec: "{}", lambda text, rec: "[]", _without("inserted"),
    _without("untouched_digest"), _short_before_row,
], ids=["truncated", "empty-object", "a-list", "no-inserted", "no-digest", "short-row"])
def test_revert_refuses_a_malformed_revert_file_before_connecting(
        pg_dsn, prod, tmp_path, capsys, monkeypatch, damage):
    revert_file = tmp_path / "revert.json"
    assert _run(pg_dsn, "--apply", "--revert-file", revert_file) == 0
    text = revert_file.read_text()
    revert_file.write_text(damage(text, json.loads(text)))
    capsys.readouterr()
    asked = _forbid_connect(monkeypatch)
    assert _run(pg_dsn, "--revert", "--revert-file", revert_file) == 2
    assert (f"revert: {revert_file} is not a schema-1 s194 revert file ("
            in capsys.readouterr().out)
    assert asked == []


def test_the_revert_file_is_fsynced_with_its_directory(pg_dsn, prod, tmp_path, monkeypatch):
    synced = []
    real = os.fsync

    def spy(fd):
        synced.append("dir" if stat.S_ISDIR(os.fstat(fd).st_mode) else "file")
        return real(fd)

    monkeypatch.setattr(s194.os, "fsync", spy)
    assert _run(pg_dsn, "--apply", "--revert-file", tmp_path / "r.json") == 0
    assert synced == ["file", "dir"]


@pytest.mark.parametrize("failing", ["file", "dir"])
def test_a_revert_file_that_cannot_be_made_durable_rolls_back(
        pg_dsn, prod, tmp_path, capsys, monkeypatch, failing):
    real = os.fsync

    def flaky(fd):
        if ("dir" if stat.S_ISDIR(os.fstat(fd).st_mode) else "file") == failing:
            raise OSError(5, "Input/output error")
        return real(fd)

    monkeypatch.setattr(s194.os, "fsync", flaky)
    before = _snapshot(pg_dsn)
    revert_file = tmp_path / "r.json"
    assert _run(pg_dsn, "--apply", "--revert-file", revert_file) == 1
    assert f"apply: could not write {revert_file}" in capsys.readouterr().out
    assert _snapshot(pg_dsn) == before
    assert sorted(p.name for p in tmp_path.iterdir() if p.name.startswith("r.json")) == []
