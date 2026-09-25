"""Tests for verify_phase5b1 gates.

Path-only gates — no Postgres access. Fixture site dirs built in tmp_path.
The one real-PDF jbook_pdf citation test uses tests/fixtures/jbooks/darpa_p24_25.pdf
and obtains the stored bbox by calling find_fact_page live.

Gate 1: citation_gate5b1 — stratified sample, 100% re-derivation required.
Gate 2: integrity_gate5b1 — per-kind set checks + manifest reconciliation.
Gate 3: coverage_report5b1 — non-gating coverage stats.
"""
from __future__ import annotations

import hashlib
import json
import shutil
from decimal import Decimal
from pathlib import Path

import duckdb
import pytest
from openpyxl import Workbook

from govbudget.export_site import (
    fact_id_jbook,
    fact_id_lda,
    fact_id_narrative,
    fact_id_workbook,
)
from govbudget.verify_phase5b1 import (
    citation_gate5b1,
    coverage_report5b1,
    integrity_gate5b1,
)

FIXTURE_PDF = Path(__file__).resolve().parent / "fixtures" / "jbooks" / "darpa_p24_25.pdf"
P40_FIXTURE_PDF = (
    Path(__file__).resolve().parent / "fixtures" / "jbooks" / "p40_resource_summary.pdf"
)


def _page_words(pdf_path: Path, page_number: int = 1) -> list[dict]:
    """Real pdfplumber words for one page — never a mocked list."""
    import pdfplumber

    with pdfplumber.open(pdf_path) as pdf:
        return pdf.pages[page_number - 1].extract_words()


def _words_naming(*texts: str) -> list[dict]:
    """A throwaway page-word list that prints exactly these tokens.

    The row-label leg accepts a BARE summary row only when the fact's pe_bli
    is printed on the page, so every _label_names_fact call now needs a page.
    """
    return [
        {"text": t, "x0": 0.0, "x1": 10.0, "top": 0.0, "bottom": 9.0} for t in texts
    ]


def _P0_1_WORDS() -> list[dict]:
    """P0-1's three real P-40 Resource Summary rows, laid out as words.

    FY26 Air Force Aircraft Procurement Vol I, p.55: Net Procurement (P-1)
    5,247.070 + Plus CY Advance Procurement 318.585 = Total Obligation
    Authority 5,565.655. The site publishes TOA; the citation resolved to Net
    Procurement. Same document, same hash, same page, same column. (The block's
    Gross/Weapon System Cost row is omitted rather than invented — the spec
    does not state its value and no test needs it.)
    """
    def row(top, label_words, value):
        out, x = [], 20.0
        for t in label_words:
            out.append({"text": t, "x0": x, "x1": x + 6.0 * len(t), "top": top,
                        "bottom": top + 9.0})
            x += 6.0 * len(t) + 2.0
        out.append({"text": value, "x0": 347.1, "x1": 390.0, "top": top,
                    "bottom": top + 9.0})
        return out

    return (
        row(208.8, ["Net", "Procurement", "(P-1)", "($", "in", "Millions)"], "5,247.070")
        + row(220.7, ["Plus", "CY", "Advance", "Procurement"], "318.585")
        + row(232.9, ["Total", "Obligation", "Authority", "($", "in", "Millions)"], "5,565.655")
    )

# ---------------------------------------------------------------------------
# Helpers — build minimal fixture site dirs
# ---------------------------------------------------------------------------


def _write_parquet(path: Path, col_defs: str, rows: list[tuple]) -> None:
    """Write a typed parquet via duckdb create-table + executemany."""
    path.parent.mkdir(parents=True, exist_ok=True)
    con = duckdb.connect()
    try:
        con.execute(f"create table _t ({col_defs})")
        if rows:
            placeholders = ", ".join("?" for _ in rows[0])
            con.executemany(f"insert into _t values ({placeholders})", rows)
        path_str = str(path).replace("'", "''")
        con.execute(f"copy _t to '{path_str}' (format parquet, compression zstd)")
    finally:
        con.close()


# ---------------------------------------------------------------------------
# Common citation column definitions (updated for 5B-3 derived-tier schema)
# ---------------------------------------------------------------------------

_CIT_COL_DEFS = (
    "fact_id varchar, kind varchar, units varchar, amount_text varchar,"
    " page_number integer, x0 double, x1 double, top_pt double, bottom_pt double,"
    " page_width double, page_height double, resolution varchar,"
    " sheet varchar, cells varchar, amount_thousands double,"
    " sha256 varchar, hosted_pdf_url varchar, official_url varchar,"
    " xml_path varchar, retrieved_at varchar,"
    " formula varchar, inputs varchar, query_body varchar, recorded_value varchar,"
    " pe_bli varchar, scenario varchar, amount_type varchar"
)


def _make_workbook(path: Path, sheet_name: str = "Exhibit R-1") -> None:
    """Create a minimal XLSX with one data row."""
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    # Header row 1
    ws.append(["Account", "Account Title", "Organization", "Budget Activity",
               "Budget Activity Title", "Line Number", "PE/BLI", "Title",
               "Include In TOA", "FY 2024 Actuals"])
    # Data row 2 (J2 = 280494)
    ws.append(["0400", "RDT&E Defense-Wide", "DARPA", "01", "Basic Research",
               "2", "0601101E", "DEFENSE RESEARCH SCIENCES", "Y", 280494])
    wb.save(path)


def _make_site_with_jbook_pdf(
    site_dir: Path,
    *,
    pdf_path: Path = FIXTURE_PDF,
) -> tuple[str, int, dict]:
    """Build a minimal site dir with a real jbook_pdf citation.

    Returns (sha256, page_number, bbox_dict) from the live find_fact_page call.
    """
    from govbudget.jbooks.provenance_pages import find_fact_page

    sha = hashlib.sha256(pdf_path.read_bytes()).hexdigest()
    amount = Decimal("280.494")
    hit = find_fact_page(pdf_path, pe_bli="0601101E", amount=amount)
    assert hit["resolution"] in ("unique", "ambiguous_first"), \
        f"fixture resolution unexpected: {hit['resolution']}"

    page_n = hit["page_number"]
    x0, top_pt = hit["x0"], hit["top_pt"]
    amount_text = hit["amount_text"]

    # sha-named PDF copy
    pdfs_dir = site_dir / "pdfs"
    pdfs_dir.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(pdf_path, pdfs_dir / f"{sha}.pdf")

    fid = fact_id_jbook(sha, "0601101E", None, "PriorYear", "280.494")

    # jbook_details.parquet
    _write_parquet(
        site_dir / "data" / "jbook_details.parquet",
        "fact_id varchar, pe_bli varchar, project_number varchar, project_title varchar,"
        " scenario varchar, amount_millions double, units varchar, xml_path varchar,"
        " org varchar, exhibit_family varchar, fiscal_year integer,"
        " document_sha256 varchar, resolution varchar",
        [(fid, "0601101E", None, "Defense Research Sciences", "PriorYear",
          280.494, "USD millions", "ProgramElement[0]",
          "DARPA", "rdte", 2026, sha, hit["resolution"])],
    )

    # citations.parquet with jbook_pdf row
    _write_parquet(
        site_dir / "citations" / "citations.parquet",
        _CIT_COL_DEFS,
        [(fid, "jbook_pdf", "USD millions", amount_text, page_n,
          float(hit["x0"]), float(hit["x1"]),
          float(hit["top_pt"]), float(hit["bottom_pt"]),
          float(hit["page_width"]), float(hit["page_height"]),
          hit["resolution"],
          None, None, None,
          sha,
          f"https://cdn.example/pdfs/{sha}.pdf#page={page_n}",
          f"https://example.mil/darpa.pdf#page={page_n}",
          None, None,
          None, None, None, None,  # formula, inputs, query_body, recorded_value
          None, None, None)],  # pe_bli, scenario, amount_type
    )

    return sha, page_n, hit


def _make_site_with_jbook_pdf_at_word(
    site_dir: Path,
    *,
    amount_text: str,
    project_number: str | None,
    project_title: str | None,
    exhibit_family: str = "rdte",
    pdf_path: Path = FIXTURE_PDF,
    pe_bli: str = "0601101E",
    write_detail: bool = True,
    null_bbox: bool = False,
    top_offset: float = 0.0,
) -> str:
    """Site whose one jbook_pdf citation highlights `amount_text` on page 1.

    Unlike _make_site_with_jbook_pdf this does not go through find_fact_page:
    it points the citation at a word chosen BY NAME, so a test can put the
    highlight on a row belonging to a different project and prove the
    row-label leg catches it. The word must be UNIQUE on the page, because
    _verify_jbook_pdf finds the first occurrence of amount_text and checks the
    stored bbox against THAT word — uniqueness is what makes the pre-existing
    x0/top_pt check pass, so the label leg is the only thing under test.
    """
    import pdfplumber

    sha = hashlib.sha256(pdf_path.read_bytes()).hexdigest()
    with pdfplumber.open(pdf_path) as pdf:
        page = pdf.pages[0]
        words = page.extract_words()
        page_w, page_h = float(page.width), float(page.height)
    hits = [w for w in words if w["text"] == amount_text]
    assert len(hits) == 1, f"fixture word {amount_text!r} is not unique: {len(hits)}"
    w = hits[0]

    pdfs_dir = site_dir / "pdfs"
    pdfs_dir.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(pdf_path, pdfs_dir / f"{sha}.pdf")

    fid = fact_id_jbook(sha, pe_bli, project_number, "PriorYear", amount_text)
    _write_parquet(
        site_dir / "data" / "jbook_details.parquet",
        "fact_id varchar, pe_bli varchar, project_number varchar, project_title varchar,"
        " scenario varchar, amount_millions double, units varchar, xml_path varchar,"
        " org varchar, exhibit_family varchar, fiscal_year integer,"
        " document_sha256 varchar, resolution varchar",
        # write_detail=False leaves the details parquet EMPTY: the citation is
        # then a jbook_pdf row with no detail to check against, which is the
        # skipped_no_detail path.
        [(fid, pe_bli, project_number, project_title, "PriorYear",
          float(amount_text), "USD millions", "ProgramElement[0]",
          "DARPA", exhibit_family, 2026, sha, "unique")] if write_detail else [],
    )
    _write_parquet(
        site_dir / "citations" / "citations.parquet",
        _CIT_COL_DEFS,
        [(fid, "jbook_pdf", "USD millions", amount_text, 1,
          None if null_bbox else float(w["x0"]), float(w["x1"]),
          None if null_bbox else float(w["top"]) + top_offset, float(w["bottom"]),
          page_w, page_h, "unique",
          None, None, None, sha,
          f"https://cdn.example/pdfs/{sha}.pdf#page=1",
          "https://example.mil/darpa.pdf#page=1",
          None, None, None, None, None, None, None, None, None)],
    )
    return fid


def _make_site_with_workbook(site_dir: Path) -> tuple[str, str]:
    """Build site dir with a workbook citation.

    Returns (sha256, fact_id).
    """
    # Create XLSX fixture
    wb_src = site_dir / "_src_r1.xlsx"
    _make_workbook(wb_src)
    sha = hashlib.sha256(wb_src.read_bytes()).hexdigest()

    wb_dir = site_dir / "workbooks"
    wb_dir.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(wb_src, wb_dir / f"{sha}.xlsx")

    fid = fact_id_workbook(sha, "R-1", 2026, "0400", "DARPA", "01", "0601101E", "fy_2024_actuals")

    # budget_lines.parquet
    _write_parquet(
        site_dir / "data" / "budget_lines.parquet",
        "fact_id varchar, exhibit varchar, fiscal_year integer, account varchar,"
        " account_title varchar, organization varchar, budget_activity varchar,"
        " budget_activity_title varchar, pe_bli varchar, title varchar,"
        " amount_type varchar, amount_thousands double, units varchar,"
        " document_sha256 varchar, source_sheet varchar, source_cells varchar",
        [(fid, "R-1", 2026, "0400", "RDT&E Defense-Wide", "DARPA", "01",
          "Basic Research", "0601101E", "DEFENSE RESEARCH SCIENCES",
          "fy_2024_actuals", 280494.0, "USD thousands",
          sha, "Exhibit R-1", "J2")],
    )

    # citations.parquet with workbook row
    _write_parquet(
        site_dir / "citations" / "citations.parquet",
        _CIT_COL_DEFS,
        [(fid, "workbook", "USD thousands", None, None,
          None, None, None, None, None, None, None,
          "Exhibit R-1", "J2", 280494.0,
          sha, None, "https://example.mil/r1.xlsx", None, None,
          None, None, None, None,  # formula, inputs, query_body, recorded_value
          None, None, None)],  # pe_bli, scenario, amount_type
    )

    return sha, fid


def _make_site_with_lda(site_dir: Path) -> tuple[str, str]:
    """Build site dir with an lda_filing citation.

    Returns (filing_uuid, fact_id).
    """
    filing_uuid = "a1b2c3d4-e5f6-7890-abcd-ef1234567890"
    pe_bli = "0601101E"
    matched_term = "darpa"
    fid = fact_id_lda(filing_uuid, pe_bli, matched_term)
    filing_url = f"https://lda.senate.gov/filings/{filing_uuid}/"

    # fct_program_lobbying.parquet  (needed for integrity gate re-derive check)
    _write_parquet(
        site_dir / "data" / "fct_program_lobbying.parquet",
        "filing_uuid varchar, pe_bli varchar, program_title varchar,"
        " matched_term varchar, description_snippet varchar,"
        " filing_url varchar, client_name varchar,"
        " family_key varchar, filing_year varchar",
        [(filing_uuid, pe_bli, "Defense Research Sciences", matched_term,
          "mentioned DARPA", filing_url, "Lockheed Martin", "lockheed", "2025")],
    )

    _write_parquet(
        site_dir / "citations" / "citations.parquet",
        _CIT_COL_DEFS,
        [(fid, "lda_filing", None, None, None,
          None, None, None, None, None, None, None,
          None, None, None,
          None, None, filing_url, None, None,
          None, None, None, None,  # formula, inputs, query_body, recorded_value
          None, None, None)],  # pe_bli, scenario, amount_type
    )

    return filing_uuid, fid


def _write_manifest(site_dir: Path, **kwargs) -> None:
    """Write a minimal manifest.json."""
    defaults = {
        "built_at": "2026-06-12T00:00:00+00:00",
        "datasets": {},
        "citations": {},
        "skipped_unresolved": 0,
        "skipped_zero_amount": 0,
        "uncited_datasets": [],
        "pdf_base_url": "/pdfs",
        "schema_version": 1,
    }
    defaults.update(kwargs)
    (site_dir / "manifest.json").write_text(json.dumps(defaults, indent=2))


# ---------------------------------------------------------------------------
# Gate 1: citation_gate5b1 — happy paths
# ---------------------------------------------------------------------------


class TestCitationGate:
    def test_jbook_pdf_happy_path(self, tmp_path):
        """Real PDF fixture — bbox within 2 pt of stored values."""
        site = tmp_path / "site"
        sha, page_n, hit = _make_site_with_jbook_pdf(site)
        _write_manifest(site, datasets={"jbook_details": 1}, citations={"jbook_pdf": 1})

        result = citation_gate5b1(site)

        assert result["ok"] is True, f"failures: {result.get('failures')}"
        assert result["sampled"] >= 1
        assert result["passed"] == result["sampled"]
        assert result["failures"] == []

    def test_workbook_happy_path(self, tmp_path):
        """Workbook cell sum must match stored amount_thousands."""
        site = tmp_path / "site"
        _make_site_with_workbook(site)
        _write_manifest(site, datasets={"budget_lines": 1}, citations={"workbook": 1})

        result = citation_gate5b1(site)

        assert result["ok"] is True, f"failures: {result.get('failures')}"
        assert result["sampled"] >= 1
        assert result["passed"] == result["sampled"]

    def test_lda_filing_happy_path(self, tmp_path):
        """LDA shape check — url starts with lda.senate.gov and contains uuid."""
        site = tmp_path / "site"
        filing_uuid, fid = _make_site_with_lda(site)
        _write_manifest(site, citations={"lda_filing": 1})

        result = citation_gate5b1(site)

        assert result["ok"] is True, f"failures: {result.get('failures')}"

    def test_empty_citations_fails(self, tmp_path):
        """Empty citations.parquet → ok=False."""
        site = tmp_path / "site"
        (site / "citations").mkdir(parents=True)
        _write_parquet(
            site / "citations" / "citations.parquet",
            _CIT_COL_DEFS,
            [],
        )
        _write_manifest(site)

        result = citation_gate5b1(site)

        assert result["ok"] is False
        assert result["sampled"] == 0

    def test_missing_site_dir_fails(self, tmp_path):
        result = citation_gate5b1(tmp_path / "nonexistent")
        assert result["ok"] is False
        assert "reason" in result

    def test_tampered_amount_text_fails(self, tmp_path):
        """Storing wrong amount_text → PDF word search fails → gate FAIL."""
        site = tmp_path / "site"
        sha, page_n, hit = _make_site_with_jbook_pdf(site)
        fid = fact_id_jbook(sha, "0601101E", None, "PriorYear", "280.494")

        # Re-write citations with a wrong amount_text
        (site / "citations").mkdir(parents=True, exist_ok=True)
        _write_parquet(
            site / "citations" / "citations.parquet",
            _CIT_COL_DEFS,
            [(fid, "jbook_pdf", "USD millions", "999999.000", page_n,
              float(hit["x0"]), float(hit["x1"]),
              float(hit["top_pt"]), float(hit["bottom_pt"]),
              float(hit["page_width"]), float(hit["page_height"]),
              hit["resolution"],
              None, None, None,
              sha,
              f"https://cdn.example/pdfs/{sha}.pdf#page={page_n}",
              f"https://example.mil/darpa.pdf#page={page_n}",
              None, None,
              None, None, None, None,
              None, None, None)],
        )
        _write_manifest(site, citations={"jbook_pdf": 1})

        result = citation_gate5b1(site)

        assert result["ok"] is False
        assert len(result["failures"]) >= 1
        assert any(fid == f[0] for f in result["failures"]), \
            f"expected {fid} in failures, got: {result['failures']}"

    def test_tampered_workbook_amount_fails(self, tmp_path):
        """Stored amount_thousands differs from cell value → gate FAIL."""
        site = tmp_path / "site"
        sha, fid = _make_site_with_workbook(site)

        # Re-write citations with wrong amount_thousands (does NOT match cell J2=280494)
        (site / "citations").mkdir(parents=True, exist_ok=True)
        _write_parquet(
            site / "citations" / "citations.parquet",
            _CIT_COL_DEFS,
            [(fid, "workbook", "USD thousands", None, None,
              None, None, None, None, None, None, None,
              "Exhibit R-1", "J2", 999999.0,   # WRONG amount
              sha, None, "https://example.mil/r1.xlsx", None, None,
              None, None, None, None,
              None, None, None)],
        )
        _write_manifest(site, citations={"workbook": 1})

        result = citation_gate5b1(site)

        assert result["ok"] is False
        assert len(result["failures"]) >= 1

    def test_lda_bad_url_fails(self, tmp_path):
        """LDA citation with non-lda.senate.gov URL → gate FAIL."""
        site = tmp_path / "site"
        fid = fact_id_lda("uuid-bad", "0601101E", "darpa")
        (site / "citations").mkdir(parents=True)
        _write_parquet(
            site / "citations" / "citations.parquet",
            _CIT_COL_DEFS,
            [(fid, "lda_filing", None, None, None,
              None, None, None, None, None, None, None,
              None, None, None,
              None, None, "https://not-lda.example.com/f/1", None, None,
              None, None, None, None,
              None, None, None)],
        )
        _write_manifest(site, citations={"lda_filing": 1})

        result = citation_gate5b1(site)

        assert result["ok"] is False
        assert len(result["failures"]) >= 1

    def test_lda_url_without_uuid_fails(self, tmp_path):
        """LDA citation whose official_url has no UUID in the path → gate FAIL."""
        site = tmp_path / "site"
        fid = fact_id_lda("no-uuid-here", "0601101E", "darpa")
        (site / "citations").mkdir(parents=True)
        _write_parquet(
            site / "citations" / "citations.parquet",
            _CIT_COL_DEFS,
            [(fid, "lda_filing", None, None, None,
              None, None, None, None, None, None, None,
              None, None, None,
              None, None, "https://lda.senate.gov/filings/no-uuid-here/", None, None,
              None, None, None, None,
              None, None, None)],
        )
        _write_manifest(site, citations={"lda_filing": 1})

        result = citation_gate5b1(site)

        assert result["ok"] is False, \
            "expected FAIL when official_url has no filing UUID"
        assert len(result["failures"]) >= 1
        # The failure message must identify the fact_id
        assert any(fid in str(f) for f in result["failures"]), \
            f"expected {fid} in failures, got: {result['failures']}"


# ---------------------------------------------------------------------------
# Stratified sampling guarantee
# ---------------------------------------------------------------------------


class TestStratifiedSampling:
    def test_stratification_guarantees_min_per_kind(self, tmp_path):
        """100 jbook rows + 2 workbook + 2 lda, sample_size=50.

        Every kind must appear in sample; workbook and lda must each get
        ALL their rows (2 each); jbook gets the remainder (46).

        The test detects the greedy bug by monkey-patching _verify_* to record
        which fact_ids were actually sampled, then checking per-kind counts.
        """
        site = tmp_path / "site"

        # Build a single valid workbook fixture (for the workbook citations)
        wb_src = site / "_r1.xlsx"
        _make_workbook(wb_src)
        sha_wb = hashlib.sha256(wb_src.read_bytes()).hexdigest()
        (site / "workbooks").mkdir(parents=True, exist_ok=True)
        shutil.copyfile(wb_src, site / "workbooks" / f"{sha_wb}.xlsx")

        # Build a single valid PDF fixture (for the jbook citations)
        sha_pdf = hashlib.sha256(FIXTURE_PDF.read_bytes()).hexdigest()
        (site / "pdfs").mkdir(parents=True, exist_ok=True)
        shutil.copyfile(FIXTURE_PDF, site / "pdfs" / f"{sha_pdf}.pdf")

        from govbudget.jbooks.provenance_pages import find_fact_page
        hit = find_fact_page(FIXTURE_PDF, pe_bli="0601101E", amount=Decimal("280.494"))

        # 100 jbook_pdf rows (same sha/pe_bli/page — same physical citation; fact_id
        # differs by a unique suffix baked into fact_id_jbook via scenario field)
        jbook_rows = []
        for i in range(100):
            fid = fact_id_jbook(sha_pdf, "0601101E", None, f"Scenario{i:04d}", "280.494")
            jbook_rows.append((
                fid, "jbook_pdf", "USD millions", hit["amount_text"],
                hit["page_number"],
                float(hit["x0"]), float(hit["x1"]),
                float(hit["top_pt"]), float(hit["bottom_pt"]),
                float(hit["page_width"]), float(hit["page_height"]),
                hit["resolution"],
                None, None, None,
                sha_pdf,
                f"https://cdn.example/pdfs/{sha_pdf}.pdf#page={hit['page_number']}",
                f"https://example.mil/darpa.pdf#page={hit['page_number']}",
                None, None,
                None, None, None, None,  # formula, inputs, query_body, recorded_value
                None, None, None,        # pe_bli, scenario, amount_type
            ))

        # 2 workbook rows
        wb_rows = []
        for i in range(2):
            fid = fact_id_workbook(sha_wb, "R-1", 2026, "0400", f"ORG{i}", "01", "0601101E", "fy_2024_actuals")
            wb_rows.append((
                fid, "workbook", "USD thousands", None, None,
                None, None, None, None, None, None, None,
                "Exhibit R-1", "J2", 280494.0,
                sha_wb, None, "https://example.mil/r1.xlsx", None, None,
                None, None, None, None,  # formula, inputs, query_body, recorded_value
                None, None, None,        # pe_bli, scenario, amount_type
            ))

        # 2 lda_filing rows (use a real UUID so the UUID check passes)
        lda_rows = []
        real_uuid = "a1b2c3d4-e5f6-7890-abcd-ef1234567890"
        for i in range(2):
            fid = fact_id_lda(f"uuid-lda-{i:03d}", "0601101E", "darpa")
            lda_rows.append((
                fid, "lda_filing", None, None, None,
                None, None, None, None, None, None, None,
                None, None, None,
                None, None, f"https://lda.senate.gov/filings/{real_uuid}/", None, None,
                None, None, None, None,  # formula, inputs, query_body, recorded_value
                None, None, None,        # pe_bli, scenario, amount_type
            ))

        all_rows = jbook_rows + wb_rows + lda_rows

        _write_parquet(site / "citations" / "citations.parquet", _CIT_COL_DEFS, all_rows)
        _write_manifest(site)

        # Build a fact_id → kind lookup for asserting per-kind coverage
        wb_fact_ids = {r[0] for r in wb_rows}
        lda_fact_ids = {r[0] for r in lda_rows}
        jbook_fact_ids = {r[0] for r in jbook_rows}

        # Monkey-patch the _verify_* functions to record sampled fact_ids
        import govbudget.verify_phase5b1 as _mod
        sampled_fact_ids: list[str] = []
        _orig_jbook = _mod._verify_jbook_pdf
        _orig_wb = _mod._verify_workbook
        _orig_lda = _mod._verify_lda

        def _record_jbook(site_dir, row, idx, *args, **kwargs):
            sampled_fact_ids.append(row[idx["fact_id"]])
            return _orig_jbook(site_dir, row, idx, *args, **kwargs)

        def _record_wb(site_dir, row, idx):
            sampled_fact_ids.append(row[idx["fact_id"]])
            return _orig_wb(site_dir, row, idx)

        def _record_lda(row, idx):
            sampled_fact_ids.append(row[idx["fact_id"]])
            return _orig_lda(row, idx)

        _mod._verify_jbook_pdf = _record_jbook
        _mod._verify_workbook = _record_wb
        _mod._verify_lda = _record_lda
        try:
            result = citation_gate5b1(site, sample_size=50)
        finally:
            _mod._verify_jbook_pdf = _orig_jbook
            _mod._verify_workbook = _orig_wb
            _mod._verify_lda = _orig_lda

        assert result["ok"] is True, f"gate failures: {result.get('failures')}"
        assert result["sampled"] == 50

        sampled_set = set(sampled_fact_ids)
        n_jbook_sampled = len(sampled_set & jbook_fact_ids)
        n_wb_sampled = len(sampled_set & wb_fact_ids)
        n_lda_sampled = len(sampled_set & lda_fact_ids)

        # workbook and lda have only 2 rows each → must both be fully sampled
        assert n_wb_sampled == 2, \
            f"workbook under-sampled: got {n_wb_sampled}/2 (greedy bug?)"
        assert n_lda_sampled == 2, \
            f"lda under-sampled: got {n_lda_sampled}/2 (greedy bug?)"
        assert n_jbook_sampled == 46, \
            f"jbook count wrong: got {n_jbook_sampled}, expected 46 (= 50 - 2 - 2)"


# ---------------------------------------------------------------------------
# Gate 2: integrity_gate5b1
# ---------------------------------------------------------------------------


class TestIntegrityGate:
    def _make_full_site(self, site_dir: Path) -> dict:
        """Build a minimal but complete site with jbook, workbook, and lda citations."""
        # ---- jbook ----
        sha_pdf = hashlib.sha256(FIXTURE_PDF.read_bytes()).hexdigest()
        from govbudget.jbooks.provenance_pages import find_fact_page
        hit = find_fact_page(FIXTURE_PDF, pe_bli="0601101E", amount=Decimal("280.494"))
        page_n = hit["page_number"]
        fid_jbook = fact_id_jbook(sha_pdf, "0601101E", None, "PriorYear", "280.494")

        pdfs_dir = site_dir / "pdfs"
        pdfs_dir.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(FIXTURE_PDF, pdfs_dir / f"{sha_pdf}.pdf")

        # ---- workbook ----
        wb_src = site_dir / "_r1.xlsx"
        _make_workbook(wb_src)
        sha_wb = hashlib.sha256(wb_src.read_bytes()).hexdigest()
        wb_dir = site_dir / "workbooks"
        wb_dir.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(wb_src, wb_dir / f"{sha_wb}.xlsx")
        fid_wb = fact_id_workbook(sha_wb, "R-1", 2026, "0400", "DARPA", "01", "0601101E", "fy_2024_actuals")

        # ---- lda ----
        fid_lda = fact_id_lda("uuid-001", "0601101E", "darpa")

        # ---- data parquets ----
        _write_parquet(
            site_dir / "data" / "jbook_details.parquet",
            "fact_id varchar, pe_bli varchar, project_number varchar, project_title varchar,"
            " scenario varchar, amount_millions double, units varchar, xml_path varchar,"
            " org varchar, exhibit_family varchar, fiscal_year integer,"
            " document_sha256 varchar, resolution varchar",
            [(fid_jbook, "0601101E", None, "Defense Research Sciences", "PriorYear",
              280.494, "USD millions", "ProgramElement[0]",
              "DARPA", "rdte", 2026, sha_pdf, hit["resolution"])],
        )
        _write_parquet(
            site_dir / "data" / "budget_lines.parquet",
            "fact_id varchar, exhibit varchar, fiscal_year integer, account varchar,"
            " account_title varchar, organization varchar, budget_activity varchar,"
            " budget_activity_title varchar, pe_bli varchar, title varchar,"
            " amount_type varchar, amount_thousands double, units varchar,"
            " document_sha256 varchar, source_sheet varchar, source_cells varchar",
            [(fid_wb, "R-1", 2026, "0400", "RDT&E Defense-Wide", "DARPA", "01",
              "Basic Research", "0601101E", "DEFENSE RESEARCH SCIENCES",
              "fy_2024_actuals", 280494.0, "USD thousands",
              sha_wb, "Exhibit R-1", "J2")],
        )
        _write_parquet(
            site_dir / "data" / "fct_program_lobbying.parquet",
            "filing_uuid varchar, pe_bli varchar, program_title varchar,"
            " matched_term varchar, description_snippet varchar,"
            " filing_url varchar, client_name varchar,"
            " family_key varchar, filing_year varchar",
            [("uuid-001", "0601101E", "Defense Research Sciences", "darpa",
              "mentioned DARPA", "https://lda.senate.gov/filings/uuid-001/",
              "Lockheed Martin", "lockheed", "2025")],
        )

        # ---- citations.parquet ----
        cit_rows = [
            (fid_jbook, "jbook_pdf", "USD millions", hit["amount_text"],
             page_n, float(hit["x0"]), float(hit["x1"]),
             float(hit["top_pt"]), float(hit["bottom_pt"]),
             float(hit["page_width"]), float(hit["page_height"]),
             hit["resolution"],
             None, None, None,
             sha_pdf,
             f"https://cdn.example/pdfs/{sha_pdf}.pdf#page={page_n}",
             f"https://example.mil/darpa.pdf#page={page_n}",
             None, None,
             None, None, None, None,  # formula, inputs, query_body, recorded_value
             None, None, None),       # pe_bli, scenario, amount_type
            (fid_wb, "workbook", "USD thousands", None, None,
             None, None, None, None, None, None, None,
             "Exhibit R-1", "J2", 280494.0,
             sha_wb, None, "https://example.mil/r1.xlsx", None, None,
             None, None, None, None,  # formula, inputs, query_body, recorded_value
             None, None, None),       # pe_bli, scenario, amount_type
            (fid_lda, "lda_filing", None, None, None,
             None, None, None, None, None, None, None,
             None, None, None,
             None, None, "https://lda.senate.gov/filings/uuid-001/", None, None,
             None, None, None, None,  # formula, inputs, query_body, recorded_value
             None, None, None),       # pe_bli, scenario, amount_type
        ]
        _write_parquet(
            site_dir / "citations" / "citations.parquet",
            _CIT_COL_DEFS,
            cit_rows,
        )

        # ---- manifest.json ----
        _write_manifest(
            site_dir,
            datasets={"jbook_details": 1, "budget_lines": 1, "fct_program_lobbying": 1},
            citations={"jbook_pdf": 1, "workbook": 1, "lda_filing": 1},
            skipped_unresolved=0,
            skipped_zero_amount=0,
        )

        return {
            "fid_jbook": fid_jbook, "fid_wb": fid_wb, "fid_lda": fid_lda,
            "sha_pdf": sha_pdf, "sha_wb": sha_wb,
        }

    def test_happy_path(self, tmp_path):
        site = tmp_path / "site"
        self._make_full_site(site)
        result = integrity_gate5b1(site)
        assert result["ok"] is True, f"failures: {result}"

    def test_orphan_jbook_citation_fails(self, tmp_path):
        """A jbook_pdf citation whose fact_id is NOT in jbook_details → FAIL."""
        site = tmp_path / "site"
        ids = self._make_full_site(site)

        # Add an extra jbook_pdf citation with an orphan fact_id
        orphan_fid = "orphan0000000000"
        sha = ids["sha_pdf"]
        cit_pq = site / "citations" / "citations.parquet"
        con = duckdb.connect()
        existing = con.execute(f"select * from read_parquet('{cit_pq}')").fetchall()
        cols = [d[0] for d in con.execute(f"describe select * from read_parquet('{cit_pq}')").fetchall()]
        con.close()

        from govbudget.jbooks.provenance_pages import find_fact_page
        hit = find_fact_page(FIXTURE_PDF, pe_bli="0601101E", amount=Decimal("280.494"))
        extra_row = (orphan_fid, "jbook_pdf", "USD millions", hit["amount_text"],
                     hit["page_number"],
                     float(hit["x0"]), float(hit["x1"]),
                     float(hit["top_pt"]), float(hit["bottom_pt"]),
                     float(hit["page_width"]), float(hit["page_height"]),
                     hit["resolution"],
                     None, None, None,
                     sha,
                     f"https://cdn.example/pdfs/{sha}.pdf#page={hit['page_number']}",
                     "https://example.mil/darpa.pdf#page=1",
                     None, None,
                     None, None, None, None,  # formula, inputs, query_body, recorded_value
                     None, None, None)        # pe_bli, scenario, amount_type
        _write_parquet(cit_pq, _CIT_COL_DEFS, existing + [extra_row])

        result = integrity_gate5b1(site)
        assert result["ok"] is False

    def test_resolved_details_row_missing_citation_fails(self, tmp_path):
        """A jbook_details row with resolution='unique' missing from citations → FAIL."""
        site = tmp_path / "site"
        ids = self._make_full_site(site)

        # Add another details row with a new fact_id (resolution=unique) that has no citation
        details_pq = site / "data" / "jbook_details.parquet"
        missing_fid = "missing00000000"
        con = duckdb.connect()
        existing = con.execute(f"select * from read_parquet('{details_pq}')").fetchall()
        extra = (missing_fid, "0602303E", None, "ICT Research", "BudgetYearOne",
                 100.0, "USD millions", "PE[1]", "DARPA", "rdte", 2026,
                 ids["sha_pdf"], "unique")
        _write_parquet(
            details_pq,
            "fact_id varchar, pe_bli varchar, project_number varchar, project_title varchar,"
            " scenario varchar, amount_millions double, units varchar, xml_path varchar,"
            " org varchar, exhibit_family varchar, fiscal_year integer,"
            " document_sha256 varchar, resolution varchar",
            existing + [extra],
        )
        con.close()

        result = integrity_gate5b1(site)
        assert result["ok"] is False

    def test_manifest_rowcount_mismatch_fails(self, tmp_path):
        """manifest.json dataset count that doesn't match actual parquet → FAIL."""
        site = tmp_path / "site"
        self._make_full_site(site)

        # Overwrite manifest with wrong count for jbook_details
        _write_manifest(
            site,
            datasets={"jbook_details": 999, "budget_lines": 1, "fct_program_lobbying": 1},
            citations={"jbook_pdf": 1, "workbook": 1, "lda_filing": 1},
        )

        result = integrity_gate5b1(site)
        assert result["ok"] is False

    def test_workbook_fact_ids_mismatch_fails(self, tmp_path):
        """workbook citations contain a fact_id not in budget_lines → FAIL."""
        site = tmp_path / "site"
        ids = self._make_full_site(site)

        # Add extra workbook citation with unknown fact_id
        extra_fid = "extraworkbook00"
        cit_pq = site / "citations" / "citations.parquet"
        con = duckdb.connect()
        existing = con.execute(f"select * from read_parquet('{cit_pq}')").fetchall()
        con.close()
        extra = (extra_fid, "workbook", "USD thousands", None, None,
                 None, None, None, None, None, None, None,
                 "Sheet1", "A1", 100.0,
                 ids["sha_wb"], None, "https://example.mil/r1.xlsx", None, None,
                 None, None, None, None,  # formula, inputs, query_body, recorded_value
                 None, None, None)        # pe_bli, scenario, amount_type
        _write_parquet(
            cit_pq,
            _CIT_COL_DEFS,
            existing + [extra],
        )

        result = integrity_gate5b1(site)
        assert result["ok"] is False

    def test_missing_site_dir_fails(self, tmp_path):
        result = integrity_gate5b1(tmp_path / "nonexistent")
        assert result["ok"] is False

    def test_manifest_skip_count_mismatch_fails(self, tmp_path):
        """manifest skipped_unresolved is higher than parquet count → integrity FAIL."""
        site = tmp_path / "site"
        self._make_full_site(site)

        # Tamper: set skipped_unresolved to 1 higher than actual (parquet has 0 unresolved)
        man = json.loads((site / "manifest.json").read_text())
        man["skipped_unresolved"] = man.get("skipped_unresolved", 0) + 1
        (site / "manifest.json").write_text(json.dumps(man, indent=2))

        result = integrity_gate5b1(site)
        assert result["ok"] is False, \
            "expected FAIL when manifest skipped_unresolved doesn't match parquet counts"
        assert any("skip" in f.lower() or "skipped" in f.lower() or "unresolved" in f.lower()
                   for f in result["failures"]), \
            f"expected skip-count failure description, got: {result['failures']}"


# ---------------------------------------------------------------------------
# Gate 3: coverage_report5b1 (non-gating)
# ---------------------------------------------------------------------------


class TestCoverageReport:
    def test_returns_non_failing_dict(self, tmp_path):
        site = tmp_path / "site"
        sha = hashlib.sha256(FIXTURE_PDF.read_bytes()).hexdigest()
        fid = fact_id_jbook(sha, "0601101E", None, "PriorYear", "280.494")
        _write_parquet(
            site / "data" / "jbook_details.parquet",
            "fact_id varchar, pe_bli varchar, project_number varchar, project_title varchar,"
            " scenario varchar, amount_millions double, units varchar, xml_path varchar,"
            " org varchar, exhibit_family varchar, fiscal_year integer,"
            " document_sha256 varchar, resolution varchar",
            [(fid, "0601101E", None, "DR Sciences", "PriorYear",
              280.494, "USD millions", "PE[0]", "DARPA", "rdte", 2026,
              sha, "unique"),
             ("fid2", "0601102E", None, "Another", "BudgetYearOne",
              0.0, "USD millions", "PE[1]", "DARPA", "rdte", 2026,
              sha, "zero_amount"),
             ("fid3", "0601103E", None, "Yet Another", "BudgetYearOne",
              50.0, "USD millions", "PE[2]", "DARPA", "rdte", 2026,
              sha, "unresolved")],
        )
        _write_manifest(site, datasets={"jbook_details": 3}, uncited_datasets=["fct_budget_trajectory"])

        report = coverage_report5b1(site)

        # Non-gating: always returns a dict (never raises on valid input)
        assert isinstance(report, dict)
        assert "unique" in report or "resolved" in report or "by_resolution" in report
        # uncited_datasets echoed
        assert "uncited_datasets" in report

    def test_missing_site_dir_returns_dict(self, tmp_path):
        """Non-gating: missing dir returns a dict with ok=False or a summary."""
        report = coverage_report5b1(tmp_path / "nonexistent")
        assert isinstance(report, dict)


# ---------------------------------------------------------------------------
# Gate 2 sub-check: citation_distinctness
# ---------------------------------------------------------------------------


class TestCitationDistinctnessCheck:
    """citation_distinctness: raw row count == count(distinct fact_id) per kind."""

    _CIT_COLS = _CIT_COL_DEFS

    def _lda_row(self, fid: str, uuid: str) -> tuple:
        url = f"https://lda.senate.gov/filings/{uuid}/"
        return (fid, "lda_filing", None, None, None,
                None, None, None, None, None, None, None,
                None, None, None, None, None, url, None, None,
                None, None, None, None,  # formula, inputs, query_body, recorded_value
                None, None, None)        # pe_bli, scenario, amount_type

    def test_duplicated_lda_fact_id_fails(self, tmp_path):
        """citations.parquet with the same lda_filing fact_id twice → FAIL."""
        site = tmp_path / "site"
        uuid = "a1b2c3d4-e5f6-7890-abcd-ef1234567890"
        fid = fact_id_lda(uuid, "0601101E", "darpa")

        # Two rows with identical fact_id (simulates join fan-out from dup filing row)
        rows = [
            self._lda_row(fid, uuid),
            self._lda_row(fid, uuid),  # duplicate
        ]
        _write_parquet(site / "citations" / "citations.parquet", self._CIT_COLS, rows)
        _write_manifest(site)

        result = integrity_gate5b1(site)

        assert result["ok"] is False, "expected FAIL when lda_filing has duplicate fact_id"
        assert result["checks"].get("citation_distinctness") is False, (
            f"citation_distinctness check should be False, got: {result['checks']}"
        )
        assert any("citation_distinctness" in f for f in result["failures"]), (
            f"expected citation_distinctness in failures, got: {result['failures']}"
        )
        assert any("lda_filing" in f for f in result["failures"]), (
            f"expected 'lda_filing' named in failure, got: {result['failures']}"
        )

    def test_all_distinct_passes(self, tmp_path):
        """citations.parquet with two different lda_filing fact_ids → PASS."""
        site = tmp_path / "site"
        uuid1 = "a1b2c3d4-e5f6-7890-abcd-ef1234567890"
        uuid2 = "b2c3d4e5-f6a7-8901-bcde-f12345678901"
        fid1 = fact_id_lda(uuid1, "0601101E", "darpa")
        fid2 = fact_id_lda(uuid2, "0601101E", "darpa")

        rows = [
            self._lda_row(fid1, uuid1),
            self._lda_row(fid2, uuid2),
        ]
        _write_parquet(site / "citations" / "citations.parquet", self._CIT_COLS, rows)
        _write_manifest(site)

        result = integrity_gate5b1(site)

        assert result["checks"].get("citation_distinctness") is True, (
            f"citation_distinctness should be True for distinct rows, got: {result}"
        )


# ---------------------------------------------------------------------------
# Derived citation tier tests (Phase 5B-3)
# ---------------------------------------------------------------------------


from govbudget.export_site import fact_id_derived
from govbudget.verify_phase5b1 import _verify_derived


def _make_derived_row(fid: str, formula: str, inputs_json: str,
                      recorded_value: str, units: str = "USD thousands",
                      retrieved_at: str = "2026-06-12T00:00:00+00:00") -> tuple:
    """Build a 27-element derived citation row (includes pe_bli/scenario/amount_type)."""
    return (fid, "derived", units, None,
            None, None, None, None, None, None, None, None,
            None, None, None, None, None, None, None, retrieved_at,
            formula, inputs_json, None, recorded_value,
            None, None, None)  # pe_bli, scenario, amount_type


class TestFactIdDerived:
    """fact_id_derived is deterministic and distinct."""

    def test_stable_and_16hex(self):
        fid = fact_id_derived("trajectory", "0601101E|DARPA", "fy2026_total")
        assert len(fid) == 16
        assert fid == fact_id_derived("trajectory", "0601101E|DARPA", "fy2026_total")

    def test_differs_by_surface(self):
        a = fact_id_derived("trajectory", "k", "m")
        b = fact_id_derived("agency", "k", "m")
        assert a != b

    def test_differs_by_key(self):
        a = fact_id_derived("trajectory", "k1", "m")
        b = fact_id_derived("trajectory", "k2", "m")
        assert a != b

    def test_differs_by_metric(self):
        a = fact_id_derived("trajectory", "k", "fy2025_total")
        b = fact_id_derived("trajectory", "k", "fy2026_total")
        assert a != b

    def test_hash_prefix(self):
        import hashlib
        fid = fact_id_derived("x", "y", "z")
        expected = hashlib.sha256("derived|x|y|z".encode()).hexdigest()[:16]
        assert fid == expected


class TestDerivedCitationGate:
    """citation_gate5b1 handles derived kind correctly."""

    def test_derived_happy_path(self, tmp_path):
        """Derived row with formula + recorded_value → PASS.

        Uses the pivot formula when inputs=[] (sum formula with empty inputs fails).
        """
        site = tmp_path / "site"
        fid = fact_id_derived("trajectory", "0601101E|DARPA", "fy2026_total")
        # When no budget_lines inputs are available, emission side uses pivot formula.
        row = _make_derived_row(
            fid,
            "trajectory pivot of budget_lines (inputs unavailable for this org/type)",
            "[]", "295000.000",
        )
        _write_parquet(site / "citations" / "citations.parquet", _CIT_COL_DEFS, [row])
        _write_manifest(site)

        result = citation_gate5b1(site)
        assert result["ok"] is True, f"failures: {result.get('failures')}"
        assert result["sampled"] >= 1

    def test_derived_missing_formula_fails(self, tmp_path):
        """Derived row with null formula → FAIL."""
        site = tmp_path / "site"
        fid = fact_id_derived("agency", "DARPA", "fy2024_total_millions")
        row = _make_derived_row(fid, "", "[]", "1234.000")  # empty formula
        _write_parquet(site / "citations" / "citations.parquet", _CIT_COL_DEFS, [row])
        _write_manifest(site)

        result = citation_gate5b1(site)
        assert result["ok"] is False, "expected FAIL for empty formula"
        assert len(result["failures"]) >= 1

    def test_derived_missing_recorded_value_fails(self, tmp_path):
        """Derived row with null recorded_value → FAIL."""
        site = tmp_path / "site"
        fid = fact_id_derived("concentration", "0601101E", "hhi")
        # Build row with recorded_value=None
        row = (fid, "derived", "HHI", None,
               None, None, None, None, None, None, None, None,
               None, None, None, None, None, None, None, "2026-06-12T00:00:00",
               "sum(share_pct^2) where obligation>0", "[]", None, None,
               None, None, None)  # recorded_value=None; pe_bli, scenario, amount_type
        _write_parquet(site / "citations" / "citations.parquet", _CIT_COL_DEFS, [row])
        _write_manifest(site)

        result = citation_gate5b1(site)
        assert result["ok"] is False, "expected FAIL for null recorded_value"

    def test_derived_with_url_inputs_passes(self, tmp_path):
        """Derived row with URL inputs (per-capita) → PASS (shape check only)."""
        site = tmp_path / "site"
        fid = fact_id_derived("state_per_capita", "CA|Education", "amount_per_capita")
        import json
        inputs = json.dumps(["https://example.com/spend", "https://example.com/pop"])
        row = _make_derived_row(fid,
                                "total_amount_usd / population",
                                inputs, "126.580000", "USD per capita")
        _write_parquet(site / "citations" / "citations.parquet", _CIT_COL_DEFS, [row])
        _write_manifest(site)

        result = citation_gate5b1(site)
        assert result["ok"] is True, f"URL-input derived should pass: {result.get('failures')}"


class TestDerivedRecompute:
    """_verify_derived recomputes trajectory difference and sum."""

    def test_trajectory_difference_recompute_pass(self, tmp_path):
        """fy2526_change = fy2026 - fy2025: correct recompute → PASS.

        Peer rows use pivot formula when no budget_lines inputs available.
        The fy2526_change row uses the two peer derived fact_ids as inputs.
        """
        site = tmp_path / "site"
        import json

        fid_fy25 = fact_id_derived("trajectory", "0601101E|DARPA", "fy2025_total")
        fid_fy26 = fact_id_derived("trajectory", "0601101E|DARPA", "fy2026_total")
        fid_chg  = fact_id_derived("trajectory", "0601101E|DARPA", "fy2526_change")

        rows = [
            # Peer rows use pivot formula (no budget_lines inputs available)
            _make_derived_row(fid_fy25,
                              "trajectory pivot of budget_lines (inputs unavailable for this org/type)",
                              "[]", "293145.000"),
            _make_derived_row(fid_fy26,
                              "trajectory pivot of budget_lines (inputs unavailable for this org/type)",
                              "[]", "295000.000"),
            # fy2526_change = fy2026 - fy2025; inputs = [fid_fy26, fid_fy25] (peer derived fids)
            _make_derived_row(fid_chg, "fy2026_total - fy2025_total",
                              json.dumps([fid_fy26, fid_fy25]), "1855.000"),
        ]
        _write_parquet(site / "citations" / "citations.parquet", _CIT_COL_DEFS, rows)
        _write_manifest(site)

        result = citation_gate5b1(site)
        assert result["ok"] is True, f"difference recompute should pass: {result.get('failures')}"

    def test_trajectory_difference_recompute_wrong_value_fails(self, tmp_path):
        """fy2526_change with wrong recorded_value → FAIL."""
        site = tmp_path / "site"
        import json

        fid_fy25 = fact_id_derived("trajectory", "0601101E|DARPA", "fy2025_total")
        fid_fy26 = fact_id_derived("trajectory", "0601101E|DARPA", "fy2026_total")
        fid_chg  = fact_id_derived("trajectory", "0601101E|DARPA", "fy2526_change")

        rows = [
            _make_derived_row(fid_fy25, "sum(budget_lines...)", "[]", "293145.000"),
            _make_derived_row(fid_fy26, "sum(budget_lines...)", "[]", "295000.000"),
            # Wrong recorded_value: should be 1855 but says 9999
            _make_derived_row(fid_chg, "fy2026_total - fy2025_total",
                              json.dumps([fid_fy26, fid_fy25]), "9999.000"),
        ]
        _write_parquet(site / "citations" / "citations.parquet", _CIT_COL_DEFS, rows)
        _write_manifest(site)

        result = citation_gate5b1(site)
        assert result["ok"] is False, "wrong difference should fail"
        assert any(fid_chg in str(f) for f in result["failures"])

    def test_derived_inputs_missing_fact_id_fails(self, tmp_path):
        """Derived row whose inputs reference a non-existent fact_id → FAIL."""
        site = tmp_path / "site"
        import json

        fid_chg = fact_id_derived("trajectory", "0601101E|DARPA", "fy2526_change")
        phantom_id = "deadbeef12345678"  # not in citations

        row = _make_derived_row(fid_chg, "fy2026_total - fy2025_total",
                                json.dumps([phantom_id, phantom_id]), "0.000")
        _write_parquet(site / "citations" / "citations.parquet", _CIT_COL_DEFS, [row])
        _write_manifest(site)

        result = citation_gate5b1(site)
        assert result["ok"] is False, "phantom input fact_id should fail"


class TestDerivedIntegrityGate:
    """integrity_gate5b1 checks derived formula+value."""

    def test_derived_formula_and_value_check_passes(self, tmp_path):
        """Derived row with formula+recorded_value → integrity PASS."""
        site = tmp_path / "site"
        fid = fact_id_derived("state_per_capita", "VA|Defense", "amount_per_capita")
        row = _make_derived_row(fid, "total_amount_usd / population",
                                "[]", "55.123")
        _write_parquet(site / "citations" / "citations.parquet", _CIT_COL_DEFS, [row])
        _write_manifest(site)

        result = integrity_gate5b1(site)
        assert result["checks"].get("derived_formula_and_value") is True, (
            f"expected pass for valid derived row: {result}"
        )

    def test_derived_missing_formula_fails_integrity(self, tmp_path):
        """Derived row with null formula → integrity FAIL."""
        site = tmp_path / "site"
        fid = fact_id_derived("entity", "some-family", "total_obligation")
        # Row with formula=None
        row = (fid, "derived", "USD", None,
               None, None, None, None, None, None, None, None,
               None, None, None, None, None, None, None, "2026-06-12",
               None, "[]", None, "5000000.000",
               None, None, None)  # formula=None; pe_bli, scenario, amount_type
        _write_parquet(site / "citations" / "citations.parquet", _CIT_COL_DEFS, [row])
        _write_manifest(site)

        result = integrity_gate5b1(site)
        assert result["checks"].get("derived_formula_and_value") is False, (
            f"expected fail for null formula: {result}"
        )


# ---------------------------------------------------------------------------
# Finding 1: derived sum recompute via budget_lines.parquet
# Tests call _verify_derived directly to avoid sampling workbook-kind rows
# that would require real xlsx files.
# ---------------------------------------------------------------------------


def _make_derived_row_tuple(fid: str, formula: str, inputs_json: str,
                            recorded_value: str) -> tuple:
    """27-element derived citation row (same layout as _CIT_COL_DEFS)."""
    return (fid, "derived", "USD thousands", None,
            None, None, None, None, None, None, None, None,
            None, None, None, None, None, None, None, "2026-06-12T00:00:00+00:00",
            formula, inputs_json, None, recorded_value,
            None, None, None)  # pe_bli, scenario, amount_type


def _build_cit_idx() -> dict:
    """Return col_idx matching _CIT_COL_DEFS order."""
    cols = [
        "fact_id", "kind", "units", "amount_text",
        "page_number", "x0", "x1", "top_pt", "bottom_pt",
        "page_width", "page_height", "resolution",
        "sheet", "cells", "amount_thousands",
        "sha256", "hosted_pdf_url", "official_url",
        "xml_path", "retrieved_at",
        "formula", "inputs", "query_body", "recorded_value",
        "pe_bli", "scenario", "amount_type",
    ]
    return {name: i for i, name in enumerate(cols)}


class TestDerivedSumViabudgetLines:
    """Rule 4c: sum(budget_lines) resolved via budget_lines.parquet when
    workbook inputs carry recorded_value=None.

    Tests call _verify_derived directly to avoid the workbook-kind
    xlsx requirement that would be triggered by citation_gate5b1 sampling.
    """

    def _make_all_cits(self, *rows: tuple) -> list[tuple]:
        return list(rows)

    def test_sum_via_budget_lines_correct_passes(self):
        """Derived sum with inputs whose recorded_value=None but
        budget_lines.parquet (fid_to_bl_amount) carries amount_thousands → PASS."""
        import json as _json
        fid_a = "aaaa000011110000"
        fid_b = "bbbb000022220000"
        fid_sum = "cccc000033330000"

        # Simulated citations: workbook rows with recorded_value=None
        cit_a = _make_derived_row_tuple(fid_a, "workbook-formula", "[]", None)
        # Override kind to workbook and recorded_value to None (already None above)
        cit_a_as_wb = list(cit_a)
        cit_a_as_wb[1] = "workbook"  # kind
        cit_a_as_wb[23] = None       # recorded_value
        cit_b_as_wb = list(_make_derived_row_tuple(fid_b, "workbook-formula", "[]", None))
        cit_b_as_wb[1] = "workbook"
        cit_b_as_wb[23] = None

        idx = _build_cit_idx()
        all_cits = [tuple(cit_a_as_wb), tuple(cit_b_as_wb)]

        # budget_lines lookup carries the actual amounts
        fid_to_bl = {fid_a: "3000.0", fid_b: "629.0"}

        # Derived row summing the two
        derived_row = _make_derived_row_tuple(
            fid_sum,
            "sum(budget_lines.amount_thousands where amount_type in (fy_2025_total, fy_2025_enacted))",
            _json.dumps([fid_a, fid_b]),
            "3629.0",
        )

        result = _verify_derived(derived_row, idx, all_cits, idx, fid_to_bl)
        assert result is None, f"correct sum via budget_lines should PASS: {result}"

    def test_sum_via_budget_lines_tampered_recorded_value_fails(self):
        """Derived sum with wrong recorded_value → FAIL 'mismatch'."""
        import json as _json
        fid_a = "aaaa000011110000"
        fid_b = "bbbb000022220000"
        fid_sum = "cccc000033330000"

        idx = _build_cit_idx()
        cit_a = list(_make_derived_row_tuple(fid_a, "wb", "[]", None))
        cit_a[1] = "workbook"; cit_a[23] = None
        cit_b = list(_make_derived_row_tuple(fid_b, "wb", "[]", None))
        cit_b[1] = "workbook"; cit_b[23] = None
        all_cits = [tuple(cit_a), tuple(cit_b)]

        fid_to_bl = {fid_a: "3000.0", fid_b: "629.0"}

        derived_row = _make_derived_row_tuple(
            fid_sum,
            "sum(budget_lines.amount_thousands where amount_type=fy_2025_total)",
            _json.dumps([fid_a, fid_b]),
            "9999.0",  # WRONG — should be 3629
        )
        result = _verify_derived(derived_row, idx, all_cits, idx, fid_to_bl)
        assert result is not None, "tampered recorded_value should FAIL"
        assert "mismatch" in result, f"expected 'mismatch' in: {result}"

    def test_sum_unresolvable_inputs_fails(self):
        """Inputs absent from both citations.recorded_value and
        budget_lines.amount_thousands → FAIL 'unresolvable'."""
        import json as _json
        fid_a = "aaaa000011110000"
        fid_b = "bbbb000022220000"
        fid_sum = "cccc000033330000"

        idx = _build_cit_idx()
        # Citations present but recorded_value=None
        cit_a = list(_make_derived_row_tuple(fid_a, "wb", "[]", None))
        cit_a[1] = "workbook"; cit_a[23] = None
        cit_b = list(_make_derived_row_tuple(fid_b, "wb", "[]", None))
        cit_b[1] = "workbook"; cit_b[23] = None
        all_cits = [tuple(cit_a), tuple(cit_b)]

        # Empty budget_lines — truly unresolvable
        fid_to_bl: dict = {}

        derived_row = _make_derived_row_tuple(
            fid_sum,
            "sum(budget_lines.amount_thousands where amount_type=fy_2025_total)",
            _json.dumps([fid_a, fid_b]),
            "1000.0",
        )
        result = _verify_derived(derived_row, idx, all_cits, idx, fid_to_bl)
        assert result is not None, "unresolvable inputs should FAIL"
        assert "unresolvable" in result, f"expected 'unresolvable' in: {result}"


class TestFlowChildrenRecompute:
    """Rule 4b0 (Phase 5H): sum(flow_children…) node facts recompute from
    their edge facts' recorded_values. Called directly like the 4c tests."""

    def _edge_cits(self, pairs):
        rows = []
        for fid, rv in pairs:
            rows.append(_make_derived_row_tuple(
                fid, "sum(contracts.federal_action_obligation where fy=2025"
                     " and flow_edge=a->b)", "[]", rv))
        return rows

    def test_flow_children_correct_sum_passes(self):
        import json as _json
        e1, e2 = "aaaa000011110000", "bbbb000022220000"
        all_cits = self._edge_cits([(e1, "600.25"), (e2, "150.50")])
        node_row = _make_derived_row_tuple(
            "cccc000033330000",
            "sum(flow_children(out_edges) where fy=2025 and flow_node=s:2025:sub:DEPT ARMY)",
            _json.dumps([e1, e2]), "750.750",
        )
        idx = _build_cit_idx()
        assert _verify_derived(node_row, idx, all_cits, idx, {}) is None

    def test_flow_children_tampered_recorded_value_fails(self):
        import json as _json
        e1, e2 = "aaaa000011110000", "bbbb000022220000"
        all_cits = self._edge_cits([(e1, "600.25"), (e2, "150.50")])
        node_row = _make_derived_row_tuple(
            "cccc000033330000",
            "sum(flow_children(out_edges) where fy=2025 and flow_node=s:2025:sub:DEPT ARMY)",
            _json.dumps([e1, e2]), "999.999",  # WRONG
        )
        idx = _build_cit_idx()
        result = _verify_derived(node_row, idx, all_cits, idx, {})
        assert result is not None and "mismatch" in result

    def test_flow_children_unresolvable_input_fails(self):
        """Edge fact present but with recorded_value=None → FAIL, never skip."""
        import json as _json
        e1 = "aaaa000011110000"
        cit = list(_make_derived_row_tuple(e1, "edge", "[]", None))
        node_row = _make_derived_row_tuple(
            "cccc000033330000",
            "sum(flow_children(out_edges) where fy=2025 and flow_node=n)",
            _json.dumps([e1]), "1.000",
        )
        idx = _build_cit_idx()
        result = _verify_derived(node_row, idx, [tuple(cit)], idx, {})
        assert result is not None and "unresolvable" in result

    def test_flow_children_name_with_dash_not_misread_as_difference(self):
        """A node whose embedded name contains ' - ' with exactly 2 inputs
        must recompute as a SUM (rule 4b0), not a difference (rule 4b)."""
        import json as _json
        e1, e2 = "aaaa000011110000", "bbbb000022220000"
        all_cits = self._edge_cits([(e1, "100.00"), (e2, "50.00")])
        node_row = _make_derived_row_tuple(
            "cccc000033330000",
            "sum(flow_children(out_edges) where fy=2025 and"
            " flow_node=s:2025:sub:DEPT OF DEFENSE - EDUCATION ACTIVITY)",
            _json.dumps([e1, e2]), "150.000",  # sum, NOT 100-50=50
        )
        idx = _build_cit_idx()
        assert _verify_derived(node_row, idx, all_cits, idx, {}) is None

    def test_flow_children_empty_inputs_fails(self):
        node_row = _make_derived_row_tuple(
            "cccc000033330000",
            "sum(flow_children(out_edges) where fy=2025 and flow_node=n)",
            "[]", "0.000",
        )
        idx = _build_cit_idx()
        result = _verify_derived(node_row, idx, [], idx, {})
        assert result is not None and "not recomputable" in result


# ---------------------------------------------------------------------------
# jbook_narrative citation kind
# ---------------------------------------------------------------------------

_NARR_SHA = "a" * 64  # fake valid sha256 hex
_NARR_PE = "0204WARSYS"
_NARR_KIND = "overview"
_NARR_XMLPATH = "ProgramElement[0]/Narrative[0]"
_NARR_URL = "https://example.mil/fy2026_justification.pdf"


def _make_narrative_cit_row(
    fid: str | None = None,
    *,
    sha: str = _NARR_SHA,
    xml_path: str = _NARR_XMLPATH,
    official_url: str = _NARR_URL,
) -> tuple:
    """Build a 27-element citation tuple for kind='jbook_narrative'."""
    fid = fid or fact_id_narrative(sha, _NARR_PE, _NARR_KIND, xml_path)
    return (
        fid, "jbook_narrative", None,  # fact_id, kind, units
        None, None, None, None, None, None, None, None,  # amount_text + bbox
        None,   # resolution
        None,   # sheet
        None,   # cells
        None,   # amount_thousands
        sha,    # sha256
        None,   # hosted_pdf_url
        official_url,
        xml_path,
        None,   # retrieved_at
        None, None, None, None,  # formula, inputs, query_body, recorded_value
        _NARR_PE, None, None,   # pe_bli, scenario, amount_type
    )


def _build_narr_cit_idx() -> dict:
    """Build column-index dict matching _CIT_COL_DEFS order."""
    cols = [
        "fact_id", "kind", "units", "amount_text", "page_number",
        "x0", "x1", "top_pt", "bottom_pt", "page_width", "page_height",
        "resolution", "sheet", "cells", "amount_thousands",
        "sha256", "hosted_pdf_url", "official_url", "xml_path", "retrieved_at",
        "formula", "inputs", "query_body", "recorded_value",
        "pe_bli", "scenario", "amount_type",
    ]
    return {name: i for i, name in enumerate(cols)}


class TestFactIdNarrative:
    def test_stable_and_16hex(self):
        fid = fact_id_narrative(_NARR_SHA, _NARR_PE, _NARR_KIND, _NARR_XMLPATH)
        assert len(fid) == 16
        assert all(c in "0123456789abcdef" for c in fid)
        assert fid == fact_id_narrative(_NARR_SHA, _NARR_PE, _NARR_KIND, _NARR_XMLPATH)

    def test_differs_by_sha(self):
        a = fact_id_narrative("a" * 64, _NARR_PE, _NARR_KIND, _NARR_XMLPATH)
        b = fact_id_narrative("b" * 64, _NARR_PE, _NARR_KIND, _NARR_XMLPATH)
        assert a != b

    def test_differs_by_pe_bli(self):
        a = fact_id_narrative(_NARR_SHA, "0204WARSYS", _NARR_KIND, _NARR_XMLPATH)
        b = fact_id_narrative(_NARR_SHA, "0604384BP", _NARR_KIND, _NARR_XMLPATH)
        assert a != b

    def test_differs_by_kind(self):
        a = fact_id_narrative(_NARR_SHA, _NARR_PE, "overview", _NARR_XMLPATH)
        b = fact_id_narrative(_NARR_SHA, _NARR_PE, "accomplishments", _NARR_XMLPATH)
        assert a != b

    def test_differs_by_xml_path(self):
        a = fact_id_narrative(_NARR_SHA, _NARR_PE, _NARR_KIND, "ProgramElement[0]/Narrative[0]")
        b = fact_id_narrative(_NARR_SHA, _NARR_PE, _NARR_KIND, "ProgramElement[0]/Narrative[1]")
        assert a != b


class TestVerifyJbookNarrative:
    """Unit tests for _verify_jbook_narrative shape rules."""

    def _call(self, row, idx=None):
        from govbudget.verify_phase5b1 import _verify_jbook_narrative
        return _verify_jbook_narrative(row, idx or _build_narr_cit_idx())

    def test_valid_passes(self):
        row = _make_narrative_cit_row()
        assert self._call(row) is None

    def test_missing_sha_fails(self):
        row = list(_make_narrative_cit_row())
        row[_build_narr_cit_idx()["sha256"]] = None
        assert self._call(tuple(row)) is not None

    def test_empty_xml_path_fails(self):
        """Empty xml_path must FAIL — narrative is not locatable."""
        row = list(_make_narrative_cit_row())
        row[_build_narr_cit_idx()["xml_path"]] = ""
        result = self._call(tuple(row))
        assert result is not None
        assert "xml_path" in result

    def test_null_xml_path_fails(self):
        row = list(_make_narrative_cit_row())
        row[_build_narr_cit_idx()["xml_path"]] = None
        assert self._call(tuple(row)) is not None

    def test_xml_path_bad_prefix_fails(self):
        """xml_path that doesn't start with ProgramElement[ or LineItem[ → FAIL."""
        row = _make_narrative_cit_row(xml_path="SomeOtherElement[0]/Narrative[0]")
        result = self._call(row)
        assert result is not None
        assert "ProgramElement[" in result or "LineItem[" in result

    def test_xml_path_program_element_passes(self):
        row = _make_narrative_cit_row(xml_path="ProgramElement[5]/Narrative[2]")
        assert self._call(row) is None

    def test_xml_path_line_item_passes(self):
        row = _make_narrative_cit_row(xml_path="LineItem[0]/Narrative[1]")
        assert self._call(row) is None

    def test_missing_official_url_fails(self):
        row = _make_narrative_cit_row(official_url="")
        result = self._call(row)
        assert result is not None
        assert "official_url" in result


class TestJbookNarrativeCitationGate:
    """citation_gate5b1 treats jbook_narrative as a verifiable kind."""

    def _write_narr_site(self, site: Path, rows: list | None = None) -> None:
        if rows is None:
            rows = [_make_narrative_cit_row()]
        (site / "citations").mkdir(parents=True, exist_ok=True)
        _write_parquet(site / "citations" / "citations.parquet", _CIT_COL_DEFS, rows)

    def test_valid_narrative_citation_passes_gate(self, tmp_path):
        site = tmp_path / "site"
        self._write_narr_site(site)
        result = citation_gate5b1(site)
        assert result["ok"] is True, f"failures: {result['failures']}"
        assert result["sampled"] == 1

    def test_empty_xml_path_fails_gate(self, tmp_path):
        """Narrative citation with empty xml_path → gate FAIL (not locatable)."""
        site = tmp_path / "site"
        row = list(_make_narrative_cit_row())
        row[_build_narr_cit_idx()["xml_path"]] = ""
        self._write_narr_site(site, [tuple(row)])
        result = citation_gate5b1(site)
        assert result["ok"] is False
        assert len(result["failures"]) >= 1

    def test_bad_xml_path_prefix_fails_gate(self, tmp_path):
        """Narrative with xml_path not starting ProgramElement[ or LineItem[ → FAIL."""
        site = tmp_path / "site"
        row = _make_narrative_cit_row(xml_path="BadPrefix[0]/Narrative[0]")
        self._write_narr_site(site, [row])
        result = citation_gate5b1(site)
        assert result["ok"] is False


class TestJbookNarrativeIntegrityCheck:
    """integrity_gate5b1 jbook_narrative_shape check."""

    def _write_narr_site_with_parquet(
        self,
        site: Path,
        cit_rows: list | None = None,
        narr_parquet_rows: list | None = None,
    ) -> None:
        """Build a minimal site with jbook_narrative citations and jbook_narratives.parquet."""
        fid = fact_id_narrative(_NARR_SHA, _NARR_PE, _NARR_KIND, _NARR_XMLPATH)
        if cit_rows is None:
            cit_rows = [_make_narrative_cit_row(fid)]
        if narr_parquet_rows is None:
            narr_parquet_rows = [
                (fid, _NARR_PE, None, _NARR_KIND, "Narrative title",
                 "Narrative body text.", _NARR_XMLPATH, "ARMY", 2026, _NARR_SHA),
            ]

        (site / "citations").mkdir(parents=True, exist_ok=True)
        _write_parquet(site / "citations" / "citations.parquet", _CIT_COL_DEFS, cit_rows)

        (site / "data").mkdir(parents=True, exist_ok=True)
        _write_parquet(
            site / "data" / "jbook_narratives.parquet",
            "fact_id varchar, pe_bli varchar, project_number varchar,"
            " kind varchar, title varchar, body varchar, xml_path varchar,"
            " org varchar, fiscal_year integer, document_sha256 varchar",
            narr_parquet_rows,
        )

    def test_valid_passes_integrity(self, tmp_path):
        site = tmp_path / "site"
        self._write_narr_site_with_parquet(site)
        result = integrity_gate5b1(site)
        assert result["checks"].get("jbook_narrative_shape") is True, result

    def test_orphan_narrative_citation_fails(self, tmp_path):
        """Citation fact_id not in jbook_narratives.parquet → integrity FAIL."""
        site = tmp_path / "site"
        orphan_fid = "dead000000000000"
        cit_rows = [_make_narrative_cit_row(orphan_fid)]
        # jbook_narratives.parquet has a DIFFERENT fact_id
        real_fid = fact_id_narrative(_NARR_SHA, _NARR_PE, _NARR_KIND, _NARR_XMLPATH)
        narr_rows = [
            (real_fid, _NARR_PE, None, _NARR_KIND, "Title", "Body",
             _NARR_XMLPATH, "ARMY", 2026, _NARR_SHA),
        ]
        self._write_narr_site_with_parquet(site, cit_rows, narr_rows)
        result = integrity_gate5b1(site)
        assert result["checks"].get("jbook_narrative_shape") is False
        assert any("orphan" in f for f in result["failures"])

    def test_missing_official_url_fails_integrity(self, tmp_path):
        """jbook_narrative citation with null official_url → integrity FAIL."""
        site = tmp_path / "site"
        fid = fact_id_narrative(_NARR_SHA, _NARR_PE, _NARR_KIND, _NARR_XMLPATH)
        row = list(_make_narrative_cit_row(fid))
        row[_build_narr_cit_idx()["official_url"]] = None
        cit_rows = [tuple(row)]
        narr_rows = [
            (fid, _NARR_PE, None, _NARR_KIND, "Title", "Body",
             _NARR_XMLPATH, "ARMY", 2026, _NARR_SHA),
        ]
        self._write_narr_site_with_parquet(site, cit_rows, narr_rows)
        result = integrity_gate5b1(site)
        assert result["checks"].get("jbook_narrative_shape") is False


# ---------------------------------------------------------------------------
# Gate 4: narrative_gate5b1 — narrative provenance re-derivation (Phase 5F §2b)
# ---------------------------------------------------------------------------

from pdf_factory import make_pdf  # noqa: E402

from govbudget.verify_phase5b1 import narrative_gate5b1  # noqa: E402

_NG_PE = "0605502TST"
_NG_XMLPATH = "ProgramElement[0]/Narrative[0]"
_NG_BODY = (
    "The Corrosion Prevention program develops coatings that reduce maintenance"
    " costs across weapon systems. Additional trailing sentences follow the"
    " distinctive opening."
)


def _ng_pdf(path: Path) -> None:
    """Two pages; the narrative opening lives ONLY on page 2."""
    make_pdf(path, [
        ["Exhibit R-1, RDT&E Program", f"{_NG_PE} 12.345"],
        ["Exhibit R-2, RDT&E Budget Item Justification",
         f"PE {_NG_PE} / CORROSION PREVENTION",
         "The Corrosion Prevention program develops coatings that reduce",
         "maintenance costs across weapon systems. Additional trailing."],
    ])


def _make_site_with_narrative_provenance(
    site_dir: Path,
    *,
    paged: bool = True,
    page_number: int | None = None,
    body_in_parquet: str = _NG_BODY,
) -> str:
    """Site dir with one jbook_narrative citation + jbook_narratives.parquet.

    paged=True stores the live find_narrative_page location (page 2);
    page_number overrides it (for planted-defect tests). Returns fact_id.
    """
    from govbudget.jbooks.provenance_pages import find_narrative_page

    pdf_src = site_dir / "_src_narr.pdf"
    site_dir.mkdir(parents=True, exist_ok=True)
    _ng_pdf(pdf_src)
    sha = hashlib.sha256(pdf_src.read_bytes()).hexdigest()
    pdfs_dir = site_dir / "pdfs"
    pdfs_dir.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(pdf_src, pdfs_dir / f"{sha}.pdf")

    fid = fact_id_narrative(sha, _NG_PE, "mission", _NG_XMLPATH)

    if paged:
        hit = find_narrative_page(pdf_src, pe_bli=_NG_PE, body=_NG_BODY,
                                  exhibit_family="rdte")
        assert hit["resolution"] == "unique" and hit["page_number"] == 2
        pn = page_number if page_number is not None else hit["page_number"]
        cit_row = (
            fid, "jbook_narrative", None, None,
            pn, hit["x0"], hit["x1"], hit["top_pt"], hit["bottom_pt"],
            hit["page_width"], hit["page_height"], hit["resolution"],
            None, None, None, sha,
            f"https://cdn.example/pdfs/{sha}.pdf#page={pn}",
            f"https://example.mil/narr.pdf#page={pn}",
            _NG_XMLPATH, None,
            None, None, None, None,
            _NG_PE, None, None,
        )
    else:
        cit_row = (
            fid, "jbook_narrative", None, None,
            None, None, None, None, None, None, None, None,
            None, None, None, sha,
            None, "https://example.mil/narr.pdf",
            _NG_XMLPATH, None,
            None, None, None, None,
            _NG_PE, None, None,
        )

    _write_parquet(site_dir / "citations" / "citations.parquet",
                   _CIT_COL_DEFS, [cit_row])
    _write_parquet(
        site_dir / "data" / "jbook_narratives.parquet",
        "fact_id varchar, pe_bli varchar, project_number varchar,"
        " kind varchar, title varchar, body varchar, xml_path varchar,"
        " org varchar, fiscal_year integer, document_sha256 varchar",
        [(fid, _NG_PE, None, "mission", "A. Mission Description",
          body_in_parquet, _NG_XMLPATH, "OSD", 2026, sha)],
    )
    return fid


class TestNarrativeGate5b1:
    def test_pre_5f_pageless_state_fails(self, tmp_path):
        """THE pre-failure the leg exists to catch: narrative citations exist
        but none carries page provenance → FAIL."""
        site = tmp_path / "site"
        _make_site_with_narrative_provenance(site, paged=False)
        result = narrative_gate5b1(site)
        assert result["ok"] is False
        assert result["resolved_total"] == 0
        assert result["narrative_total"] == 1
        assert "no narrative citation carries page provenance" in result["reason"]

    def test_resolved_narrative_rederives(self, tmp_path):
        site = tmp_path / "site"
        _make_site_with_narrative_provenance(site)
        result = narrative_gate5b1(site)
        assert result["failures"] == []
        assert result["ok"] is True
        assert result["sampled"] == 1
        assert result["passed"] == 1

    def test_wrong_page_fails(self, tmp_path):
        """Planted defect: page_number points at a page that does not contain
        the recomputed opening text."""
        site = tmp_path / "site"
        fid = _make_site_with_narrative_provenance(site, page_number=1)
        result = narrative_gate5b1(site)
        assert result["ok"] is False
        assert any(f[0] == fid and "opening text not found" in f[1]
                   for f in result["failures"])

    def test_tampered_body_fails(self, tmp_path):
        """The opening is recomputed from the exported body — a body that
        drifted from the located passage must fail, not silently pass."""
        site = tmp_path / "site"
        fid = _make_site_with_narrative_provenance(
            site,
            body_in_parquet="Completely different body text that does not"
                            " appear on the cited page at all today",
        )
        result = narrative_gate5b1(site)
        assert result["ok"] is False
        assert any(f[0] == fid for f in result["failures"])

    def test_missing_body_row_fails(self, tmp_path):
        """A paged citation whose fact_id has no jbook_narratives row cannot
        be re-derived → FAIL (never skip)."""
        site = tmp_path / "site"
        fid = _make_site_with_narrative_provenance(site)
        # rewrite the narratives parquet without the row
        _write_parquet(
            site / "data" / "jbook_narratives.parquet",
            "fact_id varchar, pe_bli varchar, project_number varchar,"
            " kind varchar, title varchar, body varchar, xml_path varchar,"
            " org varchar, fiscal_year integer, document_sha256 varchar",
            [],
        )
        result = narrative_gate5b1(site)
        assert result["ok"] is False
        assert any(f[0] == fid and "no jbook_narratives row" in f[1]
                   for f in result["failures"])

    def test_sha_mismatch_fails(self, tmp_path):
        site = tmp_path / "site"
        fid = _make_site_with_narrative_provenance(site)
        # corrupt the sha-named PDF
        cit_pq = site / "citations" / "citations.parquet"
        sha = duckdb.sql(
            f"select sha256 from read_parquet('{cit_pq}')").fetchone()[0]
        (site / "pdfs" / f"{sha}.pdf").write_bytes(b"%PDF-1.4 tampered")
        result = narrative_gate5b1(site)
        assert result["ok"] is False
        assert any(f[0] == fid and "sha256 mismatch" in f[1]
                   for f in result["failures"])

    def test_missing_site_dir(self, tmp_path):
        result = narrative_gate5b1(tmp_path / "nope")
        assert result["ok"] is False
        assert "site_dir missing" in result["reason"]


# ---------------------------------------------------------------------------
# Phase 5E Task 6: decade gate-scope extension
# (documented in docs/superpowers/reviews/5c-gates-pre-failure.txt —
#  the workbook fact universe and the derived-recompute value lookup learn
#  about data/budget_lines_decade.parquet; the difference rule gains the
#  same budget_lines fallback the sum rule already had, so book-diff facts
#  are GENUINELY recomputed, never silently shape-checked)
# ---------------------------------------------------------------------------

_BL_COL_DEFS = (
    "fact_id varchar, exhibit varchar, fiscal_year integer, account varchar,"
    " account_title varchar, organization varchar, budget_activity varchar,"
    " budget_activity_title varchar, pe_bli varchar, title varchar,"
    " amount_type varchar, amount_thousands double, units varchar,"
    " document_sha256 varchar, source_sheet varchar, source_cells varchar"
)


def _decade_bl_parquet_row(fid: str, edition: int, at: str, amount: float,
                           sha: str | None = None, cells: str = "J4") -> tuple:
    return (fid, "R-1", edition, "0400", "RDT&E Defense-Wide", "DARPA", "01",
            "Basic Research", "0601101E", "OLD LINE", at, amount,
            "USD thousands", sha or f"sha_pb{edition}", "Exhibit R-1", cells)


def _decade_wb_citation_row(fid: str, edition: int, at: str, amount: float,
                            sha: str | None = None, cells: str = "J4") -> tuple:
    return (fid, "workbook", "USD thousands", None, None,
            None, None, None, None, None, None, None,
            "Exhibit R-1", cells, amount,
            sha or f"sha_pb{edition}", None,
            "https://example.mil/old_r1.xlsx", None, None,
            None, None, None, None,
            "0601101E", None, at)


def _mk_decade_workbook(site: Path, name: str, cells: dict[str, float]) -> str:
    """Write an xlsx with the given 'Exhibit R-1' cell values, copy it to
    workbooks/{sha}.xlsx and return the real sha (citation_gate5b1's
    workbook tier re-opens the file and sums the cited cells)."""
    src = site / f"_src_{name}.xlsx"
    src.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Exhibit R-1"
    for ref, v in cells.items():
        ws[ref] = v
    wb.save(src)
    sha = hashlib.sha256(src.read_bytes()).hexdigest()
    wb_dir = site / "workbooks"
    wb_dir.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(src, wb_dir / f"{sha}.xlsx")
    return sha


class TestDecadeWorkbookSetEquality:
    """integrity_gate5b1: the workbook fact universe is
    budget_lines.parquet ∪ budget_lines_decade.parquet."""

    def _base_site(self, site: Path) -> str:
        """PB2026 workbook site (real xlsx) — returns the 2026 fid."""
        _, fid = _make_site_with_workbook(site)
        return fid

    def test_decade_rows_cited_passes(self, tmp_path):
        site = tmp_path / "site"
        fid26 = self._base_site(site)
        dfid = fact_id_workbook(
            "sha_pb2022", "R-1", 2022, "0400", "DARPA", "01",
            "0601101E", "fy_2020_actuals")
        _write_parquet(site / "data" / "budget_lines_decade.parquet",
                       _BL_COL_DEFS,
                       [_decade_bl_parquet_row(dfid, 2022, "fy_2020_actuals", 150000.0)])
        # citations: the 2026 workbook row + the decade workbook row
        _write_parquet(
            site / "citations" / "citations.parquet", _CIT_COL_DEFS,
            [(fid26, "workbook", "USD thousands", None, None,
              None, None, None, None, None, None, None,
              "Exhibit R-1", "J2", 280494.0,
              "sha_wb26", None, "https://example.mil/r1.xlsx", None, None,
              None, None, None, None, None, None, None),
             _decade_wb_citation_row(dfid, 2022, "fy_2020_actuals", 150000.0)],
        )
        _write_manifest(site)
        result = integrity_gate5b1(site)
        assert result["checks"]["workbook_set_equality"] is True, result["failures"]

    def test_uncited_decade_row_fails(self, tmp_path):
        """A budget_lines_decade row with NO workbook citation → FAIL
        (set equality is bidirectional over the union universe)."""
        site = tmp_path / "site"
        fid26 = self._base_site(site)
        dfid = fact_id_workbook(
            "sha_pb2022", "R-1", 2022, "0400", "DARPA", "01",
            "0601101E", "fy_2020_actuals")
        _write_parquet(site / "data" / "budget_lines_decade.parquet",
                       _BL_COL_DEFS,
                       [_decade_bl_parquet_row(dfid, 2022, "fy_2020_actuals", 150000.0)])
        _write_parquet(
            site / "citations" / "citations.parquet", _CIT_COL_DEFS,
            [(fid26, "workbook", "USD thousands", None, None,
              None, None, None, None, None, None, None,
              "Exhibit R-1", "J2", 280494.0,
              "sha_wb26", None, "https://example.mil/r1.xlsx", None, None,
              None, None, None, None, None, None, None)],
        )
        _write_manifest(site)
        result = integrity_gate5b1(site)
        assert result["checks"]["workbook_set_equality"] is False
        assert any("not cited" in f for f in result["failures"])

    def test_orphan_decade_citation_fails(self, tmp_path):
        """A decade workbook citation absent from BOTH parquets → FAIL."""
        site = tmp_path / "site"
        fid26 = self._base_site(site)
        _write_parquet(
            site / "citations" / "citations.parquet", _CIT_COL_DEFS,
            [(fid26, "workbook", "USD thousands", None, None,
              None, None, None, None, None, None, None,
              "Exhibit R-1", "J2", 280494.0,
              "sha_wb26", None, "https://example.mil/r1.xlsx", None, None,
              None, None, None, None, None, None, None),
             _decade_wb_citation_row("ffff000000000001", 2022,
                                     "fy_2020_actuals", 1.0)],
        )
        _write_manifest(site)
        result = integrity_gate5b1(site)
        assert result["checks"]["workbook_set_equality"] is False
        assert any("not in budget_lines" in f for f in result["failures"])


class TestDecadeDerivedRecompute:
    """citation_gate5b1 recomputes decade sums and book-diff differences
    through budget_lines_decade.parquet amounts."""

    def _site_with_diff(self, site: Path, recorded_delta: str,
                        recorded_sum: str = "90000.000") -> tuple[str, str]:
        """Decade fixture (REAL xlsx files — the workbook tier re-derives):
        a 2-row multi-source sum grain (PB2017) and a book-diff fact
        PB2020 request → PB2022 actuals. Returns (sum_fid, diff_fid)."""
        sha17 = _mk_decade_workbook(site, "pb2017",
                                    {"J4": 60000.0, "J5": 30000.0})
        sha20 = _mk_decade_workbook(site, "pb2020", {"J4": 140000.0})
        sha22 = _mk_decade_workbook(site, "pb2022", {"J4": 150000.0})
        w1 = fact_id_workbook(sha17, "R-1", 2017, "0400", "DARPA",
                              "01", "0601101E", "fy_2015_actuals")
        w2 = fact_id_workbook(sha17, "R-1", 2017, "0400", "DARPA",
                              "02", "0601101E", "fy_2015_actuals")
        w_from = fact_id_workbook(sha20, "R-1", 2020, "0400", "DARPA",
                                  "01", "0601101E", "fy_2020_total")
        w_to = fact_id_workbook(sha22, "R-1", 2022, "0400", "DARPA",
                                "01", "0601101E", "fy_2020_actuals")
        sum_fid = fact_id_derived("decade", "0601101E|2017", "fy_2015_actuals")
        diff_fid = fact_id_derived("book_diff", "0601101E|2020|2022",
                                   "request_vs_actuals")
        rows_bl = [
            _decade_bl_parquet_row(w1, 2017, "fy_2015_actuals", 60000.0,
                                   sha=sha17, cells="J4"),
            _decade_bl_parquet_row(w2, 2017, "fy_2015_actuals", 30000.0,
                                   sha=sha17, cells="J5"),
            _decade_bl_parquet_row(w_from, 2020, "fy_2020_total", 140000.0,
                                   sha=sha20),
            _decade_bl_parquet_row(w_to, 2022, "fy_2020_actuals", 150000.0,
                                   sha=sha22),
        ]
        _write_parquet(site / "data" / "budget_lines_decade.parquet",
                       _BL_COL_DEFS, rows_bl)
        cit_rows = [
            _decade_wb_citation_row(w1, 2017, "fy_2015_actuals", 60000.0,
                                    sha=sha17, cells="J4"),
            _decade_wb_citation_row(w2, 2017, "fy_2015_actuals", 30000.0,
                                    sha=sha17, cells="J5"),
            _decade_wb_citation_row(w_from, 2020, "fy_2020_total", 140000.0,
                                    sha=sha20),
            _decade_wb_citation_row(w_to, 2022, "fy_2020_actuals", 150000.0,
                                    sha=sha22),
            _make_derived_row(
                sum_fid,
                "sum(budget_lines.amount_thousands where"
                " amount_type=fy_2015_actuals and edition=2017)",
                json.dumps([w1, w2]), recorded_sum),
            _make_derived_row(
                diff_fid,
                "PB2022 FY2020 actuals - PB2020 FY2020 request"
                " (fct_book_diff request_vs_actuals)",
                json.dumps([w_to, w_from]), recorded_delta),
        ]
        _write_parquet(site / "citations" / "citations.parquet",
                       _CIT_COL_DEFS, cit_rows)
        _write_manifest(site)
        return sum_fid, diff_fid

    def test_decade_sum_and_diff_recompute_pass(self, tmp_path):
        site = tmp_path / "site"
        self._site_with_diff(site, "10000.000")
        result = citation_gate5b1(site)
        assert result["ok"] is True, f"failures: {result.get('failures')}"

    def test_wrong_sum_fails(self, tmp_path):
        site = tmp_path / "site"
        sum_fid, _ = self._site_with_diff(site, "10000.000",
                                          recorded_sum="99999.000")
        result = citation_gate5b1(site)
        assert result["ok"] is False, "tampered decade sum must FAIL"
        assert any(f[0] == sum_fid for f in result["failures"])

    def test_wrong_diff_delta_fails(self, tmp_path):
        """A book-diff fact whose recorded delta does not recompute from its
        two side amounts must FAIL — this is the proof-can-fail for the
        difference-rule fallback (pre-extension, workbook-side inputs were
        silently shape-checked and a corrupted delta would have shipped)."""
        site = tmp_path / "site"
        _, diff_fid = self._site_with_diff(site, "77777.000")
        result = citation_gate5b1(site)
        assert result["ok"] is False, "tampered book-diff delta must FAIL"
        assert any(f[0] == diff_fid and "difference" in f[1]
                   for f in result["failures"])


# ---------------------------------------------------------------------------
# Backlog #24 — budget_lines ↔ budget_lines_decade overlap equality
# ---------------------------------------------------------------------------

from govbudget.verify_phase5b1 import _load_fid_to_bl_amount

_BL_COL_DEFS = (
    "fact_id varchar, exhibit varchar, fiscal_year integer, account varchar,"
    " account_title varchar, organization varchar, budget_activity varchar,"
    " budget_activity_title varchar, pe_bli varchar, title varchar,"
    " amount_type varchar, amount_thousands double, units varchar,"
    " document_sha256 varchar, source_sheet varchar, source_cells varchar"
)


def _bl_row(fid: str, amount: float) -> tuple:
    return (fid, "R-1", 2026, "0400", "RDT&E Defense-Wide", "DARPA", "01",
            "Basic Research", "0601101E", "TITLE", "fy_2024_actuals",
            amount, "USD thousands", "sha", "Exhibit R-1", "J2")


class TestOverlapEquality:
    """5,257 fids exist in both budget_lines and budget_lines_decade
    (edition-2026 decade grains dedupe against the main export). setdefault
    alone let a divergent decade amount hide behind the budget_lines copy —
    the overlap must be exactly Decimal-equal or the gate FAILs."""

    def test_helper_equal_overlap_no_divergence(self, tmp_path):
        site = tmp_path / "site"
        _write_parquet(site / "data" / "budget_lines.parquet", _BL_COL_DEFS,
                       [_bl_row("aaaa000011110000", 280494.0)])
        _write_parquet(site / "data" / "budget_lines_decade.parquet", _BL_COL_DEFS,
                       [_bl_row("aaaa000011110000", 280494.0),
                        _bl_row("bbbb000022220000", 5.0)])
        mapping, overlap, divergences = _load_fid_to_bl_amount(site)
        assert overlap == 1
        assert divergences == []
        assert mapping["aaaa000011110000"] == "280494.0"
        assert mapping["bbbb000022220000"] == "5.0"

    def test_helper_divergent_overlap_reported(self, tmp_path):
        site = tmp_path / "site"
        _write_parquet(site / "data" / "budget_lines.parquet", _BL_COL_DEFS,
                       [_bl_row("aaaa000011110000", 280494.0)])
        _write_parquet(site / "data" / "budget_lines_decade.parquet", _BL_COL_DEFS,
                       [_bl_row("aaaa000011110000", 280495.0)])  # divergent
        _, overlap, divergences = _load_fid_to_bl_amount(site)
        assert overlap == 1
        assert divergences == [("aaaa000011110000", "280494.0", "280495.0")]

    def test_helper_no_decade_parquet_is_clean(self, tmp_path):
        site = tmp_path / "site"
        _write_parquet(site / "data" / "budget_lines.parquet", _BL_COL_DEFS,
                       [_bl_row("aaaa000011110000", 280494.0)])
        mapping, overlap, divergences = _load_fid_to_bl_amount(site)
        assert overlap == 0
        assert divergences == []
        assert len(mapping) == 1

    def test_gate_fails_on_divergent_decade_copy(self, tmp_path):
        """Proof-can-fail: an otherwise-green site whose decade parquet
        carries a divergent amount for an overlapping fid FAILs gate 1,
        naming the fid."""
        site = tmp_path / "site"
        _, fid = _make_site_with_workbook(site)
        # decade copy of the SAME fid with a tampered amount
        _write_parquet(site / "data" / "budget_lines_decade.parquet", _BL_COL_DEFS,
                       [_bl_row(fid, 280495.0)])
        result = citation_gate5b1(site)
        assert result["ok"] is False, "divergent overlap must FAIL the gate"
        assert result["overlap_fids"] == 1
        assert result["overlap_divergent"] == 1
        assert any(f[0] == fid and "divergent amount_thousands" in f[1]
                   for f in result["failures"])
        # the sampled citation itself still re-derives — passed reflects
        # sample results only, the overlap failure is additive
        assert result["passed"] == result["sampled"]

    def test_gate_passes_with_equal_decade_copy(self, tmp_path):
        site = tmp_path / "site"
        _, fid = _make_site_with_workbook(site)
        _write_parquet(site / "data" / "budget_lines_decade.parquet", _BL_COL_DEFS,
                       [_bl_row(fid, 280494.0)])
        result = citation_gate5b1(site)
        assert result["ok"] is True, result["failures"]
        assert result["overlap_fids"] == 1
        assert result["overlap_divergent"] == 0


# ---------------------------------------------------------------------------
# Gate 1 addendum: the row-label leg
# ---------------------------------------------------------------------------


class TestRowLabelLeg:
    """PM review 2026-07-30 §Systemic fix, the addendum (spec line 307):

        "Also worth adding to the existing 50-facts-per-build provenance
         spot-check: verify the **highlighted region** matches the expected row
         label, not merely that the document exists and the hash matches.
         P0-1's underlying mislabel would not have been caught by an existence
         check."

    P0-1 (spec :52-55) was two rows of ONE P-40 Resource Summary: the PDF
    citation resolved to "Net Procurement (P-1) ($ in Millions)" = 5,247.070
    while the published basis was "Total Obligation Authority" = 5,565.655,
    the two bridged by "Plus CY Advance Procurement 318.585" one row above.
    Document, hash, page and bbox all agreed. Only the ROW disagreed.

    Fixture page 1 of tests/fixtures/jbooks/darpa_p24_25.pdf carries exactly
    one "280.494" on the row "Total Program Element" (x0 235.23, top 158.47)
    and exactly one "160.158" on the row "CCS-02: MATH AND" (x0 235.23,
    top 173.97) — verified 2026-09-18.
    """

    # ── the pure helpers ────────────────────────────────────────────────
    def test_row_label_reads_the_row_and_drops_the_value_columns(self):
        from govbudget.verify_phase5b1 import _row_label

        words = [
            {"text": "Total", "x0": 20.0, "x1": 44.0, "top": 158.7, "bottom": 167.7},
            {"text": "Program", "x0": 45.0, "x1": 85.0, "top": 158.7, "bottom": 167.7},
            {"text": "Element", "x0": 86.0, "x1": 128.0, "top": 158.7, "bottom": 167.7},
            {"text": "0.000", "x0": 194.8, "x1": 227.0, "top": 158.5, "bottom": 167.5},
            {"text": "280.494", "x0": 235.2, "x1": 267.7, "top": 158.5, "bottom": 167.5},
            {"text": "999.999", "x0": 20.0, "x1": 52.0, "top": 400.0, "bottom": 409.0},
        ]
        assert _row_label(words, 235.2, 158.5) == "Total Program Element"
        # A row whose left side is nothing but value columns has no label.
        assert _row_label(words, 235.2, 400.0) == ""

    def test_label_names_fact_accepts_the_project_the_pe_and_the_summary_rows(self):
        """Each accepting branch is named, so item 4's counters can split by it."""
        from govbudget.verify_phase5b1 import _label_names_fact

        page = _words_naming("PE", "0601101E", "/", "DEFENSE", "RESEARCH", "SCIENCES")
        assert _label_names_fact(
            "CCS-02: MATH AND", "0601101E", "CCS-02", "Math and Computer Sciences", page
        ) == "project_line"
        assert _label_names_fact(
            "Total Program Element", "0601101E", None, None, page
        ) == "summary_row"
        assert _label_names_fact(
            "Net Procurement (P-1) ($ in Millions)", "JASSM0", None, None,
            _words_naming("P-1", "Line", "Item", "JASSM0"),
        ) == "summary_row"
        assert _label_names_fact(
            "153 0604840F F-35 C2D2 07 U", "0604840F", None, None, page
        ) == "pe_line"
        assert _label_names_fact(
            "CCS-02: MATH AND", "0601101E", "001",
            "Defense Technical Information Center", page,
        ) is None
        assert _label_names_fact("", "0601101E", None, None, page) is None
        # A unit-count row is not a money row: 2 citations in the 2026-09-10
        # export cite "19,452" whose TOA column reads "19.452".
        assert _label_names_fact(
            "Procurement Quantity (Units in Each)", "0312MB7000", None, None,
            _words_naming("0312MB7000"),
        ) is None

    def test_bare_summary_row_needs_the_facts_pe_printed_on_the_page(self):
        """Item 3: a summary row belongs to whatever PE the PAGE is about.

        An R-2 volume prints "Total Program Element" for every PE in it, so
        accepting the label alone lets another PE's total stand in for this
        fact — the RDT&E analog of P0-1. The page must print this fact's own
        pe_bli for the bare row to be accepted.
        """
        from govbudget.verify_phase5b1 import _label_names_fact

        ours = _words_naming("PE", "0601101E", "/", "DEFENSE", "RESEARCH", "SCIENCES")
        theirs = _words_naming("PE", "0605801KA", "/", "MISSION", "SUPPORT")
        assert _label_names_fact(
            "Total Program Element", "0601101E", None, None, ours
        ) == "summary_row"
        assert _label_names_fact(
            "Total Program Element", "0601101E", None, None, theirs
        ) is None
        # A project fact on another PE's summary row is refused the same way.
        assert _label_names_fact(
            "Total Program Element", "0601101E", "CCS-02", "Math and Computer",
            theirs,
        ) is None

    def test_pe_match_in_the_label_is_word_boundary_anchored(self):
        """Item 10: 1,640 BLIs are ≤4 all-digit chars — a bare substring of the
        row's fiscal year or line number must not stand in for the BLI."""
        from govbudget.verify_phase5b1 import _label_names_fact

        page = _words_naming("nothing", "here")
        # The R-1-line behaviour that must survive (tests:2185).
        assert _label_names_fact(
            "153 0604840F F-35 C2D2 07 U", "0604840F", None, None, page
        ) == "pe_line"
        # "0449" inside "10449" is not BLI 0449.
        assert _label_names_fact("10449 Spares and Repair", "0449", None, None, page) is None
        assert _label_names_fact("0449 Spares and Repair", "0449", None, None, page) == "pe_line"

    def test_toa_column_value_reads_the_basis_row_in_the_same_column(self):
        from govbudget.verify_phase5b1 import _toa_column_value

        words = _P0_1_WORDS()
        assert _toa_column_value(words, 347.1, 390.0) == "5,565.655"
        assert _toa_column_value(words, 10.0, 18.0) is None

    def test_toa_column_is_found_by_span_overlap_not_the_left_edge(self):
        """Item 1: P-40 value columns are RIGHT-aligned.

        Reviewer's exemplar — doc e7e1302…, page 223: the column whose right
        edge is x1 368.5 holds `407.046` at x0 343.3 and `52.191` at x0 347.1.
        One character moves x0 by 3.8 pt, past the 2.0 pt column tolerance, so
        a left-edge match finds the TOA cell only when the two numbers happen
        to be the same width — it goes blind precisely when they DISAGREE,
        which is the only case that matters.
        """
        from govbudget.verify_phase5b1 import _toa_column_value

        words = [
            {"text": "Net", "x0": 20.0, "x1": 38.0, "top": 208.8, "bottom": 217.8},
            {"text": "Procurement", "x0": 39.0, "x1": 95.0, "top": 208.8, "bottom": 217.8},
            {"text": "52.191", "x0": 347.1, "x1": 368.5, "top": 208.8, "bottom": 217.8},
            {"text": "Total", "x0": 20.0, "x1": 44.0, "top": 232.9, "bottom": 241.9},
            {"text": "Obligation", "x0": 45.0, "x1": 95.0, "top": 232.9, "bottom": 241.9},
            {"text": "Authority", "x0": 96.0, "x1": 140.0, "top": 232.9, "bottom": 241.9},
            {"text": "407.046", "x0": 343.3, "x1": 368.5, "top": 232.9, "bottom": 241.9},
        ]
        assert _toa_column_value(words, 347.1, 368.5) == "407.046"

    def test_row_basis_check_fails_the_p0_1_species(self):
        """The proof it can fail on the defect the addendum was written for."""
        from govbudget.verify_phase5b1 import _row_basis_check

        outcome, failure = _row_basis_check(
            _P0_1_WORDS(), "Net Procurement (P-1) ($ in Millions)", 347.1, 390.0,
            "5,247.070", 55,
        )
        assert outcome == "compared"
        assert failure is not None
        assert "5,247.070" in failure and "5,565.655" in failure
        assert "Total Obligation Authority" in failure
        assert "page 55" in failure

    def test_row_basis_check_passes_when_the_two_rows_carry_the_same_number(self):
        """The common real case (fact 06f8fad689ea561f, Army P-40 page 49)."""
        from govbudget.verify_phase5b1 import _row_basis_check

        words = [
            {"text": "Net", "x0": 20.0, "x1": 38.0, "top": 208.8, "bottom": 217.8},
            {"text": "Procurement", "x0": 39.0, "x1": 95.0, "top": 208.8, "bottom": 217.8},
            {"text": "(P-1)", "x0": 96.0, "x1": 120.0, "top": 208.8, "bottom": 217.8},
            {"text": "732.060", "x0": 347.1, "x1": 390.0, "top": 208.8, "bottom": 217.8},
            {"text": "Total", "x0": 20.0, "x1": 44.0, "top": 232.9, "bottom": 241.9},
            {"text": "Obligation", "x0": 45.0, "x1": 95.0, "top": 232.9, "bottom": 241.9},
            {"text": "Authority", "x0": 96.0, "x1": 140.0, "top": 232.9, "bottom": 241.9},
            {"text": "732.060", "x0": 347.1, "x1": 390.0, "top": 232.9, "bottom": 241.9},
        ]
        assert _row_basis_check(
            words, "Net Procurement (P-1)", 347.1, 390.0, "732.060", 49
        ) == ("compared", None)

    def test_row_basis_check_reports_a_missing_toa_row_rather_than_silently_skipping(self):
        """Item 2: the carve-out is counted, so an inert leg is visible."""
        from govbudget.verify_phase5b1 import _row_basis_check

        words = [
            {"text": "Net", "x0": 20.0, "x1": 38.0, "top": 208.8, "bottom": 217.8},
            {"text": "Procurement", "x0": 39.0, "x1": 95.0, "top": 208.8, "bottom": 217.8},
            {"text": "732.060", "x0": 347.1, "x1": 390.0, "top": 208.8, "bottom": 217.8},
        ]
        assert _row_basis_check(
            words, "Net Procurement (P-1)", 347.1, 390.0, "732.060", 49
        ) == ("toa_not_found", None)

    def test_row_basis_check_is_a_no_op_on_the_toa_row_itself(self):
        from govbudget.verify_phase5b1 import _row_basis_check

        assert _row_basis_check(
            _P0_1_WORDS(), "Total Obligation Authority ($ in Millions)", 347.1, 390.0,
            "5,565.655", 55,
        ) == ("on_toa_row", None)

    def test_trailing_value_columns_are_stripped_case_insensitively(self):
        """Item 8: `Continuing` and `TBD` print in whatever case the PDF uses."""
        from govbudget.verify_phase5b1 import _row_label

        def row(top, texts):
            out, x = [], 20.0
            for t in texts:
                out.append({"text": t, "x0": x, "x1": x + 6.0 * len(t), "top": top,
                            "bottom": top + 9.0})
                x += 6.0 * len(t) + 2.0
            return out

        words = row(100.0, ["Total", "Program", "Element", "CONTINUING", "tbd", "0.000"])
        words.append({"text": "280.494", "x0": 300.0, "x1": 332.0, "top": 100.0,
                      "bottom": 109.0})
        assert _row_label(words, 300.0, 100.0) == "Total Program Element"

    # ── the gate, end to end, on the real fixture PDF ───────────────────
    def test_highlight_on_another_projects_row_fails(self, tmp_path):
        """bbox, sha and page all correct; the ROW belongs to another project."""
        site = tmp_path / "site"
        fid = _make_site_with_jbook_pdf_at_word(
            site,
            amount_text="160.158",          # the CCS-02 row on page 1
            project_number="001",
            project_title="Defense Technical Information Center",
        )
        _write_manifest(site, datasets={"jbook_details": 1}, citations={"jbook_pdf": 1})
        res = citation_gate5b1(site)
        assert res["ok"] is False
        assert len(res["failures"]) == 1
        bad_fid, reason = res["failures"][0]
        assert bad_fid == fid
        assert "CCS-02" in reason
        assert "names neither" in reason

    def test_highlight_on_its_own_project_row_passes_and_is_counted(self, tmp_path):
        site = tmp_path / "site"
        _make_site_with_jbook_pdf_at_word(
            site,
            amount_text="160.158",
            project_number="CCS-02",
            project_title="Math and Computer Sciences",
        )
        _write_manifest(site, datasets={"jbook_details": 1}, citations={"jbook_pdf": 1})
        res = citation_gate5b1(site)
        assert res["ok"] is True, res["failures"]
        assert res["row_label"]["checked"] == 1
        assert res["row_label"]["fallback_summary_row"] == 0
        assert res["row_label"]["fallback_pe_line"] == 0
        assert res["row_label"]["fallback_title_match"] == 0
        assert res["row_label"]["unreadable"] == 0

    def test_project_fact_on_the_pe_row_passes_but_is_counted_weak(self, tmp_path):
        site = tmp_path / "site"
        _make_site_with_jbook_pdf_at_word(
            site,
            amount_text="280.494",          # the "Total Program Element" row
            project_number="001",
            project_title="Defense Technical Information Center",
        )
        _write_manifest(site, datasets={"jbook_details": 1}, citations={"jbook_pdf": 1})
        res = citation_gate5b1(site)
        assert res["ok"] is True, res["failures"]
        # Item 4: the branch that accepted it is named. This one is a genuine
        # summary row; the PE's own rendered R-1 line and a title match are
        # counted apart, because they are not the same weakness.
        assert res["row_label"]["fallback_summary_row"] == 1
        assert res["row_label"]["fallback_pe_line"] == 0
        assert res["row_label"]["fallback_title_match"] == 0

    def test_row_band_is_centred_on_the_matched_word_not_the_stored_top(self):
        """Item 9: stored_top may sit up to _BBOX_TOL_PT off the printed row.

        Two rows 4 pt apart: centring the band on a stored_top that is 2 pt
        high (still inside the bbox tolerance) sweeps in BOTH rows; centring
        it on the word actually found reads only its own.
        """
        from govbudget.verify_phase5b1 import _row_label

        words = [
            {"text": "Other", "x0": 20.0, "x1": 50.0, "top": 100.0, "bottom": 109.0},
            {"text": "Row", "x0": 51.0, "x1": 70.0, "top": 100.0, "bottom": 109.0},
            {"text": "Total", "x0": 20.0, "x1": 44.0, "top": 104.0, "bottom": 113.0},
            {"text": "Program", "x0": 45.0, "x1": 85.0, "top": 104.0, "bottom": 113.0},
            {"text": "Element", "x0": 86.0, "x1": 128.0, "top": 104.0, "bottom": 113.0},
            {"text": "280.494", "x0": 235.2, "x1": 267.7, "top": 104.0, "bottom": 113.0},
        ]
        assert _row_label(words, 235.2, 104.0) == "Total Program Element"
        mixed = _row_label(words, 235.2, 102.0)
        assert mixed != "Total Program Element"
        assert "Other" in mixed

    def test_offset_stored_top_inside_the_tolerance_still_passes_the_gate(self, tmp_path):
        """The gate-level half of item 9, on the real fixture."""
        site = tmp_path / "site"
        _make_site_with_jbook_pdf_at_word(
            site,
            amount_text="160.158",
            project_number="CCS-02",
            project_title="Math and Computer Sciences",
            top_offset=1.9,
        )
        _write_manifest(site, datasets={"jbook_details": 1}, citations={"jbook_pdf": 1})
        res = citation_gate5b1(site)
        assert res["ok"] is True, res["failures"]
        assert res["row_label"]["checked"] == 1

    # ── item 3, end to end: the summary row belongs to the page's PE ────
    def test_summary_row_on_another_pes_page_fails(self, tmp_path):
        """Item 3: the RDT&E analog of P0-1.

        An R-2 volume prints "Total Program Element" once per PE. Accepting
        that label with nothing tying the row to this fact's pe_bli lets
        another PE's total stand in — and it was counted as a benign
        fallback, not failed.
        """
        site = tmp_path / "site"
        fid = _make_site_with_jbook_pdf_at_word(
            site,
            amount_text="280.494",            # the "Total Program Element" row
            project_number=None,
            project_title=None,
            pe_bli="0602702E",                # NOT printed on fixture page 1
        )
        _write_manifest(site, datasets={"jbook_details": 1}, citations={"jbook_pdf": 1})
        res = citation_gate5b1(site)
        assert res["ok"] is False
        bad_fid, reason = res["failures"][0]
        assert bad_fid == fid
        assert "another program element" in reason
        assert "0602702E" in reason
        assert "280.494" in reason          # item 7: the number is in the message
        assert res["row_label"]["fallback_summary_row"] == 0

    def test_summary_row_on_its_own_pes_page_still_passes(self, tmp_path):
        """The other side of item 3 — the 2026-09-18 corpus's dominant route."""
        site = tmp_path / "site"
        _make_site_with_jbook_pdf_at_word(
            site,
            amount_text="280.494",
            project_number=None,
            project_title=None,
            pe_bli="0601101E",                # printed as "PE 0601101E / DEFENSE…"
        )
        _write_manifest(site, datasets={"jbook_details": 1}, citations={"jbook_pdf": 1})
        res = citation_gate5b1(site)
        assert res["ok"] is True, res["failures"]
        assert res["row_label"]["checked"] == 1

    # ── item 2: every skip path is counted ──────────────────────────────
    def test_citation_with_no_detail_row_is_counted_as_a_skip(self, tmp_path):
        site = tmp_path / "site"
        _make_site_with_jbook_pdf_at_word(
            site,
            amount_text="280.494",
            project_number=None,
            project_title=None,
            write_detail=False,
        )
        _write_manifest(site, datasets={"jbook_details": 0}, citations={"jbook_pdf": 1})
        res = citation_gate5b1(site)
        assert res["ok"] is True, res["failures"]
        rl = res["row_label"]
        assert rl["sampled"] == 1
        assert rl["checked"] == 0
        assert rl["skipped_no_detail"] == 1

    def test_null_bbox_is_counted_as_a_skip(self, tmp_path):
        site = tmp_path / "site"
        _make_site_with_jbook_pdf_at_word(
            site,
            amount_text="280.494",
            project_number=None,
            project_title=None,
            null_bbox=True,
        )
        _write_manifest(site, datasets={"jbook_details": 1}, citations={"jbook_pdf": 1})
        res = citation_gate5b1(site)
        assert res["ok"] is True, res["failures"]
        rl = res["row_label"]
        assert rl["sampled"] == 1
        assert rl["checked"] == 0
        assert rl["skipped_null_bbox"] == 1

    def test_procurement_page_without_a_toa_row_counts_toa_not_found(self, tmp_path):
        """The carve-out named in the module docstring, made visible.

        The R-2 fixture has no `Obligation` anywhere, so a procurement fact
        pointed at it cannot have its basis compared. That is counted, never
        failed — and the counter is the only thing that says so.
        """
        site = tmp_path / "site"
        _make_site_with_jbook_pdf_at_word(
            site,
            amount_text="280.494",
            project_number=None,
            project_title=None,
            exhibit_family="procurement",
        )
        _write_manifest(site, datasets={"jbook_details": 1}, citations={"jbook_pdf": 1})
        res = citation_gate5b1(site)
        assert res["ok"] is True, res["failures"]
        rl = res["row_label"]
        assert rl["toa_not_found"] == 1
        assert rl["basis_checked"] == 0
        assert rl["basis_compared"] == 0

    # ── item 4: the fallback counter splits by accepting branch ─────────
    def test_fallback_counters_split_by_accepting_branch(self, tmp_path):
        """A project fact matched by TITLE is not the same weakness as one
        matched by a bare summary row — the reviewer's sample of 60 found 12
        counted and only 3 of them summary rows."""
        site = tmp_path / "site"
        _make_site_with_jbook_pdf_at_word(
            site,
            amount_text="160.158",
            project_number="CCS-99",              # not the printed project number
            project_title="CCS-02: MATH AND COMPUTER SCIENCES",
        )
        _write_manifest(site, datasets={"jbook_details": 1}, citations={"jbook_pdf": 1})
        res = citation_gate5b1(site)
        assert res["ok"] is True, res["failures"]
        rl = res["row_label"]
        assert rl["fallback_title_match"] == 1
        assert rl["fallback_summary_row"] == 0
        assert rl["fallback_pe_line"] == 0


# ---------------------------------------------------------------------------
# Item 5: the procurement basis check, on a real P-40-shaped PDF
# ---------------------------------------------------------------------------


class TestProcurementBasisOnAP40Page:
    """`exhibit_family == "procurement"` had ZERO end-to-end coverage.

    No test passed `exhibit_family="procurement"` and the shipped R-2 fixture
    has no `Obligation` word, so `_row_basis_check` — the P0-1 check, the
    whole point of the addendum — was exercised only on hand-built word
    lists. tests/fixtures/jbooks/p40_resource_summary.pdf is a synthetic
    Resource Summary page written by
    tests/fixtures/jbooks/make_p40_fixture.py (the venv has no PDF writer, so
    the generator emits a minimal PDF 1.4 with Helvetica text objects);
    extraction here goes through real pdfplumber words.
    """

    def test_fixture_value_columns_are_right_aligned(self):
        """The property the fixture exists to have (and item 1's premise)."""
        words = _page_words(P40_FIXTURE_PDF)
        net = next(w for w in words if w["text"] == "52.191")
        toa = next(w for w in words if w["text"] == "407.046")
        assert abs(float(net["x1"]) - float(toa["x1"])) < 0.01   # one column
        assert abs(float(net["x0"]) - float(toa["x0"])) > 2.0    # past the tolerance

    def test_highlight_off_the_toa_row_with_a_different_basis_fails(self, tmp_path):
        """P0-1 itself: same document, same hash, same page, same column —
        only the ROW disagrees, and the two cells are different widths."""
        site = tmp_path / "site"
        fid = _make_site_with_jbook_pdf_at_word(
            site,
            amount_text="52.191",              # Net Procurement (P-1), FY2025
            project_number=None,
            project_title=None,
            exhibit_family="procurement",
            pdf_path=P40_FIXTURE_PDF,
            pe_bli="0449",
        )
        _write_manifest(site, datasets={"jbook_details": 1}, citations={"jbook_pdf": 1})
        res = citation_gate5b1(site)
        assert res["ok"] is False
        bad_fid, reason = res["failures"][0]
        assert bad_fid == fid
        assert "52.191" in reason and "407.046" in reason
        assert "Total Obligation Authority" in reason
        assert res["row_label"]["basis_compared"] == 1

    def test_highlight_on_the_toa_row_passes_and_is_counted(self, tmp_path):
        site = tmp_path / "site"
        _make_site_with_jbook_pdf_at_word(
            site,
            amount_text="407.046",             # Total Obligation Authority, FY2025
            project_number=None,
            project_title=None,
            exhibit_family="procurement",
            pdf_path=P40_FIXTURE_PDF,
            pe_bli="0449",
        )
        _write_manifest(site, datasets={"jbook_details": 1}, citations={"jbook_pdf": 1})
        res = citation_gate5b1(site)
        assert res["ok"] is True, res["failures"]
        rl = res["row_label"]
        assert rl["basis_checked"] == 1        # the basis is the highlight itself
        assert rl["basis_compared"] == 0       # nothing to compare it against
        assert rl["toa_not_found"] == 0

    def test_same_page_as_rdte_runs_no_basis_check(self, tmp_path):
        """The basis rule is procurement-only: an R-2 amount off a P-40-shaped
        page must not be compared against a TOA row."""
        site = tmp_path / "site"
        _make_site_with_jbook_pdf_at_word(
            site,
            amount_text="52.191",
            project_number=None,
            project_title=None,
            exhibit_family="rdte",
            pdf_path=P40_FIXTURE_PDF,
            pe_bli="0449",
        )
        _write_manifest(site, datasets={"jbook_details": 1}, citations={"jbook_pdf": 1})
        res = citation_gate5b1(site)
        assert res["ok"] is True, res["failures"]
        rl = res["row_label"]
        assert rl["checked"] == 1
        assert rl["basis_checked"] == 0
        assert rl["basis_compared"] == 0
        assert rl["toa_not_found"] == 0

