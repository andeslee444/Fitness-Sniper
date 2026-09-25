"""Sitewide P-1/R-1 evidence, bound to original workbook cells and PDF glyphs.

Only additive TOA facts acquire summed PDF receipts. Differences, percentages,
obligations, and R-2/P-40 detail keep their own accounting and original receipts.
"""
from __future__ import annotations

from collections import defaultdict
from ctypes import c_double
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import re
import os

import openpyxl
import pdfplumber
import pypdfium2 as pdfium

from govbudget.budget_pdf_receipts import (amount, close_line, combine_receipts, match_cell,
    page_columns, pdf_header_key, program_lines, select_glyph_geometry, title_identity, wrapped_title, inherited_activity)


def source_identity(citation: dict) -> tuple[int, str] | None:
    url = citation.get("official_url") or ""
    match = re.fullmatch(r"https://comptroller\.(?:war|defense)\.gov/Portals/45/Documents/defbudget/fy(\d{4})/(p1|r1|p1r)_display\.xlsx", url, re.I)
    if not match:
        return None
    return int(match[1]), {"p1": "P-1", "r1": "R-1", "p1r": "P-1R"}[match[2].lower()]


def source_file(record: dict, site_dir: Path) -> Path:
    path = site_dir / "pdfs" / f"{record['sha256']}.pdf"
    if not path.exists():
        import httpx
        response = httpx.get(record["url"], follow_redirects=True, timeout=90)
        response.raise_for_status()
        if hashlib.sha256(response.content).hexdigest() != record["sha256"]:
            raise ValueError(f"Government PDF changed; review required: {record['url']}")
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(response.content)
    if hashlib.sha256(path.read_bytes()).hexdigest() != record["sha256"]:
        raise ValueError(f"PDF hash mismatch: {path}")
    return path


def source_pages(record: dict, path: Path, cache_dir: Path) -> list[dict]:
    """Scan every detail page in the pinned document, not a family page list."""
    cache = cache_dir / f"{record['sha256']}-program-pages-v1.json"
    if cache.exists():
        return json.loads(cache.read_text())
    selected = []
    with pdfium.PdfDocument(path) as document:
        for i in range(len(document)):
            page = document[i]
            textpage = page.get_textpage()
            text = textpage.get_text_range()
            textpage.close()
            page.close()
            if (f"FY {record['edition']} President's Budget" in text and
                f"Exhibit {record['exhibit']}" in text and
                "(Dollars in Thousands)" in text and "Appropriation:" in text):
                selected.append(i)
    pages = []
    with pdfplumber.open(path) as document:
        for i in selected:
            page = document.pages[i]
            text = page.extract_text() or ""
            accounts = re.findall(r"\b\d{4}[A-Z]\b", text)
            if not accounts or page.rotation:
                continue
            words = page.extract_words(return_chars=True)
            # Retain exact baseline identity, not the extractor's large font dictionaries.
            for word in words:
                word["chars"] = [{"text": c["text"], "matrix": c["matrix"]} for c in word["chars"]]
            pages.append({"page_number": i + 1, "width": page.width, "height": page.height,
                          "mediabox": list(page.mediabox), "cropbox": list(page.cropbox),
                          "rotation": page.rotation, "text": text, "words": words,
                          "accounts": sorted(set(accounts))})
            page.flush_cache()
    cache.parent.mkdir(parents=True, exist_ok=True)
    staging = cache.with_suffix(f".{os.getpid()}.tmp")
    staging.write_text(json.dumps(pages, separators=(",", ":")))
    staging.replace(cache)
    return pages


class GlyphReader:
    """Reuse document/text handles across the whole workbook, then close them."""
    def __init__(self, path: Path):
        self.document = pdfium.PdfDocument(path)
        self.pages = {}
        self.cache = {}

    def geometry(self, page: dict, word: dict) -> dict:
        key = (page["page_number"], word["text"], tuple(word["chars"][0]["matrix"]))
        if key in self.cache:
            return self.cache[key]
        number = page["page_number"]
        if number not in self.pages:
            rendered = self.document[number - 1]
            self.pages[number] = rendered, rendered.get_textpage()
        rendered, textpage = self.pages[number]
        search = textpage.search(word["text"], match_case=True)
        candidates = []
        try:
            while found := search.get_next():
                index, count = found
                x, y = c_double(), c_double()
                if not pdfium.raw.FPDFText_GetCharOrigin(textpage, index, x, y):
                    raise ValueError("Missing exact source character origin")
                boxes = [textpage.get_charbox(i, loose=False) for i in range(index, index + count)]
                candidates.append({"text": textpage.get_text_range(index, count), "origin": (x.value, y.value),
                                   "box": (min(b[0] for b in boxes), min(b[1] for b in boxes), max(b[2] for b in boxes), max(b[3] for b in boxes))})
        finally:
            search.close()
        result = select_glyph_geometry(page, word, rendered.get_bbox(), candidates)
        self.cache[key] = result
        return result

    def close(self):
        for page, textpage in self.pages.values():
            textpage.close()
            page.close()
        self.document.close()


def metadata_for_row(sheet, row: int, exhibit: str) -> dict:
    if exhibit == "P-1R":
        return {"account": str(sheet[f"A{row}"].value), "activity": str(sheet[f"D{row}"].value),
                "line": "", "code": str(sheet[f"H{row}"].value), "title": str(sheet[f"I{row}"].value),
                "cost_type": str(sheet[f"J{row}"].value), "row_label": str(sheet[f"K{row}"].value)}
    return {"account": str(sheet[f"A{row}"].value), "activity": str(sheet[f"D{row}"].value),
            "line": str(sheet[f"F{row}"].value).strip(),
            "code": str(sheet[f"I{row}" if exhibit == "P-1" else f"G{row}"].value),
            "title": str(sheet[f"J{row}" if exhibit == "P-1" else f"H{row}"].value),
            "cost_type": str(sheet[f"K{row}"].value), "row_label": str(sheet[f"L{row}"].value)}


def matched_amount(page: dict, y: float, column: dict, value: float) -> dict | None:
    words = [w for w in close_line(page["words"], y) if abs(w["x1"] - column["edge"]) <= 6 and amount(w["text"]) is not None]
    if not words and value == 0:
        return {"blank_zero": True, "pdf_column_label": column["header"], "page": page}
    if len(words) == 1 and amount(words[0]["text"]) == value:
        return {"word": words[0], "page": page, "pdf_column_label": column["header"]}
    return None


def reserve_titles(page: dict) -> list[dict]:
    """Group complete P-1R titles before their indented component rows."""
    if "_reserve_titles" in page:
        return page["_reserve_titles"]
    columns = page_columns(page, "P-1R")
    if not columns:
        return []
    words = page["words"]
    sec = next((w for w in words if w["text"] == "Sec"), None)
    if sec is None:
        return []
    anchors = sorted((w for w in words if columns[0]["body_top"] < w["top"] < page["height"] - 35 and abs(w["x0"] - 35.8) < 3), key=lambda w: w["top"])
    titles = []
    for anchor in anchors:
        line = " ".join(w["text"] for w in close_line(words, anchor["top"]) if w["x1"] < sec["x0"])
        if line.startswith(("Budget Activity", "Total ")):
            continue
        if titles and anchor["top"] - titles[-1]["last_y"] < 13:
            titles[-1]["title"] += " " + line
            titles[-1]["last_y"] = anchor["top"]
        else:
            titles.append({"title": line, "top": anchor["top"], "last_y": anchor["top"]})
    page["_reserve_titles"] = titles
    return titles


def reserve_match(page: dict, row: dict, header: str, value: float, edition: int) -> dict | None:
    """P-1R requires the complete title, inherited activity and component."""
    columns = [c for c in page_columns(page, "P-1R") if c["normalized"] == pdf_header_key(header, edition=edition, exhibit="P-1R")]
    if len(columns) != 1 or row["account"] not in page["accounts"] or row["cost_type"] not in {"R", "T"}:
        return None
    words = page["words"]
    sec = next((w for w in words if w["text"] == "Sec"), None)
    if sec is None:
        return None
    titles = reserve_titles(page)
    anchors = [(page, a) for a in titles]
    previous = page.get("_previous")
    if previous and (prior_titles := reserve_titles(previous)):
        anchors.append((previous, prior_titles[-1]))
    candidates = []
    for identity_page, anchor in anchors:
        if title_identity(anchor["title"]) != title_identity(row["title"]) or inherited_activity(identity_page, anchor["top"]) != row["activity"]:
            continue
        start = anchor["last_y"] if identity_page is page else columns[0]["body_top"]
        end = min((a["top"] for a in titles if a["top"] > start + 2), default=page["height"] - 35)
        label = "National Guard" if row["cost_type"] == "T" else "Reserve"
        for word in words:
            if not start < word["top"] < end or word["text"] != label.split()[0]:
                continue
            component = " ".join(w["text"] for w in close_line(words, word["top"]) if w["x1"] < sec["x0"])
            if (component == label and inherited_activity(page, word["top"]) == row["activity"]
                    and (match := matched_amount(page, word["top"], columns[0], value))):
                if identity_page is not page:
                    match["identity_page_number"] = identity_page["page_number"]
                candidates.append(match)
    return candidates[0] if len(candidates) == 1 else None


def cost_row_match(page: dict, row: dict, header: str, value: float, edition: int) -> dict | None:
    """Match named shipbuilding components, including verified page continuations."""
    if row["cost_type"] in {"A", "B", "C", "", "None"}:
        return None
    columns = [c for c in page_columns(page, "P-1") if c["normalized"] == pdf_header_key(header, edition=edition, exhibit="P-1")]
    if len(columns) != 1:
        return None
    lines = program_lines(page)
    anchor_page = page
    anchors = [w for w in lines if w["text"] == row["line"]]
    if not anchors:
        previous = page.get("_previous")
        if previous is None:
            return None
        previous_lines = program_lines(previous)
        if not previous_lines or max(previous_lines, key=lambda w: w["top"])["text"] != row["line"]:
            return None
        anchor_page = previous
        anchors = [max(previous_lines, key=lambda w: w["top"])]
    if len(anchors) != 1:
        return None
    anchor = anchors[0]
    ident = next((w for w in anchor_page["words"] if w["text"] == "Ident"), None)
    if ident is None:
        return None
    title = wrapped_title(anchor_page["words"], anchor["top"], anchor["x1"], ident["x0"],
                          min((w["top"] for w in program_lines(anchor_page) if w["top"] > anchor["top"] + 2), default=anchor_page["height"] - 35),
                          max((w["top"] for w in program_lines(anchor_page) if w["top"] < anchor["top"] - 2), default=0))
    if title_identity(title) != title_identity(row["title"]):
        return None
    start = anchor["top"] if anchor_page is page else columns[0]["body_top"]
    end = min((w["top"] for w in lines if w["top"] > start + 2), default=page["height"] - 35)
    matches = []
    for word in page["words"]:
        if not start < word["top"] < end or abs(word["x0"] - 49.9) > 3:
            continue
        label = " ".join(w["text"] for w in close_line(page["words"], word["top"]) if w["x0"] >= word["x0"] and w["x1"] < page_columns(page, "P-1")[0]["edge"] - 45)
        if title_identity(label) == title_identity(row["row_label"]) and (match := matched_amount(page, word["top"], columns[0], value)):
            if anchor_page is not page:
                match["identity_page_number"] = anchor_page["page_number"]
            matches.append(match)
    return matches[0] if len(matches) == 1 else None


def additive_budget_formula(formula: str) -> bool:
    """Accept the exporter's sum/selection grammar, never arithmetic suffixes."""
    predicate = r"[a-z_]+=(?:'[^'\n]*'|[a-z0-9_]+)"
    selection = rf" where {predicate}(?: and {predicate})*"
    core = r"sum\(budget_lines\.amount_thousands"
    if re.fullmatch(rf"{core}(?:{selection})?\)(?:{selection})?", formula):
        return True
    tails = [
        r" for reviewed F-15 members, PB20\d{2}, FY20\d{2} (?:actuals|enacted|request); one selected scenario column per workbook line",
        r" for workbook program (?:P-1|R-1):[A-Za-z0-9:]+, PB20\d{2}, FY20\d{2} (?:actuals|enacted|request); each selected workbook input once",
        r" via annual reviewed F-15 actuals totals, FY20\d{2}–FY20\d{2}; each fiscal year once, nominal USD thousands; excludes enacted/current-year figures and requests",
    ]
    return any(re.fullmatch(core + r"\)" + tail, formula) for tail in tails)


def attach_additive_receipts(citations: dict, receipts: dict) -> None:
    visiting = set()
    leaves = {}
    for fact in receipts:
        citation = citations.get(fact, {})
        leaves[fact] = {(citation.get("sha256"), citation.get("sheet"), cell.strip())
                        for cell in citation.get("cells", "").split(",") if cell.strip()} or {fact}
    def visit(fact):
        if fact in receipts:
            return receipts[fact]
        citation = citations.get(fact, {})
        if fact in visiting or citation.get("kind") != "derived" or citation.get("units") != "USD thousands" or not additive_budget_formula(citation.get("formula") or ""):
            return None
        visiting.add(fact)
        try:
            ids = json.loads(citation.get("inputs") or "[]")
            if not isinstance(ids, list) or not ids or not all(isinstance(i, str) and re.fullmatch(r"[a-f0-9]{16}", i) for i in ids) or len(ids) != len(set(ids)):
                return None
            inputs = [visit(i) for i in ids]
            if any(r is None for r in inputs):
                return None
            input_leaves = [leaves[i] for i in ids]
            unique_leaves = set().union(*input_leaves)
            if len(unique_leaves) != sum(map(len, input_leaves)):
                return None
            expected = float(citation["recorded_value"])
            if sum(Decimal(str(r["amount_thousands"])) for r in inputs) != Decimal(str(expected)):
                raise ValueError(f"Additive budget receipt arithmetic differs: {fact}")
            receipt = combine_receipts(inputs, expected)
            documents = {d["sha256"]: d for r in inputs for d in r.get("source_documents", [])}
            receipt["source_documents"] = list(documents.values())
            receipts[fact] = receipt
            leaves[fact] = unique_leaves
            return receipt
        finally:
            visiting.remove(fact)
    for fact in citations:
        visit(fact)


def export_program_pdf_receipts(*, site_dir: Path, manifest: Path, cache_dir: Path, editions: set[int] | None = None) -> dict:
    citations = json.loads((site_dir / "json/citations.json").read_text())
    sources = {(r["edition"], r["exhibit"]): r for r in json.loads(manifest.read_text())}
    grouped = defaultdict(list)
    for fact, citation in citations.items():
        if citation["kind"] == "workbook" and (identity := source_identity(citation)) and (editions is None or identity[0] in editions):
            grouped[(identity, citation["sha256"])].append((fact, citation))
    receipts, unmatched, book_audits = {}, [], []
    for (identity, workbook_sha), entries in sorted(grouped.items()):
        edition, exhibit = identity
        record = sources.get(identity)
        if record is None:
            raise ValueError(f"Missing government PDF registry entry: {identity}")
        path = source_file(record, site_dir)
        pages = source_pages(record, path, cache_dir)
        by_account = defaultdict(list)
        previous = None
        for page in pages:
            if previous and previous["page_number"] == page["page_number"] - 1 and set(previous["accounts"]) & set(page["accounts"]):
                page["_previous"] = previous
            for account in page["accounts"]:
                by_account[account].append(page)
            previous = page
        workbook_path = site_dir / "workbooks" / f"{workbook_sha}.xlsx"
        if hashlib.sha256(workbook_path.read_bytes()).hexdigest() != workbook_sha:
            raise ValueError(f"Workbook hash mismatch: {workbook_path}")
        workbook = openpyxl.load_workbook(workbook_path, data_only=True)
        glyphs = GlyphReader(path)
        document = {"official_url": record["url"], "sha256": record["sha256"], "edition": edition, "exhibit": exhibit}
        complete_before = sum(r["complete"] for r in receipts.values())
        try:
            for fact, citation in entries:
                sheet = workbook[citation["sheet"]]
                parts, blanks, missing, values = [], 0, 0, []
                for cell in citation["cells"].split(","):
                    cell = cell.strip()
                    if not re.fullmatch(r"[A-Z]+[1-9]\d*", cell):
                        raise ValueError(f"Unsupported source cell locator: {fact}/{cell}")
                    address = sheet[cell]
                    value = float(address.value or 0)
                    values.append(Decimal(str(value)))
                    row = metadata_for_row(sheet, address.row, exhibit)
                    if citation.get("pe_bli") and str(citation["pe_bli"]) != row["code"]:
                        raise ValueError(f"Source program code differs: {fact}/{cell}")
                    header = str(sheet.cell(2, address.column).value)
                    matches = []
                    for page in by_account[row["account"]]:
                        if exhibit == "P-1R":
                            match = reserve_match(page, row, header, value, edition)
                        elif exhibit == "P-1" and row["cost_type"] not in {"A", "B", "C", "", "None"}:
                            match = cost_row_match(page, row, header, value, edition)
                        else:
                            match = match_cell(page, row, header, value, exhibit, edition)
                        if match is not None:
                            matches.append(match)
                    if len(matches) > 1:
                        # Defense-wide programs are repeated by agency later in R-1.
                        # Prefer the account-wide table, never an arbitrary first numeral.
                        account_table = [m for m in matches if m.get("page", {}).get("text", "").splitlines()[1:2] == ["Defense-Wide"]]
                        if len(account_table) == 1:
                            matches = account_table
                    if len(matches) != 1:
                        missing += 1
                        unmatched.append({"fact_id": fact, "edition": edition, "exhibit": exhibit, "cell": cell, "header": header, "value": value, **row, "matches": len(matches)})
                        continue
                    match = matches[0]
                    if match.get("blank_zero"):
                        blanks += 1
                        continue
                    word, page = match["word"], match["page"]
                    try:
                        geometry = glyphs.geometry(page, word)
                    except ValueError as error:
                        # A matching numeral alone is insufficient: retain the
                        # official book and workbook, but never guess its box.
                        missing += 1
                        unmatched.append({"fact_id": fact, "edition": edition, "exhibit": exhibit,
                                          "cell": cell, "header": header, "value": value, **row,
                                          "matches": 1, "reason": str(error)})
                        continue
                    part = {"amount_thousands": value, "program": row["title"], "line": row["line"],
                            "workbook_cell": cell, "workbook_sha256": workbook_sha, "workbook_url": citation["official_url"],
                            "workbook_sheet": citation["sheet"], "column_label": " ".join(header.split()),
                            "pdf_column_label": match["pdf_column_label"], "edition": edition, "exhibit": exhibit,
                            "sha256": record["sha256"], "hosted_pdf_url": f"/pdfs/{record['sha256']}.pdf",
                            "page_number": page["page_number"], **geometry, "resolution": "unique",
                            "amount_text": word["text"], "units": "USD thousands", "official_url": record["url"]}
                    if exhibit == "P-1":
                        part["row_label"] = str(sheet[f"L{address.row}"].value)
                    elif exhibit == "P-1R":
                        part["row_label"] = "National Guard" if row["cost_type"] == "T" else "Reserve"
                    if "identity_page_number" in match:
                        part["identity_page_number"] = match["identity_page_number"]
                    parts.append(part)
                expected = float(citation["amount_thousands"])
                if sum(values) != Decimal(str(expected)):
                    raise ValueError(f"Canonical workbook amount differs: {fact}")
                matched = sum(p["amount_thousands"] for p in parts)
                receipts[fact] = {"amount_thousands": expected, "matched_amount_thousands": matched,
                                  "complete": missing == 0 and matched == expected, "blank_zero_count": blanks,
                                  "unmatched_count": missing, "parts": parts, "source_documents": [document]}
        finally:
            workbook.close()
            glyphs.close()
        audit = {"edition": edition, "exhibit": exhibit, "pages": len(pages), "facts": len(entries),
                 "complete": sum(r["complete"] for r in receipts.values()) - complete_before}
        book_audits.append(audit)
        print(json.dumps(audit), flush=True)
    leaf_count = len(receipts)
    attach_additive_receipts(citations, receipts)
    out = site_dir / "json/budget-pdf-receipts/v2"
    out.mkdir(parents=True, exist_ok=True)
    shards = defaultdict(dict)
    for fact, receipt in receipts.items():
        shards[fact[:3]][fact] = receipt
    for i in range(4096):
        (out / f"{i:03x}.json").write_text(json.dumps(shards[f"{i:03x}"], separators=(",", ":"), sort_keys=True) + "\n")
    history_path = site_dir / "json/f15_funding_history.json"
    defaults = []
    if history_path.exists():
        history = json.loads(history_path.read_text())
        defaults = [c for p in history["points"] if p["id"] in history["default_point_ids"] for c in p["program_cells"]]
    report = {"schema_version": 2, "full_corpus": editions is None,
              "citation_sha256": hashlib.sha256((site_dir / "json/citations.json").read_bytes()).hexdigest(),
              "source_count": len(book_audits), "leaf_count": leaf_count, "receipt_count": len(receipts),
              "complete_receipts": sum(r["complete"] for r in receipts.values()), "books": book_audits,
              "default_cells": len(defaults), "default_complete": sum(receipts.get(c["fact_id"], {}).get("complete", False) for c in defaults),
              "unmatched_cells": unmatched, "shard_prefix_length": 3,
              "highlight_geometry": "tight PDFium glyph bounds verified by exact text and first-character origin"}
    (site_dir / "json/budget_pdf_receipts_audit.json").write_text(json.dumps(report, indent=2) + "\n")
    return report
