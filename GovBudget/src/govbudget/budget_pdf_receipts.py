"""Reconcile F-15 workbook cells to the same-edition DoD P-1/R-1 PDF.

Only exact matches of account, line, program, fiscal column, and amount become
highlights. Parentheses around gross procurement amounts are not negatives;
the explicit minus sign on Less: Advance Procurement is retained. Printed
blanks are recorded separately for workbook zeros and never get a made-up box.
Canonical workbook citations and facts are not modified by this sidecar export.
"""
from __future__ import annotations

import hashlib
import json
import re
import shutil
from collections import Counter, defaultdict
from ctypes import c_double
from statistics import median
from pathlib import Path

import openpyxl
import pdfplumber
import pypdfium2 as pdfium

NUMBER = re.compile(r"^\(?-?\d[\d,]*(?:\.\d+)?\)?$")


def amount(text: str) -> float | None:
    if not NUMBER.fullmatch(text):
        return None
    return float(text.replace(",", "").replace("(", "").replace(")", ""))


def viewport_geometry(page: dict, word: dict) -> dict:
    """Translate pdfplumber's absolute box into PDF.js's visible crop frame.

    Some Comptroller PDFs place their entire MediaBox/CropBox at a negative
    origin. Width alone is not enough: raw extracted x/top values otherwise
    draw the overlay above and left of the printed amount. Keep matching in
    the source coordinate frame and normalize exactly once, at export.
    """
    if page["rotation"] != 0:
        raise ValueError("Rotated PDF pages require a reviewed coordinate transform")
    left, top, right, bottom = page["cropbox"]
    width, height = right - left, bottom - top
    geometry = {
        "page_width": width,
        "page_height": height,
        "x0": word["x0"] - left,
        "x1": word["x1"] - left,
        "top_pt": word["top"] - top,
        "bottom_pt": word["bottom"] - top,
    }
    if not (0 <= geometry["x0"] < geometry["x1"] <= width and 0 <= geometry["top_pt"] < geometry["bottom_pt"] <= height):
        raise ValueError("Amount highlight is outside the PDF crop frame")
    return geometry


def select_glyph_geometry(page: dict, word: dict, native_cropbox: tuple, candidates: list[dict]) -> dict:
    """Bind tight glyph outlines to the already verified word's text origin.

    pdfminer's em box depends on the font's descent and can miss the tops of
    digits. PDFium exposes actual glyph bounds. First-character origins bind
    the two extractors to the same exact row/column, even when an adjacent
    advance-procurement allocation repeats the same amount.
    """
    em = viewport_geometry(page, word)
    chars = word.get("chars") or []
    if not chars or "".join(c["text"] for c in chars) != word["text"]:
        raise ValueError("Exact source characters are required for glyph verification")
    media_left, _, _, media_bottom = page["mediabox"]
    crop_left, crop_top, _, _ = page["cropbox"]
    source_x = chars[0]["matrix"][4] + media_left - crop_left
    source_y = media_bottom - chars[0]["matrix"][5] - crop_top
    left, bottom, right, top = native_cropbox
    if abs(right - left - em["page_width"]) > .05 or abs(top - bottom - em["page_height"]) > .05:
        raise ValueError("PDF extractors disagree about the visible page extent")
    matches = [candidate for candidate in candidates if candidate["text"] == word["text"]
               and abs(candidate["origin"][0] - left - source_x) < .05
               and abs(top - candidate["origin"][1] - source_y) < .05]
    if len(matches) != 1:
        raise ValueError(f"Expected one exact glyph origin for {word['text']}; got {len(matches)}")
    x0, y0, x1, y1 = matches[0]["box"]
    geometry = {"page_width": em["page_width"], "page_height": em["page_height"],
                "x0": x0 - left, "x1": x1 - left, "top_pt": top - y1, "bottom_pt": top - y0}
    if not (0 <= geometry["x0"] < geometry["x1"] <= geometry["page_width"] and
            0 <= geometry["top_pt"] < geometry["bottom_pt"] <= geometry["page_height"]):
        raise ValueError("Tight amount glyphs are outside the PDF crop frame")
    return geometry


def tight_glyph_geometry(pdf_path: Path, page: dict, word: dict) -> dict:
    """Read source font glyph outlines without changing or rasterizing the PDF."""
    candidates = []
    with pdfium.PdfDocument(pdf_path) as document:
        rendered_page = document[page["page_number"] - 1]
        try:
            textpage = rendered_page.get_textpage()
            try:
                search = textpage.search(word["text"], match_case=True)
                try:
                    while found := search.get_next():
                        index, count = found
                        x, y = c_double(), c_double()
                        if not pdfium.raw.FPDFText_GetCharOrigin(textpage, index, x, y):
                            raise ValueError("Could not read the source glyph baseline")
                        boxes = [textpage.get_charbox(i, loose=False) for i in range(index, index + count)]
                        candidates.append({"text": textpage.get_text_range(index, count), "origin": (x.value, y.value),
                                           "box": (min(b[0] for b in boxes), min(b[1] for b in boxes), max(b[2] for b in boxes), max(b[3] for b in boxes))})
                finally:
                    search.close()
                return select_glyph_geometry(page, word, rendered_page.get_bbox(), candidates)
            finally:
                textpage.close()
        finally:
            rendered_page.close()


def normalize_header(text: str) -> tuple[str, ...]:
    """Compare header words, preserving year and every scenario qualifier."""
    words = re.findall(r"[a-z0-9]+", text.lower())
    ignore = {"amount", "quantity", "qty", "cost"}
    aliases = {"actuals": "actual", "requests": "request", "req": "request"}
    return tuple(sorted(aliases.get(w, w) for w in words if w not in ignore))


def close_line(words: list[dict], y: float) -> list[dict]:
    return sorted((w for w in words if abs(w["top"] - y) < 2), key=lambda w: w["x0"])


def pdf_header_key(header: str, *, edition: int | None, exhibit: str) -> tuple[str, ...]:
    if edition == 2026 and exhibit == "P-1R":
        if normalize_header(header) == normalize_header("FY 2026 Request Amount"):
            return normalize_header("FY 2026 Disc Request")
        if normalize_header(header) == normalize_header("FY 2026 Reconciliation Amount"):
            return normalize_header("FY 2026 Reconcil Request")
    # Reviewed PB2025 R-1 wording difference only. Both originals carry the
    # same annualized-CR footnote (Public Law 118-35); FY2024 has no full-year
    # appropriation in this edition. Do not alias these words in other books.
    if edition == 2025 and exhibit == "R-1" and normalize_header(header) == normalize_header("FY 2024 PB Request with CR Amounts*"):
        return normalize_header("FY 2024 PB Request with CR Adjustments*")
    return normalize_header(header)


def page_columns(page: dict, exhibit: str) -> list[dict]:
    """Find labeled money columns from their headers and repeated right edges."""
    if exhibit in page.get("_columns", {}):
        return page["_columns"][exhibit]
    words = page["words"]
    no = next((w for w in words if w["text"] == "No" and w["x0"] < 100), None)
    if no is None and exhibit == "P-1R":
        no = next((w for w in words if w["text"] == "Line" and w["x0"] < 100), None)
    dollars = next((w for w in words if w["text"] == "(Dollars"), None)
    if no is None or dollars is None:
        return []
    header = [w for w in words if dollars["bottom"] < w["top"] <= no["top"] + 2]
    pairs = []
    for word in header:
        if word["text"] != "FY":
            continue
        years = [w for w in header if re.fullmatch(r"20\d\d", w["text"]) and abs(w["top"] - word["top"]) < 1 and 0 < w["x0"] - word["x1"] < 8]
        if len(years) == 1:
            pairs.append((word, years[0]))
    pairs.sort(key=lambda p: p[0]["x0"])
    if not pairs:
        return []
    first_y = min(w["top"] for w, _ in pairs)
    header = [w for w in header if w["top"] >= first_y - 1]
    centers = [(w["x0"] + y["x1"]) / 2 for w, y in pairs]
    if exhibit in {"P-1", "P-1R"}:
        costs = sorted((w for w in header if w["text"].rstrip("*") == "Cost"), key=lambda w: w["x0"])
        if len(costs) != len(pairs):
            return []
        anchors = [w["x1"] for w in costs]
    else:
        anchors = [c + 25 for c in centers]
    body = [w for w in words if w["top"] > no["bottom"] and w["top"] < page["height"] - 35 and amount(w["text"]) is not None]
    edges = []
    for anchor in anchors:
        candidates = [w for w in body if abs(w["x1"] - anchor) < 30 and w["x0"] > centers[0] - 40]
        counts = Counter(round(w["x1"], 0) for w in candidates)
        edges.append(counts.most_common(1)[0][0] if counts else None)
    offsets = [edge - anchor for edge, anchor in zip(edges, anchors) if edge is not None]
    if not offsets:
        return []
    # Entire OCO/supplemental columns can be blank. Their printed header still
    # establishes the column; inherit only the page's measured alignment.
    edges = [edge if edge is not None else anchor + median(offsets) for edge, anchor in zip(edges, anchors)]
    result = []
    for i, edge in enumerate(edges):
        left = edges[i - 1] + 2 if i else pairs[0][0]["x0"] - 30
        # Headings are centered over each full quantity/cost group. Use the
        # money edge rather than the FY word center so long labels stay whole.
        label_words = [w for w in header if left <= (w["x0"] + w["x1"]) / 2 < edge + 6 and w["text"] not in {"Qty", "Cost", "Cost*", "Quantity", "S", "e", "c", "Sec", "Se", "Act"}]
        label = " ".join(w["text"] for w in sorted(label_words, key=lambda w: (round(w["top"]), w["x0"])))
        result.append({"header": label, "normalized": normalize_header(label), "edge": edge, "body_top": no["bottom"]})
    page.setdefault("_columns", {})[exhibit] = result
    return result


def program_lines(page: dict) -> list[dict]:
    if "_lines" in page:
        return page["_lines"]
    words = page["words"]
    no = next((w for w in words if w["text"] == "No" and w["x0"] < 100), None)
    if no is None:
        return []
    page["_lines"] = [w for w in words if w["text"].isdigit() and abs(w["x1"] - no["x1"]) < 15 and w["top"] > no["bottom"] and w["top"] < page["height"] - 35]
    return page["_lines"]


def title_identity(text: str) -> str:
    """Ignore typography, while requiring the entire program title."""
    return re.sub(r"[\W_]+", "", text.casefold())


def inherited_activity(page: dict, y: float) -> str | None:
    """Read a printed activity, crossing only verified account continuations."""
    while page is not None:
        headings = [w for w in page["words"] if w["text"] == "Budget" and 150 < w["top"] < y]
        if headings:
            text = " ".join(w["text"] for w in close_line(page["words"], max(headings, key=lambda w: w["top"])["top"]))
            match = re.match(r"Budget Activity (\d+):", text)
            return match[1] if match else None
        page = page.get("_previous")
        y = float("inf")
    return None


def wrapped_title(words: list[dict], y: float, left: float, right: float, next_y: float, previous_y: float = 0, *, before: bool = True) -> str:
    """Read only contiguous title lines, stopping before the next budget row."""
    candidates = sorted((w for w in words if left < w["x0"] and w["x1"] < right and max(previous_y + 2, y - (40 if before else 2)) <= w["top"] < min(next_y - 2, y + 40)), key=lambda w: (round(w["top"]), w["x0"]))
    groups = []
    for word in candidates:
        if not groups or abs(word["top"] - groups[-1][0]["top"]) > 1.5:
            groups.append([])
        groups[-1].append(word)
    anchor = next((i for i, group in enumerate(groups) if abs(group[0]["top"] - y) < 2), None)
    if anchor is None:
        return ""
    column_left = min(w["x0"] for w in groups[anchor])
    start = end = anchor
    def continuation(index, previous):
        group = groups[index]
        full_line = close_line(words, group[0]["top"])
        return (abs(group[0]["top"] - groups[previous][0]["top"]) <= max(11.3, (group[0]["bottom"] - group[0]["top"]) * 1.4)
                and not any(w["x0"] < column_left - 1 for w in full_line)
                and group[0]["text"] not in {"Less:", "Advance", "Subsequent", "Completion"})
    while start > 0 and continuation(start - 1, start):
        start -= 1
    while end + 1 < len(groups) and continuation(end + 1, end):
        end += 1
    return " ".join(w["text"] for group in groups[start:end + 1] for w in sorted(group, key=lambda w: w["x0"]))


def match_cell(page: dict, row: dict, header: str, value: float, exhibit: str, edition: int | None = None) -> dict | None:
    """Return one exact printed amount, or an independently located blank zero."""
    text = page["text"]
    account = row["account"]
    if not re.search(rf"\b{re.escape(account)}\b", text):
        return None
    columns = [col for col in page_columns(page, exhibit) if col["normalized"] == pdf_header_key(header, edition=edition, exhibit=exhibit)]
    if len(columns) != 1:
        return None
    col = columns[0]
    lines = program_lines(page)
    candidates = [w for w in lines if int(w["text"]) == int(row["line"])]
    if len(candidates) != 1:
        return None
    line = candidates[0]
    words = page["words"]
    same = close_line(words, line["top"])
    line_text = " ".join(w["text"] for w in same)
    next_y = min((w["top"] for w in lines if w["top"] > line["top"] + 2), default=page["height"] - 35)
    previous_y = max((w["top"] for w in lines if w["top"] < line["top"] - 2), default=0)
    # P-1 omits the BLI code in the PDF; the exact title and account+line
    # identify it. R-1 prints the PE and activity, which must also agree.
    if exhibit == "R-1":
        codes = {row["code"]}
        # Public classified-program aggregate uses a nine-digit sentinel in
        # the PDF and ten digits in the workbook; never truncate real PEs.
        if row["code"] == "9999999999" and row["line"] == "999" and row["title"] == "Classified Programs":
            codes.add("999999999")
        code = next((w for w in same if w["text"] in codes), None)
        act = next((w for w in words if w["text"] == "Act"), None)
        if code is None or act is None:
            return None
        # PB2024 prints the two digits slightly left of the Act heading.
        # Keep the check inside that column, never anywhere in the row title.
        activities = [w["text"] for w in same if re.fullmatch(r"\d{2}", w["text"])
                      and act["x0"] - 12 <= w["x0"] and w["x1"] <= act["x1"] + 4]
        if activities != [row["activity"]]:
            return None
        # Editions differ: the code is printed beside either the first or
        # final title line. The exact PE/line/activity pins both layouts.
        titles = [wrapped_title(words, line["top"], code["x1"], act["x0"], next_y, previous_y, before=before)
                  for before in (False, True)]
        if not any(title_identity(title) == title_identity(row["title"]) for title in titles):
            return None
    else:
        ident = next((w for w in words if w["text"] == "Ident"), None)
        if ident is None:
            return None
        title = wrapped_title(words, line["top"], line["x1"], ident["x0"], next_y, previous_y)
        if row["line"] == "999" and row["title"] == "Classified Programs":
            if inherited_activity(page, line["top"]) != row["activity"]:
                return None
            title = " ".join(w["text"] for w in same if line["x1"] < w["x0"] and w["x1"] < ident["x0"])
        if title_identity(title) != title_identity(row["title"]):
            return None
        ident_tokens = [w["text"] for w in same if abs(w["x0"] - ident["x0"]) < 16]
        if row["cost_type"] == "C" and "A" in ident_tokens:
            return None
    y = line["top"]
    next_y = min((w["top"] for w in lines if w["top"] > y + 2), default=page["height"] - 35)
    if exhibit == "P-1" and row["cost_type"] in {"B", "C"}:
        target = "Less:" if row["cost_type"] == "B" else "Advance"
        candidates = []
        for w in words:
            if w["text"] != target:
                continue
            around = " ".join(v["text"] for v in close_line(words, w["top"]))
            if row["cost_type"] == "B" and y < w["top"] < next_y and "(PY)" in around:
                candidates.append(w)
            elif row["cost_type"] == "C" and "(CY)" in around:
                if y < w["top"] < next_y:
                    candidates.append(w)
                elif 0 < y - w["top"] < 30:
                    # PB2024 puts the CY subtotal immediately BEFORE its line.
                    # The prior line must be the same named program, not an
                    # unrelated preceding aircraft's advance-procurement sum.
                    previous = max((v for v in lines if v["top"] < w["top"]), key=lambda v: v["top"], default=None)
                    if previous:
                        prior_title = " ".join(v["text"] for v in close_line(words, previous["top"]) if previous["x1"] < v["x0"] and v["x1"] < ident["x0"])
                        if row["title"].lower() == prior_title.lower():
                            candidates.append(w)
        if len(candidates) != 1:
            return None
        y = candidates[0]["top"]
    elif exhibit == "P-1" and row["cost_type"] != "A":
        return None
    amounts = [w for w in close_line(words, y) if abs(w["x1"] - col["edge"]) <= 6 and amount(w["text"]) is not None]
    if not amounts and value == 0:
        return {"blank_zero": True, "pdf_column_label": col["header"], "page": page}
    if len(amounts) != 1 or amount(amounts[0]["text"]) != value:
        return None
    return {"word": amounts[0], "page": page, "pdf_column_label": col["header"]}


def load_sources(manifest: Path, cache_dir: Path, site_dir: Path) -> dict:
    """Re-extract only reviewed pages from checksum-pinned original PDFs."""
    sources = {}
    for record in json.loads(manifest.read_text()):
        pdf = site_dir / "pdfs" / f"{record['sha256']}.pdf"
        if not pdf.exists():
            cached = cache_dir / f"{record['edition']}_{record['exhibit'].lower().replace('-', '')}.pdf"
            if cached.exists():
                if hashlib.sha256(cached.read_bytes()).hexdigest() != record["sha256"]:
                    raise ValueError(f"Source hash mismatch: {cached}")
                pdf.parent.mkdir(parents=True, exist_ok=True)
                shutil.copyfile(cached, pdf)
            else:
                import httpx
                response = httpx.get(record["url"], follow_redirects=True, timeout=90)
                response.raise_for_status()
                if hashlib.sha256(response.content).hexdigest() != record["sha256"]:
                    raise ValueError(f"Source hash changed; review needed: {record['url']}")
                pdf.parent.mkdir(parents=True, exist_ok=True)
                pdf.write_bytes(response.content)
        if hashlib.sha256(pdf.read_bytes()).hexdigest() != record["sha256"]:
            raise ValueError(f"Source hash mismatch: {pdf}")
        pages = []
        with pdfplumber.open(pdf) as document:
            for number in record["pages"]:
                page = document.pages[number - 1]
                if page.rotation != 0:
                    raise ValueError(f"Rotated PDF source needs review: {pdf} page {number}")
                text = page.extract_text() or ""
                if f"FY {record['edition']} President's Budget" not in text or f"Exhibit {record['exhibit']}" not in text or "(Dollars in Thousands)" not in text:
                    raise ValueError(f"PDF identity mismatch: {pdf} page {number}")
                pages.append({"page_number": number, "width": page.width, "height": page.height, "mediabox": list(page.mediabox), "cropbox": list(page.cropbox), "rotation": page.rotation, "text": text, "words": page.extract_words(return_chars=True)})
        sources[(record["edition"], record["exhibit"])] = {**record, "pages": pages}
    return sources


def combine_receipts(receipts: list[dict], expected: float) -> dict:
    parts = [part for r in receipts for part in r["parts"]]
    matched = sum(p["amount_thousands"] for p in parts)
    unmatched = sum(r["unmatched_count"] for r in receipts)
    complete = all(r["complete"] for r in receipts) and matched == expected
    return {"amount_thousands": expected, "matched_amount_thousands": matched, "complete": complete, "blank_zero_count": sum(r["blank_zero_count"] for r in receipts), "unmatched_count": unmatched, "parts": parts}


def export_budget_pdf_receipts(*, site_dir: Path, manifest: Path, cache_dir: Path) -> dict:
    history = json.loads((site_dir / "json/f15_funding_history.json").read_text())
    citations = json.loads((site_dir / "json/citations.json").read_text())
    sources = load_sources(manifest, cache_dir, site_dir)
    workbooks = {}
    receipts = {}
    unmatched = []
    for point in history["points"]:
        for component in point["components"]:
            fact = component["fact_id"]
            if fact in receipts:
                continue
            citation = citations[fact]
            sha = citation["sha256"]
            if sha not in workbooks:
                path = site_dir / "workbooks" / f"{sha}.xlsx"
                if hashlib.sha256(path.read_bytes()).hexdigest() != sha:
                    raise ValueError(f"Workbook hash mismatch: {path}")
                workbooks[sha] = openpyxl.load_workbook(path, data_only=True)
            sheet = workbooks[sha][component["sheet"]]
            source = sources[(point["edition"], component["exhibit"])]
            parts = []
            blanks = 0
            missing = 0
            workbook_total = 0
            for cell in component["cells"].split(","):
                address = sheet[cell]
                row = address.row
                value = float(address.value or 0)
                workbook_total += value
                metadata = {"account": str(sheet[f"A{row}"].value), "activity": str(sheet[f"D{row}"].value), "line": str(sheet[f"F{row}"].value).strip(), "code": str(sheet[f"I{row}" if component["exhibit"] == "P-1" else f"G{row}"].value), "title": str(sheet[f"J{row}" if component["exhibit"] == "P-1" else f"H{row}"].value), "cost_type": str(sheet[f"K{row}"].value)}
                if metadata["account"] != component["account"] or metadata["code"] != component["program_id"].split(":")[-1] or metadata["activity"] != component["budget_activity"] or metadata["title"].lower() != component["title"].lower():
                    raise ValueError(f"Workbook program identity mismatch: {fact}/{cell}")
                header = str(sheet.cell(2, address.column).value)
                matches = [match for page in source["pages"] if (match := match_cell(page, metadata, header, value, component["exhibit"], point["edition"])) is not None]
                if len(matches) != 1:
                    missing += 1
                    unmatched.append({"fact_id": fact, "edition": point["edition"], "exhibit": component["exhibit"], "cell": cell, "value": value, "header": header, **metadata, "matches": len(matches)})
                    continue
                match = matches[0]
                if match.get("blank_zero"):
                    blanks += 1
                    continue
                word, page = match["word"], match["page"]
                glyph_geometry = tight_glyph_geometry(site_dir / "pdfs" / f"{source['sha256']}.pdf", page, word)
                parts.append({"amount_thousands": value, "program": metadata["title"], "line": metadata["line"], "workbook_cell": cell, "column_label": " ".join(header.split()), "pdf_column_label": match["pdf_column_label"], "edition": point["edition"], "exhibit": component["exhibit"], "sha256": source["sha256"], "hosted_pdf_url": f"/pdfs/{source['sha256']}.pdf", "page_number": page["page_number"], **glyph_geometry, "resolution": "unique", "amount_text": word["text"], "units": "USD thousands", "official_url": source["url"]})
                if component["exhibit"] == "P-1":
                    parts[-1]["row_label"] = str(sheet[f"L{row}"].value)
            expected = component["amount_thousands"]
            if workbook_total != expected:
                raise ValueError(f"Canonical workbook amount mismatch: {fact}")
            matched = sum(p["amount_thousands"] for p in parts)
            receipts[fact] = {"amount_thousands": expected, "matched_amount_thousands": matched, "complete": missing == 0 and matched == expected, "blank_zero_count": blanks, "unmatched_count": missing, "parts": parts}
    leaf_count = len(receipts)
    for point in history["points"]:
        for cell in point["program_cells"]:
            if cell["fact_id"] not in receipts:
                receipts[cell["fact_id"]] = combine_receipts([receipts[fid] for fid in cell["input_fact_ids"]], cell["amount_thousands"])
        receipts[point["fact_id"]] = combine_receipts([receipts[c["fact_id"]] for c in point["components"]], point["amount_thousands"])
    cumulative = history["cumulative"]
    inputs = json.loads(citations[cumulative["fact_id"]]["inputs"])
    receipts[cumulative["fact_id"]] = combine_receipts([receipts[fid] for fid in inputs], cumulative["amount_thousands"])
    shards = defaultdict(dict)
    for fact, receipt in receipts.items():
        shards[fact[:2]][fact] = receipt
    out = site_dir / "json/budget-pdf-receipts"
    out.mkdir(parents=True, exist_ok=True)
    for i in range(256):
        (out / f"{i:02x}.json").write_text(json.dumps(shards[f"{i:02x}"], separators=(",", ":"), sort_keys=True) + "\n")
    default_cells = [cell for p in history["points"] if p["id"] in history["default_point_ids"] for cell in p["program_cells"]]
    report = {"source_count": len(sources), "leaf_count": leaf_count, "receipt_count": len(receipts), "default_cells": len(default_cells), "default_complete": sum(receipts[c["fact_id"]]["complete"] for c in default_cells), "complete_receipts": sum(r["complete"] for r in receipts.values()), "unmatched_cells": unmatched, "highlight_geometry": "tight PDFium glyph bounds verified by exact text and first-character origin"}
    report["coordinate_frames"] = [
        {"edition": source["edition"], "exhibit": source["exhibit"], "sha256": source["sha256"], "pages": [
            {"page_number": page["page_number"], "cropbox": page["cropbox"], "rotation": page["rotation"]}
            for page in source["pages"]
        ]}
        for source in sources.values()
    ]
    (site_dir / "json/budget_pdf_receipts_audit.json").write_text(json.dumps(report, indent=2) + "\n")
    return report
