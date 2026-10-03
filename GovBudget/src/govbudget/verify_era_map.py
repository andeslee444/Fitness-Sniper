"""Families piece 1 release gate: verify-era-map (spec §8, V4).

CLI: `govbudget verify-era-map [--legs abc]`. With no arguments it checks
every leg against config.SITE_DIR, config.DUCKDB_PATH and the committed seed
(dbt/seeds/p1_era_code_decisions.csv), prints one line per leg and then
exactly one final line, `verify-era-map: PASS` or `verify-era-map: FAIL`,
and exits 1 on FAIL. It is built to fail: a missing table, workbook or seed
is a FAIL, never a skip.

  Leg (a) — workbooks. Re-read the seven PB2017–PB2023 P-1 workbooks the site
    ships (data/site/workbooks/<sha>.xlsx; the shas are pinned below and the
    file bytes must hash to them) with the loader's own parse
    (p1_loader.parse_p1_rollup). The era key set must be exactly the 6,927
    keys of the corpus (EXPECTED_ERA_KEYS) and equal p1_era_line_map's; for
    every key the account, organization, budget activity, printed code
    (line_item_code) and filed title must equal the map's, the map must cite
    the pinned workbook, and its source_cells must be the code column's cells
    of exactly the rows the parse summed. Then, independently of the parser,
    every cited cell must print the key's code, its row's 'Line Item Title'
    must equal filed_title, and its row's 'Line Number' must be the key's
    line.

  Leg (b) — decisions. No era key is undecided; every decided key's
    keys_sha_ok (the dbt-side drift hash) is true; the map was built from the
    seed on disk (each decided key's decision_id exists there with the same
    decision, ruling, pins and successor); no seed row binds zero keys; and,
    recomputed in Python with era_map.keys_sha256 over the keys each seed row
    binds, every seed row's keys_sha256 and n_keys still hold.

  Leg (c) — F-15. The map's F-15 rows — every key printing one of the
    ERA_PROGRAM_CODES codes or titled like F-15 — are exactly ERA_MEMBERS
    joined to ERA_PROGRAM_CODES (31 rows: edition, era key, budget activity,
    filed title, printed code), and each carries its printed code as
    program_key.

  Leg (d) — fact ids. Every row of the exported budget_lines_decade.parquet
    recomputes its fact_id with export_site.fact_id_workbook from its own
    columns, no fact_id repeats, and every PB2017–PB2023 P-1 row keeps its
    era key as pe_bli (spec §6.1: the lake identity is never re-keyed). At
    least one such row must exist, so the check cannot pass vacuously.

  Leg (e) — published map. The exported data/p1_era_line_map.parquet equals
    the warehouse's p1_era_line_map: the same column names in the same order,
    the same row multiset, and not empty. A warehouse whose map moved after
    the export fails here until the site is re-exported.

  Leg (f) — F-15 history pin. data/site/json/f15_funding_history.json hashes
    to the sha256 committed at S0 in tests/fixtures/f15/history.sha256 (its
    first token, so `shasum -a 256` output reads); only the S5 correction
    re-pins it.
"""
from __future__ import annotations

import csv
import hashlib
import re
from collections import Counter, defaultdict
from collections.abc import Mapping
from pathlib import Path

from govbudget import config

LEGS = ("a", "b", "c", "d", "e", "f")
LEG_NAMES = {
    "a": "workbooks", "b": "decisions", "c": "f15",
    "d": "fact ids", "e": "published map", "f": "f15 history pin",
}

# The seven era P-1 display workbooks (jbook_documents rel_path
# fy<N>/dod/p1_display.xlsx), by edition — measured 2026-10-02.
ERA_P1_WORKBOOK_SHA256: dict[int, str] = {
    2017: "fa2d6711dc766be5cdf6148ed46b158e6420fe9653613900c43add4a28b26b31",
    2018: "863ea56b3d32294c4c61f12ec6e99622d74d12cfcce6a9c7c7c1e17e439cb318",
    2019: "af88e62619ba910915a5f5b68abb44652d75359cd1f8a46b07df609b8925d501",
    2020: "f019309be2f85fd523b00c5ec419c845747d867cd5d3af10908e979b171f584f",
    2021: "73aadc1b5bee700a4d4ea3f95b7a6be65dde58ae0bff6fc4fb33bef6b7362942",
    2022: "02a59dcc426aebf241c244f26698c7dd8ed1d2594507d2ea9b514857f1cee5fd",
    2023: "710db28120491ddcbed14a5b5e8bb71d183a7fe9d032002dbca40b2e7412c504",
}
# Era keys per edition (spec §3): 6,927 in all.
EXPECTED_ERA_KEYS: dict[int, int] = {
    2017: 969, 2018: 1029, 2019: 989, 2020: 975, 2021: 1005, 2022: 993, 2023: 967,
}
F15_ERA_ROWS = 31

DEFAULT_SEED_PATH = config.ROOT / "dbt" / "seeds" / "p1_era_code_decisions.csv"

MAP_COLUMNS = (
    "edition", "account", "organization", "budget_activity", "era_key",
    "line_item_code", "filed_title", "program_key", "program_account", "program_org",
    "decision", "decision_id", "ruling", "keys_sha_ok", "successor_code",
    "source_document_sha256", "source_cells",
)

_CELL_ROW_RE = re.compile(r"^([A-Z]+)([0-9]+)$")
_PRINTED_HEADERS = ("Line Number", "Line Item", "Line Item Title")
_MAX_LISTED = 25


def _leg(ok: bool, summary: str, failures: list[str]) -> dict:
    return {"ok": ok, "detail": {"summary": summary, "failures": failures}}


def _read_map(duckdb_path: Path) -> list[dict]:
    import duckdb

    con = duckdb.connect(str(duckdb_path), read_only=True)
    try:
        cur = con.execute(
            f"select {', '.join(MAP_COLUMNS)} from p1_era_line_map order by edition, era_key")
        return [dict(zip(MAP_COLUMNS, row)) for row in cur.fetchall()]
    finally:
        con.close()


def _sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _printed_cells(xlsx_path: Path) -> tuple[dict[int, tuple], str]:
    """The era P-1 sheet's printed (Line Number, Line Item, Line Item Title) per
    worksheet row, read straight from openpyxl (not through the loader), and
    the column letter of 'Line Item'."""
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    from govbudget.jbooks.rollup_loader import norm_header

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        ws = wb["Exhibit P-1"] if "Exhibit P-1" in wb.sheetnames else wb[wb.sheetnames[0]]
        header_row, cols = None, {}
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True), start=1):
            found = {norm_header(v): j for j, v in enumerate(row) if v is not None}
            if all(h in found for h in _PRINTED_HEADERS):
                header_row, cols = i, found
                break
        if header_row is None:
            raise ValueError(
                f"{xlsx_path.name}: no header row with {_PRINTED_HEADERS} in the first 20 rows")
        idx = [cols[h] for h in _PRINTED_HEADERS]
        printed = {}
        for r, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True),
                                start=header_row + 1):
            printed[r] = tuple(row[j] if j < len(row) else None for j in idx)
        return printed, get_column_letter(cols["Line Item"] + 1)
    finally:
        wb.close()


def _workbook_keys(parse) -> dict[str, dict]:
    """{era_key: {"identity": {(account, org, ba, code, title)}, "rows": {row}}}"""
    from govbudget.jbooks.era_keys import is_era_procurement_key

    keys: dict[str, dict] = {}
    for row in parse.rows:
        if not is_era_procurement_key(row.pe_bli):
            continue  # the classified aggregate (S1b) is not an era key
        k = keys.setdefault(row.pe_bli, {"identity": set(), "rows": set()})
        k["identity"].add((row.account, row.organization, row.budget_activity,
                           row.line_item_code, row.title))
        for cell in row.source_cells:
            m = _CELL_ROW_RE.match(cell)
            if m is None:
                raise ValueError(f"{row.pe_bli}: unparseable source cell {cell!r}")
            k["rows"].add(int(m.group(2)))
    return keys


def leg_a_workbooks(
    *, site_dir: Path, map_rows: list[dict],
    workbooks: Mapping[int, str], expected_keys: Mapping[int, int],
) -> dict:
    """Leg (a): the shipped era P-1 workbooks vs p1_era_line_map."""
    from govbudget.jbooks.era_keys import era_key_anchor
    from govbudget.jbooks.p1_loader import parse_p1_rollup

    failures: list[str] = []
    # Task 20 review carry-over (Task 14 fix round): by_edition below keys
    # map_rows by (edition, era_key), so a duplicated row silently collapses
    # — the dbt grain/no_overlap tests catch this, but the release gate
    # should too, rather than quietly checking only the last copy.
    pair_counts = Counter((r["edition"], r["era_key"]) for r in map_rows)
    if len(map_rows) > len(pair_counts):
        edition, era_key = min(pair for pair, n in pair_counts.items() if n > 1)
        failures.append(
            f"p1_era_line_map has {len(map_rows)} rows but only {len(pair_counts)} distinct"
            f" (edition, era_key) pairs — e.g. {edition}/{era_key} appears"
            f" {pair_counts[(edition, era_key)]} times")
    by_edition: dict[int, dict[str, dict]] = defaultdict(dict)
    for r in map_rows:
        by_edition[r["edition"]][r["era_key"]] = r
    for edition in sorted(set(by_edition) - set(workbooks)):
        failures.append(f"{edition}: the map has {len(by_edition[edition])} keys for an"
                        " edition with no pinned era P-1 workbook")
    checked = 0
    for edition, sha in sorted(workbooks.items()):
        path = Path(site_dir) / "workbooks" / f"{sha}.xlsx"
        if not path.is_file():
            failures.append(f"{edition}: workbook missing: {path}")
            continue
        actual_sha = _sha256_file(path)
        if actual_sha != sha:
            failures.append(f"{edition}: {path.name} hashes to {actual_sha}, not its pinned sha")
            continue
        try:
            parse = parse_p1_rollup(path, exhibit="P-1", fiscal_year=edition)
            if not parse.era_line_keying:
                failures.append(f"{edition}: the parse did not use era line keying")
                continue
            wb_keys = _workbook_keys(parse)
            printed, code_col = _printed_cells(path)
        except Exception as e:  # a tripwire (EraKeyConflict) or a bad file is a FAIL
            failures.append(f"{edition}: {type(e).__name__}: {e}")
            continue
        want = expected_keys.get(edition)
        if len(wb_keys) != want:
            failures.append(f"{edition}: the workbook parses to {len(wb_keys)} era keys,"
                            f" the corpus has {want}")
        mapped = by_edition.get(edition, {})
        for key in sorted(set(wb_keys) - set(mapped)):
            failures.append(f"{edition}/{key}: in the workbook, missing from the map")
        for key in sorted(set(mapped) - set(wb_keys)):
            failures.append(f"{edition}/{key}: in the map, not in the workbook")
        for key in sorted(set(wb_keys) & set(mapped)):
            checked += 1
            wk, m = wb_keys[key], mapped[key]
            if len(wk["identity"]) != 1:
                failures.append(f"{edition}/{key}: the workbook prints {len(wk['identity'])}"
                                " identities for one key")
                continue
            account, org, ba, code, title = next(iter(wk["identity"]))
            got = (m["account"], m["organization"], m["budget_activity"],
                   m["line_item_code"], m["filed_title"])
            if got != (account, org, ba, code, title):
                failures.append(f"{edition}/{key}: workbook {(account, org, ba, code, title)!r}"
                                f" != map {got!r}")
            if m["source_document_sha256"] != sha:
                failures.append(f"{edition}/{key}: map cites {m['source_document_sha256']},"
                                f" not the pinned workbook")
            want_cells = ",".join(f"{code_col}{r}" for r in sorted(wk["rows"]))
            if m["source_cells"] != want_cells:
                failures.append(f"{edition}/{key}: map cells {m['source_cells']!r}"
                                f" != workbook rows {want_cells!r}")
            for r in sorted(wk["rows"]):
                line, printed_code, printed_title = printed.get(r, (None, None, None))
                if (str(printed_code).strip() if printed_code is not None else None) != m["line_item_code"]:
                    failures.append(f"{edition}/{key}: {code_col}{r} prints {printed_code!r},"
                                    f" the map says {m['line_item_code']!r}")
                if printed_title != m["filed_title"]:
                    failures.append(f"{edition}/{key}: row {r} title {printed_title!r}"
                                    f" != filed_title {m['filed_title']!r}")
                if (str(line).strip() if line is not None else None) != era_key_anchor(key):
                    failures.append(f"{edition}/{key}: row {r} line number {line!r}"
                                    " is not the key's line")
    total = sum(expected_keys.get(e, 0) for e in workbooks)
    return _leg(not failures,
                f"workbooks={len(workbooks)} keys_checked={checked} expected={total}"
                f" failures={len(failures)}", failures)


def _read_seed(seed_path: Path) -> dict[str, dict]:
    from govbudget.jbooks.era_map import SEED_COLUMNS

    with open(seed_path, newline="") as f:
        reader = csv.DictReader(f)
        if tuple(reader.fieldnames or ()) != tuple(SEED_COLUMNS):
            raise ValueError(f"{seed_path.name}: header {reader.fieldnames} != SEED_COLUMNS")
        rows = {}
        for row in reader:
            clean = {k: (v if v != "" else None) for k, v in row.items()}
            if clean["decision_id"] in rows:
                raise ValueError(f"{seed_path.name}: duplicate decision_id {clean['decision_id']}")
            rows[clean["decision_id"]] = clean
        return rows


def leg_b_decisions(*, map_rows: list[dict], seed_path: Path) -> dict:
    """Leg (b): nothing undecided, nothing stale, and the map is this seed's."""
    from govbudget.jbooks.era_map import keys_sha256

    failures: list[str] = []
    try:
        seed = _read_seed(Path(seed_path))
    except (OSError, ValueError) as e:
        return _leg(False, "seed unreadable", [f"{type(e).__name__}: {e}"])
    undecided = [r for r in map_rows if r["decision"] == "undecided"]
    for r in undecided:
        failures.append(f"{r['edition']}/{r['era_key']} ({r['line_item_code']}): undecided")
    bound: dict[str, list[dict]] = defaultdict(list)
    for r in map_rows:
        if r["decision"] == "undecided":
            continue
        if r["keys_sha_ok"] is not True:
            failures.append(f"{r['edition']}/{r['era_key']}: keys_sha_ok={r['keys_sha_ok']}"
                            f" under {r['decision_id']} (stale decision)")
        bound[r["decision_id"]].append(r)
    for did, rows in sorted(bound.items()):
        s = seed.get(did)
        if s is None:
            failures.append(f"{did}: in the map, not in {Path(seed_path).name} (stale build)")
            continue
        for r in rows:
            got = (r["decision"], r["ruling"], r["program_account"], r["program_org"],
                   r["successor_code"])
            want = (s["decision"], s["ruling"], s["program_account"], s["program_org"],
                    s["successor_code"])
            if got != want:
                failures.append(f"{did}: map {got!r} != seed {want!r} (stale build)")
                break
        sha = keys_sha256((r["edition"], r["era_key"], r["budget_activity"], r["filed_title"])
                          for r in rows)
        if sha != s["keys_sha256"] or str(len(rows)) != (s["n_keys"] or ""):
            failures.append(f"{did}: binds {len(rows)} keys hashing to {sha[:12]}…, the seed"
                            f" reviewed {s['n_keys']} hashing to {(s['keys_sha256'] or '')[:12]}…")
    for did in sorted(set(seed) - set(bound)):
        failures.append(f"{did}: in the seed, binds no era key")
    return _leg(not failures,
                f"keys={len(map_rows)} undecided={len(undecided)} decisions={len(seed)}"
                f" failures={len(failures)}", failures)


def _looks_f15(title: str | None) -> bool:
    t = (title or "").upper()
    return "F-15" in t or "F15" in t


def leg_c_f15(*, map_rows: list[dict]) -> dict:
    """Leg (c): the map's F-15 rows are ERA_MEMBERS x ERA_PROGRAM_CODES."""
    from govbudget.f15_funding_history import ERA_MEMBERS, ERA_PROGRAM_CODES

    failures: list[str] = []
    expected = set()
    for edition, members in ERA_MEMBERS.items():
        for line, ba, title in members:
            code = ERA_PROGRAM_CODES.get(edition, {}).get(line)
            if code is None:
                failures.append(f"{edition}/L{line}: ERA_MEMBERS line has no ERA_PROGRAM_CODES code")
                continue
            expected.add((edition, f"3010F-AF-L{line}", ba, title, code))
    if len(expected) != F15_ERA_ROWS:
        failures.append(f"ERA_MEMBERS x ERA_PROGRAM_CODES has {len(expected)} rows, not {F15_ERA_ROWS}")
    codes = {c for per in ERA_PROGRAM_CODES.values() for c in per.values()}
    f15 = [r for r in map_rows if r["line_item_code"] in codes or _looks_f15(r["filed_title"])]
    actual = {(r["edition"], r["era_key"], r["budget_activity"], r["filed_title"],
               r["line_item_code"]) for r in f15}
    for row in sorted(expected - actual):
        failures.append(f"missing from the map: {row!r}")
    for row in sorted(actual - expected):
        failures.append(f"in the map, not in ERA_MEMBERS x ERA_PROGRAM_CODES: {row!r}")
    for r in f15:
        if r["program_key"] != r["line_item_code"]:
            failures.append(f"{r['edition']}/{r['era_key']}: program_key {r['program_key']!r}"
                            f" != printed code {r['line_item_code']!r} ({r['decision']})")
    return _leg(not failures,
                f"expected={len(expected)} map_rows={len(actual)} failures={len(failures)}",
                failures)


# --- Legs d–f (Task 20): the published export against its own inputs -------

#: The F-15 history pin committed at S0 (Task 4). Its first whitespace-
#: separated token is the sha256 hex digest of
#: data/site/json/f15_funding_history.json (`shasum -a 256` output reads).
#: Resolved from this file, not config.ROOT, so a test can swap it.
F15_HISTORY_PIN = Path(__file__).resolve().parents[2] / "tests" / "fixtures" / "f15" / "history.sha256"

#: The last edition whose P-1 rows carry era keys as pe_bli (spec §3).
_ERA_P1_LAST_EDITION = 2023


def leg_d_fact_ids(*, site_dir: Path) -> dict:
    """Leg (d): budget_lines_decade's fact ids recompute; era rows keep era keys."""
    import duckdb

    from govbudget.export_site import fact_id_workbook
    from govbudget.jbooks.era_keys import is_era_procurement_key

    pq = Path(site_dir) / "data" / "budget_lines_decade.parquet"
    if not pq.is_file():
        return _leg(False, "budget_lines_decade.parquet missing", [f"missing {pq}"])
    con = duckdb.connect()
    try:
        rows = con.execute(
            "select fact_id, document_sha256, exhibit, fiscal_year, account,"
            " organization, budget_activity, pe_bli, amount_type"
            " from read_parquet(?)", [str(pq)]).fetchall()
    finally:
        con.close()
    failures: list[str] = []
    mismatched = 0
    for fid, sha, exhibit, fy, account, org, ba, pe_bli, amount_type in rows:
        if fact_id_workbook(sha, exhibit, fy, account, org, ba, pe_bli, amount_type) != fid:
            mismatched += 1
            failures.append(f"{fid}: does not recompute with fact_id_workbook from its own row"
                            f" ({exhibit} FY{fy} {pe_bli} {amount_type})")
    counts = Counter(r[0] for r in rows)
    for fid, n in sorted(counts.items()):
        if n > 1:
            failures.append(f"{fid}: fact_id repeats on {n} rows")
    era = [r for r in rows
           if r[2] == "P-1" and r[3] is not None and int(r[3]) <= _ERA_P1_LAST_EDITION]
    re_keyed = [r for r in era if not is_era_procurement_key(r[7])]
    for r in re_keyed:
        failures.append(f"{r[0]}: P-1 FY{r[3]} row has pe_bli {r[7]!r}, not an era key"
                        " (spec §6.1: the lake identity is never re-keyed)")
    if not era:
        failures.append("no PB2017–PB2023 P-1 row: the era-key check would pass vacuously")
    return _leg(not failures,
                f"rows={len(rows)} era_p1_rows={len(era)} mismatched={mismatched}"
                f" duplicate_fact_ids={len(rows) - len(counts)}"
                f" era_rows_without_era_key={len(re_keyed)}", failures)


def leg_e_published_map(*, site_dir: Path, duckdb_path: Path) -> dict:
    """Leg (e): the published p1_era_line_map.parquet is the warehouse's map."""
    import duckdb

    pq = Path(site_dir) / "data" / "p1_era_line_map.parquet"
    if not pq.is_file():
        return _leg(False, "p1_era_line_map.parquet missing", [f"missing {pq}"])
    src = str(pq)
    try:
        con = duckdb.connect(str(duckdb_path), read_only=True)
    except duckdb.Error as e:
        return _leg(False, "warehouse unreadable", [f"cannot open {duckdb_path}: {e}"])
    try:
        try:
            mart_cols = [d[0] for d in con.execute(
                "select * from p1_era_line_map limit 0").description]
        except duckdb.CatalogException:
            return _leg(False, "p1_era_line_map missing from the warehouse",
                        [f"{duckdb_path}: no table p1_era_line_map"])
        pub_cols = [d[0] for d in con.execute(
            "select * from read_parquet(?) limit 0", [src]).description]
        if pub_cols != mart_cols:
            return _leg(False, "columns differ",
                        [f"published columns {pub_cols} != warehouse columns {mart_cols}"])
        mart_rows = con.execute("select count(*) from p1_era_line_map").fetchone()[0]
        pub_rows = con.execute("select count(*) from read_parquet(?)", [src]).fetchone()[0]
        only_mart = con.execute(
            "select count(*) from (select * from p1_era_line_map"
            " except all select * from read_parquet(?))", [src]).fetchone()[0]
        only_pub = con.execute(
            "select count(*) from (select * from read_parquet(?)"
            " except all select * from p1_era_line_map)", [src]).fetchone()[0]
    finally:
        con.close()
    failures: list[str] = []
    if mart_rows == 0:
        failures.append("the warehouse p1_era_line_map is empty")
    if only_mart:
        failures.append(f"{only_mart} warehouse row(s) are not in the published parquet: re-export")
    if only_pub:
        failures.append(f"{only_pub} published row(s) are not in the warehouse map: re-export")
    return _leg(not failures,
                f"mart_rows={mart_rows} published_rows={pub_rows}"
                f" only_in_mart={only_mart} only_published={only_pub}", failures)


def leg_f_history_pin(*, site_dir: Path) -> dict:
    """Leg (f): f15_funding_history.json is byte-identical to the S0 pin."""
    history = Path(site_dir) / "json" / "f15_funding_history.json"
    if not history.is_file():
        return _leg(False, "f15_funding_history.json missing", [f"missing {history}"])
    if not F15_HISTORY_PIN.is_file():
        return _leg(False, "pin missing", [f"missing pin {F15_HISTORY_PIN}"])
    tokens = F15_HISTORY_PIN.read_text().split()
    pin = tokens[0].lower() if tokens else ""
    if not re.fullmatch(r"[0-9a-f]{64}", pin):
        return _leg(False, "pin unreadable",
                    [f"{F15_HISTORY_PIN} does not start with a sha256 hex digest"])
    actual = _sha256_file(history)
    failures = [] if actual == pin else [
        f"{history.name} hashes to {actual}, the S0 pin says {pin}"
        " (only the S5 correction re-pins it)"]
    return _leg(not failures, f"pinned={pin} actual={actual}", failures)


def run_verify_era_map(
    *, site_dir: Path, duckdb_path: Path, seed_path: Path, legs: str = "abcdef",
) -> dict:
    """Run the selected legs. Returns {"legs": {leg: {"ok", "detail"}}, "verdict"}."""
    unknown = sorted(set(legs) - set(LEGS))
    if unknown or not legs:
        raise ValueError(f"legs must be letters from {''.join(LEGS)}; got {legs!r}")
    import duckdb

    try:
        map_rows, map_error = _read_map(Path(duckdb_path)), None
    except duckdb.Error as e:
        map_rows, map_error = [], f"cannot read p1_era_line_map from {duckdb_path}: {e}"
    results: dict[str, dict] = {}
    for leg in (x for x in LEGS if x in legs):
        if leg in "abc" and map_error:
            results[leg] = _leg(False, "p1_era_line_map unreadable", [map_error])
        elif leg == "a":
            results[leg] = leg_a_workbooks(
                site_dir=Path(site_dir), map_rows=map_rows,
                workbooks=ERA_P1_WORKBOOK_SHA256, expected_keys=EXPECTED_ERA_KEYS)
        elif leg == "b":
            results[leg] = leg_b_decisions(map_rows=map_rows, seed_path=Path(seed_path))
        elif leg == "c":
            results[leg] = leg_c_f15(map_rows=map_rows)
        elif leg == "d":
            results[leg] = leg_d_fact_ids(site_dir=Path(site_dir))
        elif leg == "e":
            results[leg] = leg_e_published_map(site_dir=Path(site_dir),
                                               duckdb_path=Path(duckdb_path))
        else:  # f
            results[leg] = leg_f_history_pin(site_dir=Path(site_dir))
    verdict = "PASS" if all(r["ok"] for r in results.values()) else "FAIL"
    return {"legs": results, "verdict": verdict}


def format_report(result: dict) -> list[str]:
    """One line per leg plus its first failures; the CLI prints the verdict after."""
    lines = []
    for leg, r in result["legs"].items():
        lines.append(f"leg {leg} {LEG_NAMES[leg]}: {r['detail']['summary']}"
                     f" → {'PASS' if r['ok'] else 'FAIL'}")
        for failure in r["detail"]["failures"][:_MAX_LISTED]:
            lines.append(f"  FAIL {failure}")
        extra = len(r["detail"]["failures"]) - _MAX_LISTED
        if extra > 0:
            lines.append(f"  … and {extra} more")
    return lines
