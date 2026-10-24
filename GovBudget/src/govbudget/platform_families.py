"""The platform-family registry: formats, vocabularies and the loader (families piece 2).

Spec: docs/superpowers/specs/2026-10-20-families-piece2-registry-design.md
§7 (the three registry files under data-seeds/: platform_families.json, the
frozen 49-column platform_family_members.csv and platform_family_era_lines.csv)
and §9 (the loader). Stage 1 reads the files alone, collects every fault and
raises one RegistryError whose lines read
'platform-families: <file>[:<row id>]: <reason>'; it never strips or otherwise
normalizes a value.

Module-level imports are the standard library, govbudget.config and
govbudget.jbooks.era_map (with govbudget.jbooks.era_keys, which era_map itself
imports) only: stage 1 runs on every import of the F-15 facade,
program_pdf_receipts and export_site. Stage 2 (§9.2: the lake, read-only, and
the doc:/src: files under the site directory) imports duckdb,
govbudget.export_site, govbudget.jbooks.collision_keys,
govbudget.jbooks.provenance_pages and openpyxl inside the functions that use
them.
"""
from __future__ import annotations

import codecs
import copy
import csv
import hashlib
import io
import json
import re
import unicodedata
from collections import defaultdict
from collections.abc import Callable, Iterable, Mapping, Sequence
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from functools import partial
from pathlib import Path

from govbudget import config
from govbudget.jbooks import era_keys, era_map

# ---------------------------------------------------------------------------
# Paths
# ---------------------------------------------------------------------------

REGISTRY_JSON = config.ROOT / "data-seeds" / "platform_families.json"
MEMBERS_CSV = config.ROOT / "data-seeds" / "platform_family_members.csv"
ERA_LINES_CSV = config.ROOT / "data-seeds" / "platform_family_era_lines.csv"
REVIEW_ROOT = config.RESEARCH_DIR / "platform_families"


@dataclass(frozen=True)
class RegistryPaths:
    json: Path
    members: Path
    era_lines: Path


DEFAULT_PATHS = RegistryPaths(REGISTRY_JSON, MEMBERS_CSV, ERA_LINES_CSV)

# ---------------------------------------------------------------------------
# Columns and key orders (spec §7.2, §7.3, §7.4)
# ---------------------------------------------------------------------------

MEMBER_COLUMNS = (  # spec §7.3, frozen order
    "member_row_id", "platform_family", "member_id", "exhibit", "account", "organization", "code",
    "role", "verdict", "project_number", "appropriation", "budget_activities", "titles_seen",
    "display_title", "page_slug", "valid_from_edition", "valid_to_edition",
    "coverage_expected_from_edition", "expected_next_edition", "era_decision_ids", "drift_sha256",
    "member_order", "program_tag", "purpose_tag", "purpose", "variant_ids", "variant_scope",
    "scope_note", "services_using", "paying_org", "predecessor_member_id", "successor_member_id",
    "link_basis", "link_evidence", "evidence", "quantity_bearing", "quantity_row", "quantity_notes",
    "jbook_anchor", "contract_link_policy", "msar_account", "msar_funding_source_page",
    "msar_shared_flag", "allocation_source", "allocation_years", "decided_by", "ruling_id",
    "decided_on", "note")
ERA_LINE_COLUMNS = ("platform_family", "member_id", "edition", "era_key", "budget_activity",
                    "filed_title", "decision_id")
REVIEW_COLUMNS = MEMBER_COLUMNS + ("reason", "headline_actuals_k")
PREFILL_COLUMNS = tuple(c for c in MEMBER_COLUMNS if c not in {
    "member_row_id", "titles_seen", "budget_activities", "drift_sha256", "decided_by", "ruling_id",
    "decided_on"}) + ("reason",)
ON_PARTIAL_CORPUS = ("skip", "raise")
RULING_KINDS = ("class", "batch")
HEADLINE_MEASURES = ("cumulative_actuals",)
GUARD_EXHIBITS = ("P-1", "R-1")
EDITIONS = range(2017, 2027)
SEED_BASIS_ALLOWED = frozenset({"R-DEC-FAM-F15-SEED"})
F15_SEED_RULING_ID = "R-DEC-FAM-F15-SEED"
TOP_KEYS = ("schema_version", "rulings", "sources", "families")
RULING_KEYS = ("ruling_id", "kind", "platform_family", "decided_by", "decided_on", "owner_reply",
               "spec", "text", "seed_basis")
SOURCE_KEYS = ("source_id", "title", "publisher", "url", "as_of", "sha256", "bytes", "archived",
               "archive_path", "tier", "verified_by")           # optional: sha256, bytes, archive_path, verified_by
FAMILY_KEYS = ("platform_family", "status", "published_ruling_id", "short_name", "history_dataset",
               "on_partial_corpus", "headline", "scope_note", "coverage_notes", "review_guard",
               "program_notes", "variants", "program_tags", "purpose_tags", "caveat_ids",
               "exclusions", "aliases", "contract_routes", "official_estimates", "art")  # optional: published_ruling_id
HEADLINE_KEYS = ("measure", "first_fy", "last_fy")
GUARD_KEYS = ("exhibits", "title_substrings")
PROGRAM_NOTE_KEYS = ("note", "note_fact_id", "expected_receipt")
RECEIPT_KEYS = ("kind", "page_number", "pe_bli", "sha256")
VARIANT_KEYS = ("id", "note")                                   # note optional
PROGRAM_TAG_KEYS = ("tag", "mdap_code")                         # mdap_code str or null
EXCLUSION_KEYS = ("exclusion_id", "kind", "label", "size_text", "evidence")
CONTRACT_ROUTE_KEYS = ("program_codes", "description_pattern", "description_exclusions", "basis")  # basis: action|award
OFFICIAL_ESTIMATE_KEYS = ("msar_program_number", "msar_source_ids", "gao_program_key", "wsaa_exit_ruling")
ART_KEYS = ("src", "alt", "caption", "generated")              # generated true needs "AI-generated" in caption

# ---------------------------------------------------------------------------
# Vocabularies (spec §7.2, §7.3)
# ---------------------------------------------------------------------------

ROLES = ("counted", "context", "project_slice", "excluded")
VERDICTS = ("approved", "deferred", "rejected")
EXHIBITS = ("P-1", "R-1", "P-1R")
STATUSES = ("draft", "ratified", "published")
SOURCE_TIERS = ("G", "PG", "C", "X")
APPROPRIATIONS = ("procurement", "rdte")
PURPOSES = ("buy", "upgrade", "develop", "support")
VARIANT_SCOPES = ("variant", "shared")
LINK_BASES = ("book_stated", "era_map", "inferred")
CONTRACT_LINK_POLICIES = ("production+AP", "modifications_only", "development", "none")
SERVICES = ("USAF", "USN", "USMC", "USA", "USSF", "SOCOM", "DW")
PAYING_ORGS = ("AF", "NAVY", "USMC", "ARMY", "USSF", "SOCOM", "DW")
CAVEAT_IDS = ("classified", "construction", "operating_and_support", "spares", "foreign",
              "cr_years", "nominal_dollars")
EXCLUSION_KINDS = ("outside_scope", "not_a_budget_line", "shared_line")
CONTRACT_ROUTE_BASES = ("action", "award")

# Every regex in this module is applied with .fullmatch (a bare `$` would
# accept a trailing newline).
PLATFORM_FAMILY_RE = re.compile(r"^[a-z0-9]+(-[a-z0-9]+)*$")
SHORT_NAME_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9 /-]*$")
HISTORY_DATASET_RE = re.compile(r"^[a-z0-9]+_funding_history$")
RULING_ID_RE = re.compile(r"^R-DEC-FAM-[A-Z0-9]+(-[A-Z0-9]+)+$")
SOURCE_ID_RE = re.compile(r"^[a-z0-9][a-z0-9-]*$")
MEMBER_ID_RE = re.compile(r"^(P-1|R-1|P-1R):[0-9]{4}[A-Z]+:[A-Z0-9]+:[A-Z0-9]+$")

_ISO_DATE_RE = re.compile(r"^[0-9]{4}-[0-9]{2}-[0-9]{2}$")
_SHA256_RE = re.compile(r"^[0-9a-f]{64}$")
_EDITION_RE = re.compile(r"^[0-9]{4}$")
_POSITIVE_RE = re.compile(r"^[1-9][0-9]*$")
_TITLE_ITEM_RE = re.compile(r"^([0-9]{4}): (.+)$")
_QUANTITY_ROW_RE = re.compile(r"^[A-Z]$")
_JBOOK_ANCHOR_RE = re.compile(
    r"^doc:[0-9a-f]{64}#p[1-9][0-9]*;xml=[^;]+;units=(Thousands|Millions)$")

# The inverse of era_map.ERA_ORG_TO_MODERN: a modern spelling -> its era spelling.
_MODERN_TO_ERA = {modern: era for era, modern in era_map.ERA_ORG_TO_MODERN.items()}

# Member columns by how Member.from_row parses them.
_INT_COLUMNS = frozenset({"valid_from_edition", "valid_to_edition",
                          "coverage_expected_from_edition", "expected_next_edition",
                          "member_order"})
_LIST_COLUMNS = ("budget_activities", "era_decision_ids", "variant_ids", "services_using",
                 "allocation_years")
_REF_COLUMNS = ("evidence", "link_evidence")
_BOOL_COLUMNS = ("quantity_bearing", "msar_shared_flag")
_REQUIRED_MEMBER_COLUMNS = ("platform_family", "member_id", "exhibit", "account", "organization",
                            "code", "role", "verdict", "appropriation", "display_title")
_ALWAYS_VOCABULARY = (("role", ROLES), ("verdict", VERDICTS), ("exhibit", EXHIBITS),
                      ("appropriation", APPROPRIATIONS))
_OPTIONAL_VOCABULARY = (("purpose", PURPOSES), ("variant_scope", VARIANT_SCOPES),
                        ("link_basis", LINK_BASES),
                        ("contract_link_policy", CONTRACT_LINK_POLICIES),
                        ("paying_org", PAYING_ORGS))


# ---------------------------------------------------------------------------
# Types
# ---------------------------------------------------------------------------


class RegistryError(ValueError):
    """Every stage-1 fault at once, one 'platform-families: …' line each."""

    def __init__(self, errors: Sequence[str]):
        self.errors = list(errors)
        super().__init__("\n".join(self.errors))


@dataclass(frozen=True)
class Ruling:
    ruling_id: str
    kind: str
    platform_family: str
    decided_by: str
    decided_on: str
    owner_reply: str
    spec: str
    text: str
    seed_basis: bool


@dataclass(frozen=True)
class Source:
    source_id: str
    title: str
    publisher: str
    url: str
    as_of: str
    sha256: str | None
    bytes: int | None
    archived: bool
    archive_path: str | None
    tier: str
    verified_by: tuple[str, ...]


def _split_ref_items(text: str) -> tuple[list[str], str | None]:
    """split_refs' scan: the refs read so far and its fault (None when the
    whole list splits)."""
    refs: list[str] = []
    if text == "":
        return refs, None
    for seg in text.split(";"):
        if seg.startswith("q="):
            if refs and refs[-1].startswith("doc:") and ";q=" not in refs[-1]:
                refs[-1] = f"{refs[-1]};{seg}"
                continue
            return refs, "dangling q="
        if seg == "":
            return refs, "empty list item"
        refs.append(seg)
    return refs, None


def split_refs(text: str) -> tuple[str, ...]:
    """A ';'-list of typed evidence refs (§7.5). A `q=` segment belongs to the
    `doc:` ref before it (a quote is the last part of a doc ref), so it is
    re-attached; prefixes are not validated here (parse_ref does)."""
    refs, fault = _split_ref_items(text)
    if fault is not None:
        raise ValueError(fault)
    return tuple(refs)


# The typed-ref grammar (spec §7.5), applied with .fullmatch. A `q=` quote may
# not contain ';', the list separator.
_REF_RES = {
    "era": re.compile(r"^era:([^|;]+\|[^|;]+\|[^|;]*\|[0-9]{4}-[0-9]{4})$"),
    "fact": re.compile(r"^fact:([0-9a-f]{16})$"),
    "doc": re.compile(r"^doc:([0-9a-f]{64})#p([1-9][0-9]*)(?:;q=([^;]+))?$"),
    "src": re.compile(r"^src:([a-z0-9][a-z0-9-]*)(?:#p([1-9][0-9]*))?$"),
    "ev": re.compile(r"^ev:(\S+)$"),
}
# A src: ref to a source of one of these tiers is an independent membership
# signal (§7.5); so is a doc: ref. At stage 1 a fact: ref never is (R1).
INDEPENDENT_TIERS = frozenset({"G", "PG"})


def parse_ref(ref: str) -> tuple[str, dict]:
    """One typed evidence ref -> (prefix, parts); ValueError on anything else.
    Parts: era {decision_id}; fact {fact_id}; doc {sha256, page, quote (None
    without q=)}; src {source_id, page (None without #p)}; ev {id}."""
    prefix = ref.partition(":")[0]
    pattern = _REF_RES.get(prefix)
    if pattern is None:
        raise ValueError(f"malformed ref {ref!r}: the prefix is not one of {', '.join(_REF_RES)}")
    m = pattern.fullmatch(ref)
    if m is None:
        raise ValueError(f"malformed ref {ref!r}: it does not match {pattern.pattern}")
    if prefix == "era":
        return prefix, {"decision_id": m.group(1)}
    if prefix == "fact":
        return prefix, {"fact_id": m.group(1)}
    if prefix == "doc":
        return prefix, {"sha256": m.group(1), "page": int(m.group(2)), "quote": m.group(3)}
    if prefix == "src":
        return prefix, {"source_id": m.group(1), "page": int(m.group(2)) if m.group(2) else None}
    return prefix, {"id": m.group(1)}


def _quote_hint(prev: str | None, item: str) -> str:
    """Review Focus 2: a ';' typed inside a q= quote ends the quote there, so
    the fault shows on the item after a quoted doc: ref (a ';'-list column) or
    inside the doc: item itself (an exclusion's evidence list)."""
    if prev is not None and prev.startswith("doc:") and ";q=" in prev:
        return " (the quote of the doc: ref before it contains ';', which a q= quote may not)"
    if item.startswith("doc:") and ";" in item.partition(";q=")[2]:
        return " (its quote contains ';', which a q= quote may not)"
    return ""


def _parse_int(text: str) -> int | None:
    if text == "":
        return None
    if not (text.isascii() and text.isdigit() and str(int(text)) == text):
        raise ValueError(f"not an integer: {text!r}")
    return int(text)


def _parse_list(text: str) -> tuple[str, ...]:
    if text == "":
        return ()
    items = tuple(text.split(";"))
    if "" in items:
        raise ValueError(f"empty list item in {text!r}")
    return items


def _parse_bool(text: str) -> bool | None:
    if text == "":
        return None
    if text not in ("true", "false"):
        raise ValueError(f"not 'true' or 'false': {text!r}")
    return text == "true"


def _parse_titles(text: str) -> tuple[tuple[int, str], ...]:
    if text == "":
        return ()
    out = []
    for item in text.split("; "):
        m = _TITLE_ITEM_RE.fullmatch(item)
        if not m:
            raise ValueError(f"titles_seen item {item!r} is not '<edition>: <title>'")
        out.append((int(m.group(1)), m.group(2)))
    return tuple(out)


@dataclass(frozen=True)
class Member:
    """One row of platform_family_members.csv, one field per MEMBER_COLUMNS name."""

    member_row_id: str
    platform_family: str
    member_id: str
    exhibit: str
    account: str
    organization: str
    code: str
    role: str
    verdict: str
    project_number: str
    appropriation: str
    budget_activities: tuple[str, ...]
    titles_seen: tuple[tuple[int, str], ...]
    display_title: str
    page_slug: str
    valid_from_edition: int | None
    valid_to_edition: int | None
    coverage_expected_from_edition: int | None
    expected_next_edition: int | None
    era_decision_ids: tuple[str, ...]
    drift_sha256: str
    member_order: int | None
    program_tag: str
    purpose_tag: str
    purpose: str
    variant_ids: tuple[str, ...]
    variant_scope: str
    scope_note: str
    services_using: tuple[str, ...]
    paying_org: str
    predecessor_member_id: str
    successor_member_id: str
    link_basis: str
    link_evidence: tuple[str, ...]
    evidence: tuple[str, ...]
    quantity_bearing: bool | None
    quantity_row: str
    quantity_notes: str
    jbook_anchor: str
    contract_link_policy: str
    msar_account: str
    msar_funding_source_page: str
    msar_shared_flag: bool | None
    allocation_source: str
    allocation_years: tuple[str, ...]
    decided_by: str
    ruling_id: str
    decided_on: str
    note: str

    @classmethod
    def from_row(cls, row: Mapping[str, str]) -> Member:
        """Parse a row over MEMBER_COLUMNS; ValueError on an unparsable value."""
        values: dict[str, object] = {}
        for col in MEMBER_COLUMNS:
            if col not in row:
                raise ValueError(f"missing column {col!r}")
            text = row[col]
            if not isinstance(text, str):
                raise ValueError(f"{col} must be a string, got {text!r}")
            if col in _INT_COLUMNS:
                values[col] = _parse_int(text)
            elif col in _LIST_COLUMNS:
                values[col] = _parse_list(text)
            elif col in _REF_COLUMNS:
                values[col] = split_refs(text)
            elif col in _BOOL_COLUMNS:
                values[col] = _parse_bool(text)
            elif col == "titles_seen":
                values[col] = _parse_titles(text)
            else:
                values[col] = text
        return cls(**values)

    def title_for(self, edition: int) -> str | None:
        return dict(self.titles_seen).get(edition)

    def as_row(self) -> dict[str, str]:
        """The exact inverse of from_row for every row stage 1 accepts."""
        row: dict[str, str] = {}
        for col in MEMBER_COLUMNS:
            value = getattr(self, col)
            if col == "titles_seen":
                row[col] = "; ".join(f"{e}: {t}" for e, t in value)
            elif col in _LIST_COLUMNS or col in _REF_COLUMNS:
                row[col] = ";".join(value)
            elif value is None:
                row[col] = ""
            elif isinstance(value, bool):
                row[col] = "true" if value else "false"
            else:
                row[col] = str(value)
        return row


@dataclass(frozen=True)
class EraLine:
    platform_family: str
    member_id: str
    edition: int
    era_key: str
    budget_activity: str
    filed_title: str
    decision_id: str

    def as_row(self) -> dict[str, str]:
        return {"platform_family": self.platform_family, "member_id": self.member_id,
                "edition": str(self.edition), "era_key": self.era_key,
                "budget_activity": self.budget_activity, "filed_title": self.filed_title,
                "decision_id": self.decision_id}


@dataclass(frozen=True)
class Headline:
    measure: str
    first_fy: int
    last_fy: int


@dataclass(frozen=True)
class ReviewGuard:
    exhibits: tuple[str, ...]
    title_substrings: tuple[str, ...]


@dataclass(frozen=True)
class ProgramNote:
    note: str
    note_fact_id: str
    expected_receipt: Mapping[str, object]


def _counted_order(member: Member) -> tuple[bool, int]:
    return (member.member_order is None, member.member_order or 0)


@dataclass(frozen=True)
class Family:
    """A family object of platform_families.json plus its CSV rows."""

    platform_family: str
    status: str
    published_ruling_id: str | None
    short_name: str
    history_dataset: str
    on_partial_corpus: str
    headline: Headline
    scope_note: str
    coverage_notes: tuple[str, ...]
    review_guard: ReviewGuard
    program_notes: Mapping[str, ProgramNote]
    variants: tuple[Mapping[str, str], ...]
    program_tags: tuple[Mapping[str, str | None], ...]
    purpose_tags: tuple[str, ...]
    caveat_ids: tuple[str, ...]
    exclusions: tuple[Mapping[str, object], ...]
    aliases: tuple[str, ...]
    contract_routes: Mapping[str, object] | None
    official_estimates: Mapping[str, object] | None
    art: Mapping[str, object] | None
    members: tuple[Member, ...]
    era_lines: tuple[EraLine, ...]

    @property
    def counted(self) -> tuple[Member, ...]:
        """Approved counted rows in member_order."""
        return tuple(sorted((m for m in self.members if m.verdict == "approved" and m.role == "counted"),
                            key=_counted_order))

    @property
    def code_keyed(self) -> tuple[Member, ...]:
        """D5: the counted rows with code-keyed lake rows (a non-empty titles_seen)."""
        return tuple(m for m in self.counted if m.titles_seen)

    @property
    def context(self) -> tuple[Member, ...]:
        """Approved context rows by member_id."""
        return tuple(sorted((m for m in self.members if m.verdict == "approved" and m.role == "context"),
                            key=lambda m: m.member_id))

    @property
    def route(self) -> str:
        return f"/families/{self.platform_family}/"

    @property
    def history_path(self) -> str:
        return f"json/{self.history_dataset}.json"


@dataclass(frozen=True)
class Registry:
    schema_version: int
    rulings: Mapping[str, Ruling]
    sources: Mapping[str, Source]
    families: tuple[Family, ...]
    doc: dict

    def __getitem__(self, platform_family: str) -> Family:
        for fam in self.families:
            if fam.platform_family == platform_family:
                return fam
        raise KeyError(platform_family)

    def published(self) -> tuple[Family, ...]:
        return tuple(f for f in self.families if f.status == "published")


# ---------------------------------------------------------------------------
# Stage 1 (spec §9.1): the JSON
# ---------------------------------------------------------------------------

_STRS = "a list of strings"
_STR_OR_NULL = "a string or null"
_OBJ_OR_NULL = "an object or null"
_TYPE_WORDS = {str: "a string", bool: "a boolean", int: "an integer", list: "a list",
               dict: "an object"}

_RULING_TYPES = {"ruling_id": str, "kind": str, "platform_family": str, "decided_by": str,
                 "decided_on": str, "owner_reply": str, "spec": str, "text": str,
                 "seed_basis": bool}
_SOURCE_TYPES = {"source_id": str, "title": str, "publisher": str, "url": str, "as_of": str,
                 "sha256": str, "bytes": int, "archived": bool, "archive_path": str, "tier": str,
                 "verified_by": _STRS}
_FAMILY_TYPES = {"platform_family": str, "status": str, "published_ruling_id": str,
                 "short_name": str, "history_dataset": str, "on_partial_corpus": str,
                 "headline": dict, "scope_note": str, "coverage_notes": _STRS,
                 "review_guard": dict, "program_notes": dict, "variants": list,
                 "program_tags": list, "purpose_tags": _STRS, "caveat_ids": _STRS,
                 "exclusions": list, "aliases": _STRS, "contract_routes": _OBJ_OR_NULL,
                 "official_estimates": _OBJ_OR_NULL, "art": _OBJ_OR_NULL}
_HEADLINE_TYPES = {"measure": str, "first_fy": int, "last_fy": int}
_GUARD_TYPES = {"exhibits": _STRS, "title_substrings": _STRS}
_PROGRAM_NOTE_TYPES = {"note": str, "note_fact_id": str, "expected_receipt": dict}
_RECEIPT_TYPES = {"kind": str, "page_number": int, "pe_bli": str, "sha256": str}
_VARIANT_TYPES = {"id": str, "note": str}
_PROGRAM_TAG_TYPES = {"tag": str, "mdap_code": _STR_OR_NULL}
_EXCLUSION_TYPES = {"exclusion_id": str, "kind": str, "label": str, "size_text": str,
                    "evidence": _STRS}
_CONTRACT_ROUTE_TYPES = {"program_codes": _STRS, "description_pattern": str,
                         "description_exclusions": _STRS, "basis": str}
_OFFICIAL_ESTIMATE_TYPES = {"msar_program_number": _STR_OR_NULL, "msar_source_ids": _STRS,
                            "gao_program_key": _STR_OR_NULL, "wsaa_exit_ruling": _STR_OR_NULL}
_ART_TYPES = {"src": str, "alt": str, "caption": str, "generated": bool}

_Add = Callable[[str], None]
_Rows = list[tuple[str, Mapping[str, str]]]     # (row id, row), CSV order


def _has_type(value: object, kind: object) -> bool:
    if kind is int:
        return type(value) is int          # never a bool
    if kind is bool:
        return type(value) is bool         # never an int
    if kind == _STRS:
        return isinstance(value, list) and all(isinstance(v, str) for v in value)
    if kind == _STR_OR_NULL:
        return value is None or isinstance(value, str)
    if kind == _OBJ_OR_NULL:
        return value is None or isinstance(value, dict)
    return isinstance(value, kind)


def _check_shape(obj: object, keys: Sequence[str], types: Mapping[str, object], add: _Add, *,
                 label: str = "", optional: frozenset[str] = frozenset()) -> bool:
    """Exact keys (required = all but `optional`) and JSON types. False when
    `obj` is not an object at all."""
    prefix = f"{label}: " if label else ""
    if not isinstance(obj, dict):
        add(f"{label or 'entry'} must be an object, got {obj!r}")
        return False
    for key in obj:
        if key not in keys:
            add(f"{prefix}unknown key {key!r}")
    for key in keys:
        if key not in obj:
            if key not in optional:
                add(f"{prefix}missing key {key!r}")
        elif not _has_type(obj[key], types[key]):
            word = types[key] if isinstance(types[key], str) else _TYPE_WORDS[types[key]]
            add(f"{prefix}{key} must be {word}, got {obj[key]!r}")
    return True


def _iso_date(value: object) -> bool:
    if not isinstance(value, str) or not _ISO_DATE_RE.fullmatch(value):
        return False
    try:
        date.fromisoformat(value)
    except ValueError:
        return False
    return True


def _entry_id(obj: object, key: str, fallback: str) -> str:
    value = obj.get(key) if isinstance(obj, Mapping) else None
    return value if isinstance(value, str) else fallback


def _strs(obj: Mapping, key: str) -> list[str]:
    """The value as a list of strings, or [] when it is anything else."""
    value = obj.get(key)
    return value if _has_type(value, _STRS) else []


def _objects(obj: Mapping, key: str) -> list[dict]:
    value = obj.get(key)
    return [v for v in value if isinstance(v, dict)] if isinstance(value, list) else []


def _check_unique(seen: set, value: object, key: str, add: _Add) -> None:
    if not isinstance(value, str):
        return
    if value in seen:
        add(f"duplicate {key} {value!r}")
    seen.add(value)


def _check_rulings(rulings: list, add_at: Callable[[str, str], None]) -> None:
    seen: set[str] = set()
    for i, r in enumerate(rulings):
        rid = _entry_id(r, "ruling_id", f"rulings[{i}]")

        def add(reason: str, rid: str = rid) -> None:
            add_at(rid, reason)

        if not _check_shape(r, RULING_KEYS, _RULING_TYPES, add):
            continue
        ruling_id = r.get("ruling_id")
        if isinstance(ruling_id, str) and not RULING_ID_RE.fullmatch(ruling_id):
            add(f"ruling_id {ruling_id!r} does not match {RULING_ID_RE.pattern}")
        _check_unique(seen, ruling_id, "ruling_id", add)
        kind = r.get("kind")
        if isinstance(kind, str) and kind not in RULING_KINDS:
            add(f"kind {kind!r} is not one of {', '.join(RULING_KINDS)}")
        decided_by = r.get("decided_by")
        if isinstance(decided_by, str) and decided_by != "owner":
            add(f"decided_by {decided_by!r} must be 'owner'")
        reply = r.get("owner_reply")
        if isinstance(reply, str) and not reply.strip():
            add(f"owner_reply {reply!r} is blank")
        if isinstance(r.get("decided_on"), str) and not _iso_date(r["decided_on"]):
            add(f"decided_on {r['decided_on']!r} is not an ISO date (YYYY-MM-DD)")
        if r.get("seed_basis") is True and ruling_id not in SEED_BASIS_ALLOWED:
            add(f"seed_basis true is allowed only on {', '.join(sorted(SEED_BASIS_ALLOWED))}, "
                f"not on {ruling_id!r}")


def _check_sources(sources: list, add_at: Callable[[str, str], None]) -> None:
    seen: set[str] = set()
    optional = frozenset({"sha256", "bytes", "archive_path", "verified_by"})
    for i, s in enumerate(sources):
        sid = _entry_id(s, "source_id", f"sources[{i}]")

        def add(reason: str, sid: str = sid) -> None:
            add_at(sid, reason)

        if not _check_shape(s, SOURCE_KEYS, _SOURCE_TYPES, add, optional=optional):
            continue
        source_id = s.get("source_id")
        if isinstance(source_id, str) and not SOURCE_ID_RE.fullmatch(source_id):
            add(f"source_id {source_id!r} does not match {SOURCE_ID_RE.pattern}")
        _check_unique(seen, source_id, "source_id", add)
        tier = s.get("tier")
        if isinstance(tier, str) and tier not in SOURCE_TIERS:
            add(f"tier {tier!r} is not one of {', '.join(SOURCE_TIERS)}")
        sha = s.get("sha256")
        if isinstance(sha, str) and not _SHA256_RE.fullmatch(sha):
            add(f"sha256 {sha!r} is not 64 lower-case hex")
        if isinstance(s.get("as_of"), str) and not _iso_date(s["as_of"]):
            add(f"as_of {s['as_of']!r} is not an ISO date (YYYY-MM-DD)")
        if s.get("archived") is True and not s.get("archive_path"):
            add(f"archive_path is required when archived is true (got {s.get('archive_path')!r})")


def _check_family_fields(fam: dict, add: _Add) -> None:
    """Rule 1 for the family's nested objects and lists, rules 5, 18 (the
    vocabularies' own duplicates), 26 and 27."""
    status = fam.get("status")
    if isinstance(status, str) and status not in STATUSES:
        add(f"status {status!r} is not one of {', '.join(STATUSES)}")
    opc = fam.get("on_partial_corpus")
    if isinstance(opc, str) and opc not in ON_PARTIAL_CORPUS:
        add(f"on_partial_corpus {opc!r} is not one of {', '.join(ON_PARTIAL_CORPUS)}")
    for caveat in _strs(fam, "caveat_ids"):
        if caveat not in CAVEAT_IDS:
            add(f"caveat_ids item {caveat!r} is not one of {', '.join(CAVEAT_IDS)}")

    headline = fam.get("headline")
    if isinstance(headline, dict) and _check_shape(headline, HEADLINE_KEYS, _HEADLINE_TYPES, add,
                                                   label="headline"):
        measure = headline.get("measure")
        if isinstance(measure, str) and measure not in HEADLINE_MEASURES:
            add(f"headline measure {measure!r} is not one of {', '.join(HEADLINE_MEASURES)}")
        first, last = headline.get("first_fy"), headline.get("last_fy")
        if type(first) is int and type(last) is int and first > last:
            add(f"headline first_fy {first!r} is after last_fy {last!r}")

    guard = fam.get("review_guard")
    if isinstance(guard, dict) and _check_shape(guard, GUARD_KEYS, _GUARD_TYPES, add,
                                                label="review_guard"):
        subs = guard.get("title_substrings")
        if _has_type(subs, _STRS):
            if not subs:
                add("review_guard title_substrings [] is empty")
            for sub in subs:
                if sub == "" or sub != sub.upper():
                    add(f"review_guard title_substrings item {sub!r} is not a non-empty "
                        "upper-case string")
        exhibits = guard.get("exhibits")
        if _has_type(exhibits, _STRS) and (not exhibits or not set(exhibits) <= set(GUARD_EXHIBITS)):
            add(f"review_guard exhibits {exhibits!r} must be a non-empty subset of "
                f"{', '.join(GUARD_EXHIBITS)}")

    notes = fam.get("program_notes")
    if isinstance(notes, dict):
        for member_id, note in notes.items():
            label = f"program_notes[{member_id!r}]"
            if _check_shape(note, PROGRAM_NOTE_KEYS, _PROGRAM_NOTE_TYPES, add, label=label):
                receipt = note.get("expected_receipt")
                if isinstance(receipt, dict):
                    _check_shape(receipt, RECEIPT_KEYS, _RECEIPT_TYPES, add,
                                 label=f"{label}.expected_receipt")

    for key, keys, types, optional in (
            ("variants", VARIANT_KEYS, _VARIANT_TYPES, frozenset({"note"})),
            ("program_tags", PROGRAM_TAG_KEYS, _PROGRAM_TAG_TYPES, frozenset()),
            ("exclusions", EXCLUSION_KEYS, _EXCLUSION_TYPES, frozenset())):
        items = fam.get(key)
        if isinstance(items, list):
            for j, item in enumerate(items):
                _check_shape(item, keys, types, add, label=f"{key}[{j}]", optional=optional)

    seen_exclusions: set[str] = set()
    for j, exclusion in enumerate(_objects(fam, "exclusions")):
        kind = exclusion.get("kind")
        if isinstance(kind, str) and kind not in EXCLUSION_KINDS:
            add(f"exclusions[{j}] kind {kind!r} is not one of {', '.join(EXCLUSION_KINDS)}")
        _check_unique(seen_exclusions, exclusion.get("exclusion_id"), "exclusion_id", add)

    seen_variants: set[str] = set()
    for variant in _objects(fam, "variants"):
        _check_unique(seen_variants, variant.get("id"), "variants id", add)
    seen_tags: set[str] = set()
    for tag in _objects(fam, "program_tags"):
        _check_unique(seen_tags, tag.get("tag"), "program_tags tag", add)
    seen_purposes: set[str] = set()
    for purpose_tag in _strs(fam, "purpose_tags"):
        _check_unique(seen_purposes, purpose_tag, "purpose_tags item", add)

    for key, keys, types in (("contract_routes", CONTRACT_ROUTE_KEYS, _CONTRACT_ROUTE_TYPES),
                             ("official_estimates", OFFICIAL_ESTIMATE_KEYS, _OFFICIAL_ESTIMATE_TYPES),
                             ("art", ART_KEYS, _ART_TYPES)):
        if isinstance(fam.get(key), dict):
            _check_shape(fam[key], keys, types, add, label=key)
    routes = fam.get("contract_routes")
    if isinstance(routes, dict) and isinstance(routes.get("basis"), str) \
            and routes["basis"] not in CONTRACT_ROUTE_BASES:
        add(f"contract_routes basis {routes['basis']!r} is not one of {', '.join(CONTRACT_ROUTE_BASES)}")
    art = fam.get("art")
    if isinstance(art, dict) and art.get("generated") is True and (
            not isinstance(art.get("caption"), str) or "AI-generated" not in art["caption"]):
        add(f"art generated true needs a caption containing 'AI-generated' (caption {art.get('caption')!r})")


def _check_families(families: list, add_at: Callable[[str, str], None]) -> dict[str, dict]:
    """Rules 1–5, 18 (family side), 26, 27. Returns the well-formed family
    objects by platform_family (first occurrence) for the row checks."""
    by_id: dict[str, dict] = {}
    seen = {"platform_family": set(), "short_name": set(), "history_dataset": set()}
    patterns = {"platform_family": PLATFORM_FAMILY_RE, "short_name": SHORT_NAME_RE,
                "history_dataset": HISTORY_DATASET_RE}
    for i, fam in enumerate(families):
        pf = _entry_id(fam, "platform_family", f"families[{i}]")

        def add(reason: str, pf: str = pf) -> None:
            add_at(pf, reason)

        if not _check_shape(fam, FAMILY_KEYS, _FAMILY_TYPES, add,
                            optional=frozenset({"published_ruling_id"})):
            continue
        for key, pattern in patterns.items():
            value = fam.get(key)
            if isinstance(value, str) and not pattern.fullmatch(value):
                add(f"{key} {value!r} does not match {pattern.pattern}")
            _check_unique(seen[key], value, key, add)
        _check_family_fields(fam, add)
        if isinstance(fam.get("platform_family"), str):
            by_id.setdefault(fam["platform_family"], fam)
    return by_id


def _check_doc(doc: object, add_at: Callable[[str | None, str], None]) -> dict[str, dict]:
    if not isinstance(doc, dict):
        add_at(None, f"the top level must be a JSON object, got {type(doc).__name__}")
        return {}
    for key in doc:
        if key not in TOP_KEYS:
            add_at(None, f"unknown key {key!r}")
    for key in TOP_KEYS:
        if key not in doc:
            add_at(None, f"missing key {key!r}")
    if "schema_version" in doc and not (type(doc["schema_version"]) is int and doc["schema_version"] == 1):
        add_at(None, f"schema_version must be 1, got {doc['schema_version']!r}")
    for key in ("rulings", "sources", "families"):
        if key in doc and not isinstance(doc[key], list):
            add_at(None, f"{key} must be a list, got {doc[key]!r}")
    if isinstance(doc.get("rulings"), list):
        _check_rulings(doc["rulings"], add_at)
    if isinstance(doc.get("sources"), list):
        _check_sources(doc["sources"], add_at)
    if isinstance(doc.get("families"), list):
        return _check_families(doc["families"], add_at)
    return {}


# ---------------------------------------------------------------------------
# Stage 1 (spec §9.1): the CSV rows
# ---------------------------------------------------------------------------


def _row_columns(row: Mapping, columns: Sequence[str], add: _Add) -> bool:
    """Rule 8: exactly the frozen columns, in order, every value a string."""
    keys = tuple(row.keys())
    if keys != tuple(columns):
        missing = [c for c in columns if c not in row]
        extra = [k for k in keys if k not in columns]
        detail = "; ".join(filter(None, (
            f"missing {', '.join(map(repr, missing))}" if missing else "",
            f"unexpected {', '.join(map(repr, extra))}" if extra else ""))) or "out of order"
        add(f"columns are not the frozen {len(columns)}-column header ({detail})")
        return False
    bad = [(k, v) for k, v in row.items() if not isinstance(v, str)]
    for key, value in bad:
        add(f"column {key} must be a string, got {value!r}")
    return not bad


def _check_editions(row: Mapping[str, str], add: _Add) -> None:
    """Rule 12."""
    bounds: dict[str, int] = {}
    for col in ("valid_from_edition", "valid_to_edition"):
        text = row[col]
        if not _EDITION_RE.fullmatch(text) or int(text) not in EDITIONS:
            add(f"{col} {text!r} is not an edition in {EDITIONS[0]}–{EDITIONS[-1]}")
        else:
            bounds[col] = int(text)
    frm, to = bounds.get("valid_from_edition"), bounds.get("valid_to_edition")
    span = frm is not None and to is not None
    if span and frm > to:
        add(f"valid_from_edition {frm!r} is after valid_to_edition {to!r}")
    cov = row["coverage_expected_from_edition"]
    if cov:
        if not _EDITION_RE.fullmatch(cov):
            add(f"coverage_expected_from_edition {cov!r} is not a 4-digit edition")
        elif span and not frm <= int(cov) <= to:
            add(f"coverage_expected_from_edition {cov!r} is outside [{frm}, {to}]")
    nxt = row["expected_next_edition"]
    if nxt:
        if not _EDITION_RE.fullmatch(nxt):
            add(f"expected_next_edition {nxt!r} is not a 4-digit edition")
        elif to is not None and int(nxt) <= to:
            add(f"expected_next_edition {nxt!r} is not after valid_to_edition {to!r}")
    order = row["member_order"]
    if order and not _POSITIVE_RE.fullmatch(order):
        add(f"member_order {order!r} is not a positive integer")
    titles = row["titles_seen"]
    if titles:
        prev = None
        for item in titles.split("; "):
            m = _TITLE_ITEM_RE.fullmatch(item)
            if not m:
                add(f"titles_seen item {item!r} is not '<edition>: <title>'")
                break
            edition = int(m.group(1))
            if prev is not None and edition <= prev:
                add(f"titles_seen {titles!r}: editions are not strictly ascending ({prev} then {edition})")
                break
            if span and not frm <= edition <= to:
                add(f"titles_seen {titles!r}: edition {edition} is outside [{frm}, {to}]")
                break
            prev = edition


def _check_member_values(row: Mapping[str, str], add: _Add) -> None:
    """Rules 10 and 11."""
    member_id = row["member_id"]
    if not MEMBER_ID_RE.fullmatch(member_id):
        add(f"member_id {member_id!r} does not match {MEMBER_ID_RE.pattern}")
    else:
        for col, part in zip(("exhibit", "account", "organization", "code"), member_id.split(":")):
            if row[col] != part:
                add(f"{col} {row[col]!r} disagrees with member_id {member_id!r} ({part!r})")
    org = row["organization"]
    if org in _MODERN_TO_ERA:
        add(f"organization {org!r} is a modern spelling; the member uses the era spelling "
            f"{_MODERN_TO_ERA[org]!r}")

    for col in _REQUIRED_MEMBER_COLUMNS:
        if row[col] == "":
            add(f"{col} is required (got '')")
    for col, vocab in _ALWAYS_VOCABULARY + _OPTIONAL_VOCABULARY:
        if row[col] != "" and row[col] not in vocab:
            add(f"{col} {row[col]!r} is not one of {', '.join(vocab)}")
    if row["purpose"] == "" and row["page_slug"] != "":
        add(f"purpose is blank but page_slug {row['page_slug']!r} is set")
    if row["exhibit"] == "P-1R" and row["role"] != "context":
        add(f"exhibit 'P-1R' is allowed only with role 'context', not {row['role']!r}")

    for col in _LIST_COLUMNS:
        if row[col] and "" in row[col].split(";"):
            add(f"{col} {row[col]!r} has an empty list item")
    for item in filter(None, row["services_using"].split(";")):
        if item not in SERVICES:
            add(f"services_using item {item!r} is not one of {', '.join(SERVICES)}")
    for item in filter(None, row["allocation_years"].split(";")):
        if not _EDITION_RE.fullmatch(item):
            add(f"allocation_years item {item!r} is not a 4-digit year")
    for col in _REF_COLUMNS:
        refs, fault = _split_ref_items(row[col])
        if fault is not None:
            add(f"{col} {row[col]!r}: {fault}{_quote_hint(refs[-1] if refs else None, '')}")
    for col in _BOOL_COLUMNS:
        if row[col] not in ("", "true", "false"):
            add(f"{col} {row[col]!r} must be blank, 'true' or 'false'")

    drift = row["drift_sha256"]
    if drift == "":
        if row["verdict"] == "approved":
            add("drift_sha256 is required on an approved row (got '')")
    elif not _SHA256_RE.fullmatch(drift):
        add(f"drift_sha256 {drift!r} is not 64 lower-case hex")
    if row["quantity_row"] and not _QUANTITY_ROW_RE.fullmatch(row["quantity_row"]):
        add(f"quantity_row {row['quantity_row']!r} is not one upper-case letter")
    page = row["msar_funding_source_page"]
    if page and not _POSITIVE_RE.fullmatch(page):
        add(f"msar_funding_source_page {page!r} is not a positive integer")
    anchor = row["jbook_anchor"]
    if anchor and not _JBOOK_ANCHOR_RE.fullmatch(anchor):
        add(f"jbook_anchor {anchor!r} does not match {_JBOOK_ANCHOR_RE.pattern}")


def _check_role_rules(row: Mapping[str, str], add: _Add) -> None:
    """Rule 17."""
    role = row["role"]
    if role == "project_slice" and not row["project_number"]:
        add("project_number is required on a project_slice row (got '')")
    if role != "project_slice" and row["project_number"]:
        add(f"project_number {row['project_number']!r} is allowed only on a project_slice row, "
            f"not role {role!r}")
    if row["variant_scope"] == "shared" and not row["scope_note"]:
        add("scope_note is required on a shared row (got '')")
    elif role == "context" and not row["scope_note"]:
        add("scope_note is required on a context row (got '')")
    if role == "context":
        for col in ("coverage_expected_from_edition", "member_order"):
            if row[col]:
                add(f"{col} {row[col]!r} is not allowed on a context row")
    else:
        for col in ("allocation_source", "allocation_years"):
            if row[col]:
                add(f"{col} {row[col]!r} is allowed only on a context row, not role {role!r}")
    if role != "counted":
        for col in ("quantity_bearing", "quantity_row", "quantity_notes"):
            if row[col]:
                add(f"{col} {row[col]!r} is allowed only on a counted row, not role {role!r}")


def _check_tags(row: Mapping[str, str], fam: Mapping, add: _Add) -> None:
    """Rule 18 (row side): tags and variants from the family's own lists."""
    pf = row["platform_family"]
    variant_ids = {v.get("id") for v in _objects(fam, "variants")}
    for item in filter(None, row["variant_ids"].split(";")):
        if item not in variant_ids:
            add(f"variant_ids item {item!r} is not a variants id of {pf!r}")
    tag = row["program_tag"]
    if tag and tag not in {t.get("tag") for t in _objects(fam, "program_tags")}:
        add(f"program_tag {tag!r} is not in the program_tags of {pf!r}")
    purpose_tag = row["purpose_tag"]
    if purpose_tag and purpose_tag not in _strs(fam, "purpose_tags"):
        add(f"purpose_tag {purpose_tag!r} is not in the purpose_tags of {pf!r}")


def _check_members(rows: Sequence[Mapping[str, str]], families: Mapping[str, dict],
                   add_at: Callable[[str, str], None]) -> _Rows:
    """The per-row rules. Returns the rows that pass rule 8 (exactly the frozen
    columns, every value a string) with their row ids, in CSV order: the rows
    the cross-row rules read."""
    seen: set[str] = set()
    valid: _Rows = []
    for i, row in enumerate(rows):
        # The row id is member_row_id as written; without one, the row's line
        # in the CSV (the header is line 1).
        rid = _entry_id(row, "member_row_id", "") or f"line {i + 2}"

        def add(reason: str, rid: str = rid) -> None:
            add_at(rid, reason)

        if not isinstance(row, Mapping):
            add(f"row must be a mapping of the frozen columns, got {row!r}")
            continue
        if not _row_columns(row, MEMBER_COLUMNS, add):
            continue
        valid.append((rid, row))
        expected = (f"{row['platform_family']}|{row['member_id']}|"
                    f"{row['valid_from_edition']}-{row['valid_to_edition']}")
        if row["member_row_id"] != expected:
            add(f"member_row_id {row['member_row_id']!r} is not {expected!r} "
                "(platform_family|member_id|valid_from_edition-valid_to_edition)")
        _check_unique(seen, row["member_row_id"], "member_row_id", add)
        _check_member_values(row, add)
        _check_editions(row, add)
        _check_role_rules(row, add)
        fam = families.get(row["platform_family"])
        if fam is not None:
            _check_tags(row, fam, add)
    return valid


def _check_era_rows(rows: Sequence[Mapping[str, str]], add_at: Callable[[str, str], None]) -> _Rows:
    """Rules 8 and 12 per era row; returns the rows that pass rule 8 with their
    row ids, in CSV order."""
    valid: _Rows = []
    for i, row in enumerate(rows):
        parts = [row.get(k) for k in ("platform_family", "edition", "era_key")] \
            if isinstance(row, Mapping) else []
        rid = "|".join(parts) if parts and all(isinstance(p, str) for p in parts) else f"line {i + 2}"

        def add(reason: str, rid: str = rid) -> None:
            add_at(rid, reason)

        if not isinstance(row, Mapping):
            add(f"row must be a mapping of the frozen columns, got {row!r}")
            continue
        if not _row_columns(row, ERA_LINE_COLUMNS, add):
            continue
        valid.append((rid, row))
        edition = row["edition"]
        if not _EDITION_RE.fullmatch(edition) or int(edition) not in EDITIONS:
            add(f"edition {edition!r} is not an edition in {EDITIONS[0]}–{EDITIONS[-1]}")
    return valid


# ---------------------------------------------------------------------------
# Stage 1 (spec §9.1): the rules across rows and files
# ---------------------------------------------------------------------------

_ERA_REF_DECISIONS = ("same_program", "history_only")
_FACT_ID_RE = re.compile(r"^[0-9a-f]{16}$")


@dataclass(frozen=True)
class _Cross:
    """What the cross-row rules read. They run beside the per-row rules on
    whatever is well-formed enough to compare, so every value is guarded."""

    json_name: str
    rulings: Mapping[str, dict]               # by ruling_id, first occurrence
    sources: Mapping[str, dict]               # by source_id, first occurrence
    families: Mapping[str, dict]              # by platform_family, JSON order
    rows: _Rows                               # member rows that pass rule 8, CSV order
    decisions: Mapping[str, Mapping[str, str]]  # the era seed by decision_id


def _by_id(entries: object, key: str) -> dict[str, dict]:
    out: dict[str, dict] = {}
    if isinstance(entries, list):
        for entry in entries:
            if isinstance(entry, dict) and isinstance(entry.get(key), str):
                out.setdefault(entry[key], entry)
    return out


def _span(row: Mapping[str, str]) -> tuple[int, int] | None:
    """[valid_from_edition, valid_to_edition], or None when rule 12 refuses it."""
    frm, to = row["valid_from_edition"], row["valid_to_edition"]
    if not (_EDITION_RE.fullmatch(frm) and _EDITION_RE.fullmatch(to)) or int(frm) > int(to):
        return None
    return int(frm), int(to)


def _intersect(a: tuple[int, int], b: tuple[int, int]) -> bool:
    return a[0] <= b[1] and b[0] <= a[1]


def _is_counted(row: Mapping[str, str]) -> bool:
    return row["verdict"] == "approved" and row["role"] == "counted"


def _tier(cx: _Cross, source_id: str) -> str | None:
    """The source's tier, or None when it is unknown or not a string (rule 1
    reports a mistyped tier)."""
    tier = cx.sources.get(source_id, {}).get("tier")
    return tier if isinstance(tier, str) else None


def _check_era_ref(item: str, decision_id: str, label: str, cx: _Cross, add: _Add,
                   member: Mapping[str, str] | None) -> None:
    """Rule 20 for one era: ref; `member` is the citing row (None for an
    exclusion, which has no code, account or organization to compare, but
    whose era: refs must still name a same_program or history_only decision)."""
    d = cx.decisions.get(decision_id)
    if d is None:
        add(f"{label} ref {item!r} does not resolve to a decision of the era seed "
            "(p1_era_code_decisions)")
        return
    if d.get("decision") not in _ERA_REF_DECISIONS:
        add(f"{label} ref {item!r} resolves to decision {d.get('decision')!r}, not "
            f"{' or '.join(_ERA_REF_DECISIONS)}")
    if member is None:
        return
    code, account, org = member["code"], member["account"], member["organization"]
    if (d.get("line_item_code"), d.get("account")) != (code, account):
        add(f"{label} ref {item!r} is for line_item_code {d.get('line_item_code')!r}, account "
            f"{d.get('account')!r}; the member's are {code!r}, {account!r}")
    if code in era_map.ORG_SPLIT_CODES:
        if d.get("organization") != org:
            add(f"{label} ref {item!r} has organization {d.get('organization')!r}; code {code!r} is "
                f"organization-split, so it must be the member's {org!r}")
    elif d.get("organization") != "":
        add(f"{label} ref {item!r} has organization {d.get('organization')!r}; it must be blank "
            f"(code {code!r} is not organization-split)")


def _check_src_ref(item: str, source_id: str, label: str, cx: _Cross, add: _Add, *,
                   refuse_x: bool) -> None:
    """Rules 7 and 20 for one src: ref."""
    source = cx.sources.get(source_id)
    if source is None:
        add(f"{label} ref {item!r} names no source in sources[]")
        return
    tier = _tier(cx, source_id)
    if tier == "X" and refuse_x:
        add(f"{label} ref {item!r} cites tier X source {source_id!r}, which no approved row or "
            "exclusion may cite")
    elif tier == "C" and not any(_tier(cx, v) in INDEPENDENT_TIERS
                                 for v in _strs(source, "verified_by")):
        add(f"{label} ref {item!r} cites tier C source {source_id!r}, whose verified_by "
            f"{source.get('verified_by')!r} names no tier G or PG source")


def _check_ref_list(items: Sequence[str], label: str, cx: _Cross, add: _Add, *,
                    member: Mapping[str, str] | None, refuse_x: bool) -> list[tuple[str, dict]]:
    """Rules 7 and 20 on one ref list (a member's column or an exclusion's
    evidence). Every fault quotes the ref; returns the refs that parse."""
    parsed: list[tuple[str, dict]] = []
    prev = None
    for item in items:
        try:
            prefix, parts = parse_ref(item)
        except ValueError as exc:
            add(f"{label}: {exc}{_quote_hint(prev, item)}")
        else:
            parsed.append((prefix, parts))
            if prefix == "era":
                _check_era_ref(item, parts["decision_id"], label, cx, add, member)
            elif prefix == "src":
                _check_src_ref(item, parts["source_id"], label, cx, add, refuse_x=refuse_x)
        prev = item
    return parsed


def _independent(parsed: Sequence[tuple[str, dict]], cx: _Cross) -> bool:
    """§7.5's bar at stage 1: a doc: ref, or a src: ref to a tier G or PG source."""
    return any(prefix == "doc"
               or (prefix == "src" and _tier(cx, parts["source_id"]) in INDEPENDENT_TIERS)
               for prefix, parts in parsed)


def _check_family_links(cx: _Cross, member_rows: Sequence[object],
                        add_at: Callable[[str, str], None]) -> None:
    """Rules 5 (the ruling link), 7 and 20 (exclusions), 23 (a family without
    rows) and 24."""
    with_rows = {row.get("platform_family") for row in member_rows
                 if isinstance(row, Mapping) and isinstance(row.get("platform_family"), str)}
    for pf, fam in cx.families.items():
        add = partial(add_at, pf)
        status = fam.get("status")
        rid = fam.get("published_ruling_id")
        if status == "published":
            if "published_ruling_id" not in fam:
                add("a published family needs published_ruling_id (the ruling that published it)")
            elif isinstance(rid, str):
                ruling = cx.rulings.get(rid)
                if ruling is None:
                    add(f"published_ruling_id {rid!r} is not a ruling in rulings")
                elif ruling.get("platform_family") != pf:
                    add(f"published_ruling_id {rid!r} is a ruling for "
                        f"{ruling.get('platform_family')!r}, not {pf!r}")
        elif status == "draft" and "published_ruling_id" in fam:
            add(f"published_ruling_id {rid!r} is not allowed on a draft family (only ratify moves a "
                "family out of draft)")
        if status in ("ratified", "published") and pf not in with_rows:
            add(f"a {status} family has no member rows (only a draft may have none)")

        exclusions = fam.get("exclusions")
        for j, exclusion in enumerate(exclusions if isinstance(exclusions, list) else ()):
            if isinstance(exclusion, dict):
                _check_ref_list(_strs(exclusion, "evidence"), f"exclusions[{j}] evidence", cx, add,
                                member=None, refuse_x=True)

        notes = fam.get("program_notes")
        if isinstance(notes, dict):   # expected_receipt's exact keys are rule 1's (_check_family_fields)
            counted = {row["member_id"] for _, row in cx.rows
                       if row["platform_family"] == pf and _is_counted(row)}
            for member_id, note in notes.items():
                label = f"program_notes[{member_id!r}]"
                if member_id not in counted:
                    add(f"program_notes key {member_id!r} is not the member_id of an approved "
                        f"counted row of {pf!r}")
                if not isinstance(note, dict):
                    continue
                text, fact_id = note.get("note"), note.get("note_fact_id")
                if isinstance(text, str) and not text.strip():
                    add(f"{label} note {text!r} is blank")
                if isinstance(fact_id, str) and not _FACT_ID_RE.fullmatch(fact_id):
                    add(f"{label} note_fact_id {fact_id!r} is not 16 lower-case hex")


def _check_row_links(cx: _Cross, add_at: Callable[[str, str], None]) -> None:
    """Per member row: rules 7 and 20 (its refs and the independent-signal
    bar), 21, 22 and 23 (a row without a family)."""
    for rid, row in cx.rows:
        add = partial(add_at, rid)
        pf = row["platform_family"]
        if pf not in cx.families:
            add(f"platform_family {pf!r} is not a family in {cx.json_name}")
        approved = row["verdict"] == "approved"
        ruling = cx.rulings.get(row["ruling_id"])

        parsed: dict[str, list[tuple[str, dict]]] = {}
        for col in _REF_COLUMNS:
            items, fault = _split_ref_items(row[col])
            if fault is None:      # a split fault is rule 11's line (_check_member_values)
                parsed[col] = _check_ref_list(items, col, cx, add, member=row, refuse_x=approved)
        # The bar reads the membership evidence only (link_evidence evidences a
        # link); a column that does not split already carries its fault.
        evidence = parsed.get("evidence")
        seed_basis = ruling is not None and ruling.get("seed_basis") is True
        if evidence is not None and _is_counted(row) and not seed_basis \
                and not _independent(evidence, cx):
            add("an approved counted row needs an independent membership signal in evidence: a "
                "doc: ref, or a src: ref to a tier G or PG source (era: and fact: refs do not count; "
                f"its ruling {row['ruling_id']!r} has no seed_basis)")

        for decision_id in filter(None, row["era_decision_ids"].split(";")):
            d = cx.decisions.get(decision_id)
            if d is None:
                add(f"era_decision_ids item {decision_id!r} does not resolve to a decision of the era "
                    "seed (p1_era_code_decisions)")
            elif (d.get("line_item_code"), d.get("account")) != (row["code"], row["account"]):
                add(f"era_decision_ids item {decision_id!r} is for line_item_code "
                    f"{d.get('line_item_code')!r}, account {d.get('account')!r}; the member's are "
                    f"{row['code']!r}, {row['account']!r}")

        if not approved:
            continue
        if row["decided_by"] != "owner":
            add(f"decided_by {row['decided_by']!r} must be 'owner' on an approved row")
        if ruling is None:
            add(f"ruling_id {row['ruling_id']!r} is not a ruling in rulings")
            continue
        if ruling.get("platform_family") != pf:
            add(f"ruling_id {row['ruling_id']!r} is a ruling for {ruling.get('platform_family')!r}, "
                f"not {pf!r}")
        if ruling.get("decided_on") != row["decided_on"]:
            add(f"decided_on {row['decided_on']!r} differs from the decided_on "
                f"{ruling.get('decided_on')!r} of ruling {row['ruling_id']!r}")


def _has_quoted_doc(text: str) -> bool:
    """Rule 19's book_stated test: a doc: ref carrying q= among the refs."""
    for item in _split_ref_items(text)[0]:
        try:
            prefix, parts = parse_ref(item)
        except ValueError:
            continue
        if prefix == "doc" and parts["quote"] is not None:
            return True
    return False


def _check_family_rows(pf: str, rows: _Rows, add_at: Callable[[str, str], None]) -> None:
    """Rules 13, 15, 16 and 19 over one family's rows; a fault between two
    rows goes on the later one by CSV order."""
    member_ids = {row["member_id"] for _, row in rows}
    n = sum(1 for _, row in rows if _is_counted(row))
    spans: dict[str, list[tuple[str, tuple[int, int]]]] = defaultdict(list)
    codes: dict[str, tuple[str, str]] = {}       # code -> (member_id, row id) of its first counted row
    orders: dict[int, str] = {}                  # member_order -> row id
    for rid, row in rows:
        add = partial(add_at, rid)
        member_id, span = row["member_id"], _span(row)
        if span is not None:
            for other_rid, other in spans[member_id]:
                if _intersect(span, other):
                    add(f"member_id {member_id!r} overlaps row {other_rid!r}: editions "
                        f"{span[0]}–{span[1]} and {other[0]}–{other[1]} intersect")
            spans[member_id].append((rid, span))

        order = row["member_order"]
        if _is_counted(row):
            first_id, first_rid = codes.setdefault(row["code"], (member_id, rid))
            if first_id != member_id:
                add(f"code {row['code']!r} is also the code of approved counted member {first_id!r} "
                    f"(row {first_rid!r}); two counted members of one family may not share a code")
            if order == "" or _POSITIVE_RE.fullmatch(order):    # _check_editions reports any other
                value = int(order) if order else 0
                if not 1 <= value <= n:
                    add(f"member_order {order!r} is not in 1..{n} (contiguous over the {n} approved "
                        f"counted rows of {pf!r})")
                elif value in orders:
                    add(f"member_order {order!r} repeats row {orders[value]!r}")
                else:
                    orders[value] = rid
        elif order and row["role"] != "context":    # a context row's member_order is rule 17's line
            add(f"member_order {order!r} is set on a {row['verdict']} {row['role']} row; only approved "
                "counted rows carry one")

        for col in ("predecessor_member_id", "successor_member_id"):
            if row[col] and row[col] not in member_ids:
                add(f"{col} {row[col]!r} is not the member_id of a row of {pf!r}")
        if (row["predecessor_member_id"] or row["successor_member_id"]) and row["link_basis"] == "":
            add("link_basis is required when predecessor_member_id or successor_member_id is set "
                "(got '')")
        if row["link_basis"] == "book_stated" and not _has_quoted_doc(row["link_evidence"]):
            add(f"link_basis 'book_stated' needs a doc: ref with a quote (q=) in link_evidence "
                f"(got {row['link_evidence']!r})")


def _check_row_sets(cx: _Cross, add_at: Callable[[str, str], None]) -> None:
    """Rules 13, 15, 16, 19 within each family; rule 14 across families, on
    the later family's row (JSON order), naming the earlier family."""
    by_family: dict[str, _Rows] = defaultdict(list)
    for rid, row in cx.rows:
        by_family[row["platform_family"]].append((rid, row))
    for pf, rows in by_family.items():
        _check_family_rows(pf, rows, add_at)

    counted: dict[str, list[tuple[str, str, tuple[int, int]]]] = defaultdict(list)
    for pf in cx.families:
        mine = []
        for rid, row in by_family.get(pf, ()):
            span = _span(row)
            if not _is_counted(row) or span is None:
                continue
            for other_pf, other_rid, other in counted[row["member_id"]]:
                if _intersect(span, other):
                    add_at(rid, f"member_id {row['member_id']!r} is also counted in family "
                                f"{other_pf!r} (row {other_rid!r}, editions {other[0]}–{other[1]}); "
                                "a member is counted in one family only")
            mine.append((row["member_id"], (pf, rid, span)))
        for member_id, entry in mine:
            counted[member_id].append(entry)


def _check_era_links(cx: _Cross, era_rows: _Rows, add_at: Callable[[str, str], None]) -> None:
    """Rule 25."""
    counted: dict[tuple[str, str], list[tuple[str, Mapping[str, str], tuple[int, int]]]] = \
        defaultdict(list)
    for rid, row in cx.rows:
        span = _span(row)
        if _is_counted(row) and span is not None:
            counted[(row["platform_family"], row["member_id"])].append((rid, row, span))

    seen: set[tuple[str, str, str]] = set()
    for rid, row in era_rows:
        add = partial(add_at, rid)
        pf, member_id, era_key = row["platform_family"], row["member_id"], row["era_key"]
        key = (pf, row["edition"], era_key)
        if key in seen:
            add(f"duplicate era line (platform_family, edition, era_key) {key!r}")
        seen.add(key)
        m = era_keys._ERA_KEY_RE.fullmatch(era_key)
        if m is None:
            add(f"era_key {era_key!r} does not match {era_keys._ERA_KEY_RE.pattern}")
        rows = counted.get((pf, member_id), [])
        if not rows:
            add(f"member {member_id!r} is not an approved counted row of {pf!r}")
            continue
        if not _EDITION_RE.fullmatch(row["edition"]):
            continue                                  # rule 12's line (_check_era_rows)
        edition = int(row["edition"])
        member = next((entry for entry in rows if entry[2][0] <= edition <= entry[2][1]), None)
        if member is None:
            spans = ", ".join(f"{s[0]}–{s[1]}" for _, _, s in rows)
            add(f"edition {edition} is outside the editions of approved counted member "
                f"{member_id!r} ({spans})")
            continue
        member_rid, member_row, _ = member
        if m is not None:
            if m.group(1) != member_row["account"]:
                add(f"era_key {era_key!r} has account {m.group(1)!r}; the member's account is "
                    f"{member_row['account']!r}")
            if m.group(2) != member_row["organization"]:
                add(f"era_key {era_key!r} has organization {m.group(2)!r}; the member's organization "
                    f"is {member_row['organization']!r}")
        if row["decision_id"] not in set(filter(None, member_row["era_decision_ids"].split(";"))):
            add(f"decision_id {row['decision_id']!r} is not in the era_decision_ids of member row "
                f"{member_rid!r}")

    # One (edition, era_key) bound in two families: the later family's row
    # (JSON order) carries the fault. sorted() is stable, so CSV order holds
    # within a family.
    position = {pf: i for i, pf in enumerate(cx.families)}
    first: dict[tuple[str, str], str] = {}
    for rid, row in sorted((r for r in era_rows if r[1]["platform_family"] in position),
                           key=lambda r: position[r[1]["platform_family"]]):
        bound = (row["edition"], row["era_key"])
        earlier = first.setdefault(bound, row["platform_family"])
        if earlier != row["platform_family"]:
            add_at(rid, f"(edition, era_key) {bound!r} is already bound by family {earlier!r}; an era "
                        "line is counted in one family only")


# ---------------------------------------------------------------------------
# parse / load
# ---------------------------------------------------------------------------


def _line(file: str, row: str | None, reason: str) -> str:
    return f"platform-families: {file}:{row}: {reason}" if row is not None \
        else f"platform-families: {file}: {reason}"


def _family(fam: dict, members: Sequence[Member], era_lines: Sequence[EraLine]) -> Family:
    headline, guard = fam["headline"], fam["review_guard"]
    return Family(
        platform_family=fam["platform_family"], status=fam["status"],
        published_ruling_id=fam.get("published_ruling_id"), short_name=fam["short_name"],
        history_dataset=fam["history_dataset"], on_partial_corpus=fam["on_partial_corpus"],
        headline=Headline(headline["measure"], headline["first_fy"], headline["last_fy"]),
        scope_note=fam["scope_note"], coverage_notes=tuple(fam["coverage_notes"]),
        review_guard=ReviewGuard(tuple(guard["exhibits"]), tuple(guard["title_substrings"])),
        program_notes={mid: ProgramNote(n["note"], n["note_fact_id"], dict(n["expected_receipt"]))
                       for mid, n in fam["program_notes"].items()},
        variants=tuple(dict(v) for v in fam["variants"]),
        program_tags=tuple(dict(t) for t in fam["program_tags"]),
        purpose_tags=tuple(fam["purpose_tags"]), caveat_ids=tuple(fam["caveat_ids"]),
        exclusions=tuple(copy.deepcopy(e) for e in fam["exclusions"]),
        aliases=tuple(fam["aliases"]),
        contract_routes=copy.deepcopy(fam["contract_routes"]),
        official_estimates=copy.deepcopy(fam["official_estimates"]),
        art=copy.deepcopy(fam["art"]),
        members=tuple(members), era_lines=tuple(era_lines))


def _era_line(row: Mapping[str, str]) -> EraLine:
    return EraLine(row["platform_family"], row["member_id"], int(row["edition"]), row["era_key"],
                   row["budget_activity"], row["filed_title"], row["decision_id"])


def parse_registry(doc: dict, member_rows: Sequence[Mapping[str, str]],
                   era_rows: Sequence[Mapping[str, str]], *,
                   era_decisions: Mapping[str, Mapping[str, str]],
                   paths: RegistryPaths = DEFAULT_PATHS) -> Registry:
    """Stage 1 on in-memory files (pure): the per-row rules and the rules
    across rows and files (rules 1–27) all run, every fault is collected (the
    JSON's lines, then the members', then the era lines'), then one
    RegistryError is raised; otherwise the Registry is built. `era_decisions`
    is the era seed keyed by decision_id."""
    json_name, members_name, era_name = paths.json.name, paths.members.name, paths.era_lines.name
    json_errors: list[str] = []
    member_errors: list[str] = []
    era_errors: list[str] = []

    def add_json(row: str | None, reason: str) -> None:
        json_errors.append(_line(json_name, row, reason))

    def add_member(row: str, reason: str) -> None:
        member_errors.append(_line(members_name, row, reason))

    def add_era(row: str, reason: str) -> None:
        era_errors.append(_line(era_name, row, reason))

    families = _check_doc(doc, add_json)
    rows = _check_members(member_rows, families, add_member)
    era = _check_era_rows(era_rows, add_era)
    top = doc if isinstance(doc, dict) else {}
    cx = _Cross(json_name=json_name, rulings=_by_id(top.get("rulings"), "ruling_id"),
                sources=_by_id(top.get("sources"), "source_id"), families=families, rows=rows,
                decisions=era_decisions)
    _check_family_links(cx, member_rows, add_json)
    _check_row_links(cx, add_member)
    _check_row_sets(cx, add_member)
    _check_era_links(cx, era, add_era)
    errors = json_errors + member_errors + era_errors
    if errors:
        raise RegistryError(errors)

    members: dict[str, list[Member]] = defaultdict(list)
    for row in member_rows:
        members[row["platform_family"]].append(Member.from_row(row))
    era_lines: dict[str, list[EraLine]] = defaultdict(list)
    for row in era_rows:
        era_lines[row["platform_family"]].append(_era_line(row))
    sources = {}
    for s in doc["sources"]:
        sources[s["source_id"]] = Source(
            source_id=s["source_id"], title=s["title"], publisher=s["publisher"], url=s["url"],
            as_of=s["as_of"], sha256=s.get("sha256"), bytes=s.get("bytes"), archived=s["archived"],
            archive_path=s.get("archive_path"), tier=s["tier"],
            verified_by=tuple(s.get("verified_by", ())))
    return Registry(
        schema_version=doc["schema_version"],
        rulings={r["ruling_id"]: Ruling(**{k: r[k] for k in RULING_KEYS}) for r in doc["rulings"]},
        sources=sources,
        families=tuple(_family(f, members[f["platform_family"]], era_lines[f["platform_family"]])
                       for f in doc["families"]),
        doc=copy.deepcopy(doc))


def _read_text(path: Path) -> tuple[str | None, list[str]]:
    """The file as text, or None with its byte-level faults."""
    if not path.is_file():
        return None, [f"missing (no file at {path})"]
    data = path.read_bytes()
    faults = []
    if data.startswith(codecs.BOM_UTF8):
        faults.append("starts with a UTF-8 BOM")
    text = None
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError as exc:
        faults.append(f"is not UTF-8 ({exc.reason} at byte {exc.start})")
    if b"\r" in data:
        faults.append("contains a carriage return (CRLF or CR line endings)")
    if not data:
        faults.append("is empty")
    elif not data.endswith(b"\n"):
        faults.append("does not end with a newline")
    elif data.endswith(b"\n\n"):
        faults.append("ends with a trailing blank line")
    return (None if faults else text), faults


def _reject_constant(name: str) -> object:
    raise ValueError(f"{name} is not JSON")


def _parse_json(text: str) -> tuple[object, list[str]]:
    duplicates: list[str] = []

    def pairs_hook(pairs: list[tuple[str, object]]) -> dict:
        obj: dict = {}
        for key, value in pairs:
            if key in obj:
                duplicates.append(key)
            obj[key] = value
        return obj

    try:
        doc = json.loads(text, object_pairs_hook=pairs_hook, parse_constant=_reject_constant)
    except json.JSONDecodeError as exc:
        return None, [f"is not valid JSON ({exc.msg} at line {exc.lineno} column {exc.colno})"]
    except ValueError as exc:
        return None, [f"is not valid JSON ({exc})"]
    if duplicates:
        return None, [f"duplicate key {key!r} in a JSON object" for key in duplicates]
    return doc, []


def _header_fault(header: list[str], columns: Sequence[str]) -> str | None:
    if tuple(header) == tuple(columns):
        return None
    if len(header) != len(columns):
        return f"header has {len(header)} columns, the frozen header has {len(columns)}"
    i = next(i for i, (a, b) in enumerate(zip(header, columns)) if a != b)
    return (f"header column {i + 1} is {header[i]!r}, the frozen header has {columns[i]!r} "
            f"(the {len(columns)}-column tuple, in order)")


def _parse_csv(text: str, columns: Sequence[str], file: str,
               row_ref: Callable[[list[str]], str | None]) -> tuple[list[dict[str, str]] | None, list[str]]:
    """Rows over the frozen header, or None with the file's lines. Fields are
    kept exactly as written (no skipinitialspace, no strip)."""
    reader = csv.reader(io.StringIO(text, newline=""), strict=True)
    rows: list[dict[str, str]] = []
    faults: list[str] = []
    try:
        header = next(reader, [])
        fault = _header_fault(header, columns)
        if fault:
            return None, [_line(file, None, fault)]
        for fields in reader:
            if len(fields) != len(columns):
                ref = row_ref(fields) or f"line {reader.line_num}"
                faults.append(_line(file, ref, f"{len(fields)} fields, the header has {len(columns)}"))
                continue
            rows.append(dict(zip(columns, fields)))
    except csv.Error as exc:
        return None, [_line(file, None, f"is not valid CSV (line {reader.line_num}: {exc})")]
    return (None if faults else rows), faults


def _member_ref(fields: list[str]) -> str | None:
    return fields[0] if fields and fields[0] else None


def _era_ref(fields: list[str]) -> str | None:
    return f"{fields[0]}|{fields[2]}|{fields[3]}" if len(fields) >= 4 else None


def load_registry(paths: RegistryPaths = DEFAULT_PATHS, *,
                  era_decisions_path: Path | None = None) -> Registry:
    """Read the three files and run stage 1. File-level faults (missing, BOM,
    not UTF-8, carriage returns, empty, no final newline, a trailing blank
    line, bad JSON, duplicate JSON keys, a CSV header other than the frozen
    tuple, a row whose field count differs from the header's) are all raised
    together before parse_registry runs."""
    errors: list[str] = []
    texts: dict[Path, str | None] = {}
    for path in (paths.json, paths.members, paths.era_lines):
        text, faults = _read_text(path)
        errors.extend(_line(path.name, None, fault) for fault in faults)
        texts[path] = text
    doc: object = None
    if texts[paths.json] is not None:
        doc, faults = _parse_json(texts[paths.json])
        errors.extend(_line(paths.json.name, None, fault) for fault in faults)
    member_rows = era_rows = None
    if texts[paths.members] is not None:
        member_rows, faults = _parse_csv(texts[paths.members], MEMBER_COLUMNS, paths.members.name,
                                         _member_ref)
        errors.extend(faults)
    if texts[paths.era_lines] is not None:
        era_rows, faults = _parse_csv(texts[paths.era_lines], ERA_LINE_COLUMNS,
                                      paths.era_lines.name, _era_ref)
        errors.extend(faults)
    if errors:
        raise RegistryError(errors)
    seed = era_map.read_seed(era_decisions_path or era_map.default_seed_path())
    decisions = {row["decision_id"]: row for row in seed}
    return parse_registry(doc, member_rows, era_rows, era_decisions=decisions, paths=paths)


# ---------------------------------------------------------------------------
# Writers (D11) and the family hash (D10)
# ---------------------------------------------------------------------------


def _ordered(obj: object, keys: Sequence[str], label: str) -> dict:
    if not isinstance(obj, dict):
        raise ValueError(f"{label} must be an object, got {obj!r}")
    for key in obj:
        if key not in keys:
            raise ValueError(f"{label}: unknown key {key!r}")
    return {key: obj[key] for key in keys if key in obj}


def _canonical_family(fam: object) -> dict:
    out = _ordered(fam, FAMILY_KEYS, "families[]")
    for key, keys in (("headline", HEADLINE_KEYS), ("review_guard", GUARD_KEYS),
                      ("contract_routes", CONTRACT_ROUTE_KEYS),
                      ("official_estimates", OFFICIAL_ESTIMATE_KEYS), ("art", ART_KEYS)):
        if out.get(key) is not None:
            out[key] = _ordered(out[key], keys, key)
    for key, keys in (("variants", VARIANT_KEYS), ("program_tags", PROGRAM_TAG_KEYS),
                      ("exclusions", EXCLUSION_KEYS)):
        if key in out:
            out[key] = [_ordered(item, keys, f"{key}[]") for item in out[key]]
    if "program_notes" in out:
        notes = {}
        for member_id, note in sorted(out["program_notes"].items()):
            note = _ordered(note, PROGRAM_NOTE_KEYS, "program_notes")
            if "expected_receipt" in note:
                note["expected_receipt"] = _ordered(note["expected_receipt"], RECEIPT_KEYS,
                                                    "expected_receipt")
            notes[member_id] = note
        out["program_notes"] = notes
    return out


def _canonical_doc(doc: dict) -> dict:
    """Every object's keys in its key tuple's order (program_notes by
    member_id); families and list items keep their given order."""
    out = _ordered(doc, TOP_KEYS, "the top level")
    if "rulings" in out:
        out["rulings"] = [_ordered(r, RULING_KEYS, "rulings[]") for r in out["rulings"]]
    if "sources" in out:
        out["sources"] = [_ordered(s, SOURCE_KEYS, "sources[]") for s in out["sources"]]
    if "families" in out:
        out["families"] = [_canonical_family(f) for f in out["families"]]
    return out


def registry_texts(paths: RegistryPaths, *, doc: dict, members: Sequence[Mapping[str, str]],
                   era_lines: Sequence[Mapping[str, str]]) -> dict[Path, str]:
    """The three registry files' texts (D11). Does not validate: callers run
    parse_registry first."""
    return {
        paths.json: json.dumps(_canonical_doc(doc), indent=2, ensure_ascii=False) + "\n",
        paths.members: era_map._csv_text(MEMBER_COLUMNS,
                                         sorted(members, key=lambda r: r["member_row_id"])),
        paths.era_lines: era_map._csv_text(
            ERA_LINE_COLUMNS,
            sorted(era_lines, key=lambda r: (r["platform_family"], r["edition"], r["era_key"]))),
    }


def write_registry(paths: RegistryPaths, *, doc: dict, members: Sequence[Mapping[str, str]],
                   era_lines: Sequence[Mapping[str, str]]) -> None:
    """The three files in one era_map._atomic_write_texts call."""
    era_map._atomic_write_texts(registry_texts(paths, doc=doc, members=members, era_lines=era_lines))


def family_hash(doc: dict, platform_family: str) -> str:
    """D10: sha256 over the family object as written plus the sources its
    exclusions cite by `src:` ref (sorted by source_id)."""
    fam = next((f for f in doc["families"] if f.get("platform_family") == platform_family), None)
    if fam is None:
        raise KeyError(platform_family)
    cited = {ref[len("src:"):].split("#", 1)[0]
             for exclusion in fam.get("exclusions") or ()
             for ref in exclusion.get("evidence") or ()
             if ref.startswith("src:")}
    sources = sorted((s for s in doc.get("sources") or () if s.get("source_id") in cited),
                     key=lambda s: s["source_id"])
    text = json.dumps({"family": fam, "sources": sources}, sort_keys=True, separators=(",", ":"),
                      ensure_ascii=False)
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


# ---------------------------------------------------------------------------
# Stage 2 (spec §9.2): the lake, read-only
# ---------------------------------------------------------------------------


class LakeReader:
    """Read-only queries over the warehouse and its jbooks parquet (open_lake).

    No query filters editions, so a PB2027 row reaches stage 2 and fails item
    1 (R2; §17 risk 7). Every budget_lines query reads its edition as
    cast(fiscal_year as integer); results are distinct and sorted. A
    duckdb.IOException (another writer's lock) is never caught here.
    """

    def __init__(self, con, budget_lines: Path, documents: Path):
        self._con = con
        self.budget_lines = budget_lines
        self.documents = documents

    def __enter__(self) -> LakeReader:
        return self

    def __exit__(self, *exc: object) -> None:
        self.close()

    def close(self) -> None:
        self._con.close()

    def _budget_lines_sql(self) -> str:
        return f"read_parquet('{era_map._sql_path(self.budget_lines)}')"

    def _dicts(self, sql: str, params: Sequence[object] = ()) -> list[dict]:
        cur = self._con.execute(sql, list(params))
        names = [d[0] for d in cur.description]
        return [dict(zip(names, row)) for row in cur.fetchall()]

    def code_rows(self, exhibit: str, account: str, organizations: set[str], code: str) -> list[dict]:
        """The code-keyed lines printing `code` (pe_bli) in the exhibit and
        account under any of the organizations: {edition, organization,
        budget_activity, title}, by (edition, budget_activity, title,
        organization)."""
        orgs = sorted(organizations)
        if not orgs:
            return []
        return self._dicts(
            "select distinct cast(fiscal_year as integer) as edition, organization, budget_activity, title"
            f" from {self._budget_lines_sql()}"
            " where exhibit = ? and account = ? and pe_bli = ?"
            f" and organization in ({', '.join('?' * len(orgs))})"
            " order by edition, budget_activity, title, organization",
            [exhibit, account, code, *orgs])

    def era_map_rows(self, code: str, account: str, organization: str | None) -> list[dict]:
        """p1_era_line_map lines printing `code` in the account (and the
        organization, when given), every decision, by (edition, era_key)."""
        sql = ("select distinct cast(edition as integer) as edition, account, organization,"
               " budget_activity, era_key, line_item_code, filed_title, program_key, decision,"
               " decision_id from p1_era_line_map where line_item_code = ? and account = ?")
        params: list[object] = [code, account]
        if organization is not None:
            sql += " and organization = ?"
            params.append(organization)
        return self._dicts(sql + " order by edition, era_key, decision_id", params)

    def guard_rows(self, guard: ReviewGuard) -> list[dict]:
        """The lines of the guard's exhibits whose upper-cased title contains
        any guard substring: {edition, exhibit, account, organization,
        budget_activity, pe_bli, title, line_item_code}, by (edition, exhibit,
        pe_bli, account, organization)."""
        exhibits, subs = sorted(guard.exhibits), list(guard.title_substrings)
        if not exhibits or not subs:
            return []
        return self._dicts(
            "select distinct cast(fiscal_year as integer) as edition, exhibit, account, organization,"
            " budget_activity, pe_bli, title, line_item_code"
            f" from {self._budget_lines_sql()}"
            f" where exhibit in ({', '.join('?' * len(exhibits))})"
            f" and ({' or '.join(['contains(upper(title), ?)'] * len(subs))})"
            " order by edition, exhibit, pe_bli, account, organization, budget_activity, title,"
            " line_item_code",
            [*exhibits, *subs])

    def collision_codes(self) -> frozenset[str]:
        """The codes dim_programs reuses across accounts or organizations
        (collision_keys.classify_shared_keys)."""
        from govbudget.jbooks.collision_keys import classify_shared_keys

        rows = self._con.execute("select pe_bli, account, org from dim_programs").fetchall()
        return frozenset(classify_shared_keys(rows))

    def series_codes(self) -> frozenset[str]:
        """Every pe_bli of fct_decade_series; empty when the table is absent
        (a small export's warehouse)."""
        import duckdb

        try:
            rows = self._con.execute("select distinct pe_bli from fct_decade_series").fetchall()
        except duckdb.CatalogException:
            return frozenset()
        return frozenset(pe for (pe,) in rows)

    def actuals(self, keys: Iterable[tuple[str, int]]) -> dict[tuple[str, int], Decimal]:
        """Per requested (pe_bli, edition): the summed fct_decade_series
        actuals at fy = edition - 2 of that edition. A key with no row is
        absent (so is every key when the table is absent, as series_codes)."""
        import duckdb

        wanted = set(keys)
        if not wanted:
            return {}
        codes = sorted({pe for pe, _ in wanted})
        try:
            rows = self._con.execute(
                "select pe_bli, edition_year, sum(amount) from fct_decade_series"
                " where amount_type_kind = 'actuals' and fy = edition_year - 2"
                f" and pe_bli in ({', '.join('?' * len(codes))})"
                " group by pe_bli, edition_year order by pe_bli, edition_year", codes).fetchall()
        except duckdb.CatalogException:
            return {}
        return {(pe, int(ed)): Decimal(str(v)) for pe, ed, v in rows
                if (pe, int(ed)) in wanted and v is not None}

    def program_series(self, code: str) -> dict[tuple[int, str], Decimal]:
        """fct_program_decade_series for one program_key under the edition
        rule (actuals from edition fy + 2, enacted from fy + 1, request from
        fy), summed per (fy, amount_type_kind) in $K; {} when nothing matches
        or the table is absent."""
        import duckdb

        try:
            rows = self._con.execute(
                "select fy, amount_type_kind, sum(amount_thousands) from fct_program_decade_series"
                " where program_key = ? and ((amount_type_kind = 'actuals' and edition_year = fy + 2)"
                " or (amount_type_kind = 'enacted' and edition_year = fy + 1)"
                " or (amount_type_kind = 'request' and edition_year = fy))"
                " group by fy, amount_type_kind order by fy, amount_type_kind", [code]).fetchall()
        except duckdb.CatalogException:
            return {}
        return {(int(fy), kind): Decimal(str(v)) for fy, kind, v in rows if v is not None}


def open_lake(duckdb_path: Path) -> LakeReader:
    """One read-only DuckDB connection (it reads the jbooks parquet too). A
    second writer's lock surfaces as duckdb.IOException, uncaught."""
    import duckdb

    from govbudget.export_site import _stage_parquet_path

    budget_lines = _stage_parquet_path(duckdb_path, "jbooks", "budget_lines.parquet")
    documents = _stage_parquet_path(duckdb_path, "jbooks", "documents.parquet")
    if budget_lines is None or documents is None:
        base = Path(duckdb_path).parent
        roots = (base / "parquet" / "jbooks", base.parent / "parquet" / "jbooks",
                 base.parent / "data" / "parquet" / "jbooks")
        raise FileNotFoundError(
            f"platform-families: budget_lines.parquet and documents.parquet must both be beside "
            f"{duckdb_path} (looked in {', '.join(map(str, roots))})")
    con = duckdb.connect(str(duckdb_path), read_only=True)
    return LakeReader(con, budget_lines, documents)


@dataclass(frozen=True)
class MemberMeasure:
    """A member row's lake-derived columns (§7.3 rows 12, 13 and 21), its era
    lines and its headline actuals ($K)."""

    titles_seen: tuple[tuple[int, str], ...]
    budget_activities: tuple[str, ...]
    editions: tuple[int, ...]
    drift_sha256: str
    era_lines: tuple[EraLine, ...]
    headline_actuals_k: Decimal | None


def _era_org(member: Member) -> str | None:
    """The organization filter of the member's era lines: its own for an
    organization-split code, none otherwise."""
    return member.organization if member.code in era_map.ORG_SPLIT_CODES else None


def _member_lines(lake: LakeReader, member: Member) -> tuple[list[dict], list[dict]]:
    """(the code-keyed lines, every p1_era_line_map line printing the code)."""
    orgs = {member.organization, era_map.ERA_ORG_TO_MODERN.get(member.organization, member.organization)}
    return (lake.code_rows(member.exhibit, member.account, orgs, member.code),
            lake.era_map_rows(member.code, member.account, _era_org(member)))


def _in_editions(member: Member, edition: int) -> bool:
    return member.valid_from_edition <= edition <= member.valid_to_edition


def _bound_era_rows(member: Member, map_rows: Sequence[dict]) -> list[dict]:
    """The map lines under the member's era_decision_ids, any edition."""
    return [r for r in map_rows if r["decision_id"] in member.era_decision_ids]


def _drift_keys(member: Member, code_lines: Sequence[dict],
                era_rows: Sequence[dict]) -> set[tuple[int, str, str, str]]:
    """§7.3 row 21: the distinct lines the row binds within its editions."""
    keys = {(r["edition"], member.code, r["budget_activity"], r["title"])
            for r in code_lines if _in_editions(member, r["edition"])}
    keys |= {(r["edition"], r["era_key"], r["budget_activity"], r["filed_title"])
             for r in era_rows if _in_editions(member, r["edition"])}
    return keys


def _measure(lake: LakeReader, member: Member, headline: Headline, code_lines: Sequence[dict],
             map_rows: Sequence[dict]) -> MemberMeasure:
    era_rows = _bound_era_rows(member, map_rows)
    inside_code = [r for r in code_lines if _in_editions(member, r["edition"])]
    inside_era = [r for r in era_rows if _in_editions(member, r["edition"])]
    keys = _drift_keys(member, code_lines, era_rows)
    amounts = lake.actuals({(line, edition) for edition, line, _ba, _title in keys
                            if headline.first_fy <= edition - 2 <= headline.last_fy})
    return MemberMeasure(
        titles_seen=tuple(sorted({(r["edition"], r["title"]) for r in inside_code})),
        budget_activities=tuple(sorted({r["budget_activity"] for r in inside_code + inside_era})),
        editions=tuple(sorted({r["edition"] for r in [*code_lines, *era_rows]})),
        drift_sha256=era_map.keys_sha256(keys),
        era_lines=tuple(EraLine(member.platform_family, member.member_id, r["edition"], r["era_key"],
                                r["budget_activity"], r["filed_title"], r["decision_id"])
                        for r in era_rows),
        headline_actuals_k=sum(amounts.values(), Decimal(0)) if amounts else None)


def measure_member(lake: LakeReader, member: Member, *, headline: Headline) -> MemberMeasure:
    """Measure one member row against the lake (propose fills its lake
    columns from this; stage 2 item 2 recomputes its drift_sha256)."""
    return _measure(lake, member, headline, *_member_lines(lake, member))


# doc: quotes (§7.5, Review Focus 2): NFKC, then every hyphen/dash and the
# minus sign to '-', curly single and double quotes to straight ones, then
# whitespace runs to one space. Case is kept.
_QUOTE_FOLD = str.maketrans({
    **dict.fromkeys(map(chr, range(0x2010, 0x2016)), "-"), "−": "-",
    **dict.fromkeys("‘’‚‛", "'"),
    **dict.fromkeys("“”„‟", '"'),
})
_PAGE_TEXTS: dict[Path, tuple[str, ...]] = {}       # page texts per file, for the life of the process


def normalize_quote_text(text: str) -> str:
    """The form both a doc: quote and its page text take before the
    substring test."""
    return " ".join(unicodedata.normalize("NFKC", text).translate(_QUOTE_FOLD).split())


def _workbook_page_texts(path: Path) -> tuple[str, ...]:
    """Page n of a workbook is its n-th worksheet: every non-empty cell's
    value, row-major, joined by one space."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return tuple(" ".join(str(v) for row in ws.iter_rows(values_only=True) for v in row
                              if v is not None and str(v) != "")
                     for ws in wb.worksheets)
    finally:
        wb.close()


def _doc_page_texts(path: Path) -> tuple[str, ...]:
    if path not in _PAGE_TEXTS:
        if path.suffix == ".pdf":
            from govbudget.jbooks.provenance_pages import _page_texts

            _PAGE_TEXTS[path] = tuple(_page_texts(path))
        else:
            _PAGE_TEXTS[path] = _workbook_page_texts(path)
    return _PAGE_TEXTS[path]


def resolve_doc_ref(ref: str, *, site_dir: Path) -> str | None:
    """None when the doc: ref resolves (§7.5): the sha-pinned PDF
    (pdfs/<sha>.pdf) or workbook (workbooks/<sha>.xlsx) exists, the page is
    in range and the quote, if any, is on that page once both are normalized
    (normalize_quote_text). Otherwise the reason."""
    try:
        prefix, parts = parse_ref(ref)
    except ValueError as exc:
        return str(exc)
    if prefix != "doc":
        return f"{ref!r} is not a doc: ref"
    sha, page, quote = parts["sha256"], parts["page"], parts["quote"]
    head = f"doc:{sha}#p{page}"
    site = Path(site_dir)
    pdf, workbook = site / "pdfs" / f"{sha}.pdf", site / "workbooks" / f"{sha}.xlsx"
    if pdf.is_file():
        texts = _doc_page_texts(pdf)
    elif workbook.is_file():
        texts = _doc_page_texts(workbook)
    else:
        return f"{head}: no file at {pdf} or {workbook}"
    if not 1 <= page <= len(texts):
        return f"{head}: page out of range (1–{len(texts)})"
    if quote is not None and normalize_quote_text(quote) not in normalize_quote_text(texts[page - 1]):
        return f"{head}: quote not found: {quote!r}"
    return None


def corpus_gate(family: Family, series_codes: frozenset[str]) -> bool:
    """D5, the partial-corpus gate: True when every code-keyed counted code is
    in fct_decade_series; otherwise False for on_partial_corpus 'skip', and a
    RegistryError for 'raise'."""
    missing = sorted({m.code for m in family.code_keyed} - series_codes)
    if not missing:
        return True
    if family.on_partial_corpus == "skip":
        return False
    raise RegistryError([_line(REGISTRY_JSON.name, family.platform_family,
                               f"{family.short_name} family history requires every counted workbook "
                               f"member: missing {', '.join(missing)}")])


@dataclass
class LakeReport:
    """Stage 2's result (R4): failures and stale block; undecided never does."""

    failures: list[str]
    undecided: list[str]
    stale: list[str]

    def raise_if_failed(self) -> None:
        if self.failures or self.stale:
            raise RegistryError(self.failures + self.stale)


def _guard_text(row: Mapping[str, object]) -> str:
    return (f"PB{row['edition']} {row['exhibit']} {row['account']}/{row['organization']} "
            f"{row['pe_bli']} {row['title']!r}")


def _covers(member: Member, row: Mapping[str, object]) -> bool:
    """Item 5: the member row covers the guard-matching lake row. An era-keyed
    row is matched by its printed line_item_code, a code-keyed one by pe_bli;
    the organization may be the member's spelling or its modern one."""
    pe_bli = row["pe_bli"]
    code = row["line_item_code"] if era_keys.is_era_procurement_key(pe_bli) else pe_bli
    orgs = {member.organization, era_map.ERA_ORG_TO_MODERN.get(member.organization, member.organization)}
    return ((member.exhibit, member.account, member.code) == (row["exhibit"], row["account"], code)
            and row["organization"] in orgs and _in_editions(member, row["edition"]))


def _refs(member: Member) -> list[tuple[str, str]]:
    """(column, ref) for every ref a member row carries."""
    alloc = _split_ref_items(member.allocation_source)[0]
    return ([("evidence", r) for r in member.evidence] + [("link_evidence", r) for r in member.link_evidence]
            + [("allocation_source", r) for r in alloc])


def _exclusion_refs(family: Family) -> list[tuple[int, str]]:
    return [(j, ref) for j, exclusion in enumerate(family.exclusions)
            for ref in exclusion.get("evidence") or ()]


def _none_as_empty(key: tuple) -> tuple:
    """A sort key for era-line tuples whose text fields may be NULL."""
    return tuple("" if v is None else v for v in key)


class _FamilyLake:
    """check_lake for one family: report lines in Task 2a's format."""

    def __init__(self, family: Family, report: LakeReport):
        self.family = family
        self.report = report

    def json_failure(self, reason: str, row: str | None = None) -> None:
        self.report.failures.append(_line(REGISTRY_JSON.name, row or self.family.platform_family, reason))

    def member_failure(self, member: Member, reason: str) -> None:
        self.report.failures.append(_line(MEMBERS_CSV.name, member.member_row_id, reason))

    def era_failure(self, edition: int, era_key: str, reason: str) -> None:
        row = f"{self.family.platform_family}|{edition}|{era_key}"
        self.report.failures.append(_line(ERA_LINES_CSV.name, row, reason))

    def doc_refs(self, site_dir: Path, members: Sequence[Member]) -> None:
        """Item 7."""
        for member in members:
            for col, ref in _refs(member):
                if ref.startswith("doc:"):
                    reason = resolve_doc_ref(ref, site_dir=site_dir)
                    if reason is not None:
                        self.member_failure(member, f"{col}: {reason}")
        for j, ref in _exclusion_refs(self.family):
            if ref.startswith("doc:"):
                reason = resolve_doc_ref(ref, site_dir=site_dir)
                if reason is not None:
                    self.json_failure(f"exclusions[{j}] evidence: {reason}")

    def code_keyed(self, member: Member, code_lines: Sequence[dict]) -> None:
        """Item 1."""
        expected = dict(member.titles_seen)
        printed = {r["edition"] for r in code_lines}
        reasons: list[str] = []
        for edition, title in member.titles_seen:
            if edition not in printed:
                reasons.append(f"titles_seen lists PB{edition} {title!r}, but the lake has no code-keyed "
                               f"{member.exhibit} {member.account} {member.code} line in PB{edition} "
                               "(missing)")
        for r in code_lines:
            edition, title = r["edition"], r["title"]
            if not _in_editions(member, edition):
                reasons.append(f"code-keyed line PB{edition} {member.code} {title!r} is outside the member's "
                               f"editions {member.valid_from_edition}–{member.valid_to_edition}")
            elif edition not in expected:
                reasons.append(f"code-keyed line PB{edition} {member.code} {title!r} is in an edition "
                               "titles_seen omits")
            elif title != expected[edition]:
                reasons.append(f"code-keyed line PB{edition} {member.code} has title {title!r}; titles_seen "
                               f"has {expected[edition]!r} for PB{edition}")
        for reason in dict.fromkeys(reasons):
            self.member_failure(member, reason)

    def roster(self, member: Member, measure: MemberMeasure, map_rows: Sequence[dict]) -> None:
        """Item 3: the member's era-lines rows equal the map's lines under its
        era_decision_ids, whose program_key and line_item_code are its code."""
        mine = {(e.edition, e.era_key, e.budget_activity, e.filed_title, e.decision_id)
                for e in self.family.era_lines if e.member_id == member.member_id}
        measured = {(e.edition, e.era_key, e.budget_activity, e.filed_title, e.decision_id)
                    for e in measure.era_lines}
        for key in sorted(mine - measured, key=_none_as_empty):
            self.era_failure(key[0], key[1], f"roster: the era-lines row (BA {key[2]!r}, filed title "
                                             f"{key[3]!r}, decision {key[4]!r}) of {member.member_id!r} is not "
                                             "a p1_era_line_map line under its era_decision_ids")
        for key in sorted(measured - mine, key=_none_as_empty):
            self.era_failure(key[0], key[1], f"roster: the p1_era_line_map line (BA {key[2]!r}, filed title "
                                             f"{key[3]!r}, decision {key[4]!r}) under the era_decision_ids of "
                                             f"{member.member_id!r} is not in the era-lines file")
        for r in _bound_era_rows(member, map_rows):
            if (r["program_key"], r["line_item_code"]) != (member.code, member.code):
                self.era_failure(r["edition"], r["era_key"],
                                 f"roster: the p1_era_line_map line has program_key {r['program_key']!r} and "
                                 f"line_item_code {r['line_item_code']!r}; both must be the code "
                                 f"{member.code!r} of {member.member_id!r}")

    def unmapped(self, member: Member, map_rows: Sequence[dict]) -> None:
        """Item 4."""
        bound = {(e.edition, e.era_key) for e in self.family.era_lines}
        for r in map_rows:
            if (r["edition"], r["era_key"]) in bound or str(r["decision"] or "").startswith("exclude_"):
                continue
            self.member_failure(member, f"unmapped era line PB{r['edition']} {r['era_key']} prints code "
                                        f"{member.code} (decision {r['decision']!r}, {r['decision_id']!r}); "
                                        "it is neither in the era-lines file nor under an exclude_* decision")

    def guard(self, guard_rows: Sequence[dict]) -> None:
        """Item 5."""
        fam = self.family
        for row in guard_rows:
            covering = [m for m in fam.members if _covers(m, row)]
            if any(m.verdict in ("approved", "rejected") for m in covering):
                continue
            deferred = [m.member_row_id for m in covering if m.verdict == "deferred"]
            if deferred:
                self.report.undecided.append(_line(
                    REGISTRY_JSON.name, fam.platform_family,
                    f"undecided guard match {_guard_text(row)}: covered only by deferred row "
                    f"{', '.join(deferred)}"))
            else:
                self.json_failure(f"Unreviewed {fam.short_name} workbook identity: "
                                  f"{row['edition']}/{row['pe_bli']}")

    def sources(self, registry: Registry, site_dir: Path, members: Sequence[Member]) -> None:
        """Item 8 (a published family only)."""
        cited = set()
        for ref in [r for m in members for _, r in _refs(m)] + [r for _, r in _exclusion_refs(self.family)]:
            if ref.startswith("src:"):
                try:
                    cited.add(parse_ref(ref)[1]["source_id"])
                except ValueError:
                    continue
        for sid, source in registry.sources.items():
            if sid not in cited:
                continue
            if not source.archived:
                self.json_failure(f"source {sid!r}, cited by published family "
                                  f"{self.family.platform_family!r}, is not archived", row=sid)
                continue
            path = Path(site_dir) / (source.archive_path or "")
            if source.sha256 is None:
                self.json_failure(f"source {sid!r} has no sha256, so its archived copy {path} is unverifiable",
                                  row=sid)
            elif not path.is_file():
                self.json_failure(f"the archived copy {path} of source {sid!r} is missing (expected sha256 "
                                  f"{source.sha256})", row=sid)
            else:
                actual = hashlib.sha256(path.read_bytes()).hexdigest()
                if actual != source.sha256:
                    self.json_failure(f"the archived copy {path} of source {sid!r} has sha256 {actual}, "
                                      f"expected {source.sha256}", row=sid)


def check_lake(registry: Registry, *, lake: LakeReader, site_dir: Path,
               families: Sequence[str] | None = None) -> LakeReport:
    """Stage 2 (§9.2) for every family, or only those named, in JSON order.
    A draft family gets item 7 on its exclusions' doc: refs and reports its
    guard matches as undecided; a ratified or published one gets items 1–7,
    9 and 10, and item 8 too when published (R4: a recorded deferral is
    undecided, a moved drift_sha256 is stale)."""
    if families is None:
        chosen = registry.families
    else:
        named = {registry[pf].platform_family for pf in families}
        chosen = tuple(f for f in registry.families if f.platform_family in named)
    report = LakeReport([], [], [])
    lake_sets: dict[str, frozenset[str]] = {}
    for fam in chosen:
        check = _FamilyLake(fam, report)
        guard_rows = lake.guard_rows(fam.review_guard)
        if fam.status == "draft":
            check.doc_refs(site_dir, ())
            for row in guard_rows:
                report.undecided.append(_line(REGISTRY_JSON.name, fam.platform_family,
                                              f"undecided guard match {_guard_text(row)}"))
            continue
        if not lake_sets:
            lake_sets.update(collision=lake.collision_codes(), series=lake.series_codes())
        members = sorted(fam.members, key=lambda m: m.member_row_id)
        counted = [m for m in members if m.verdict == "approved" and m.role == "counted"]
        counted_ids = {m.member_row_id for m in counted}
        measured_rows = [m for m in members if m.drift_sha256 or m.member_row_id in counted_ids]
        lines = {m.member_row_id: _member_lines(lake, m) for m in measured_rows}
        measures = {m.member_row_id: _measure(lake, m, fam.headline, *lines[m.member_row_id])
                    for m in measured_rows}

        for m in counted:                                                    # item 1
            check.code_keyed(m, lines[m.member_row_id][0])
        for m in measured_rows:                                              # item 2
            measured = measures[m.member_row_id]
            if m.drift_sha256 and measured.drift_sha256 != m.drift_sha256:
                code_lines, map_rows = lines[m.member_row_id]
                n = len(_drift_keys(m, code_lines, _bound_era_rows(m, map_rows)))
                report.stale.append(_line(
                    MEMBERS_CSV.name, m.member_row_id,
                    f"drift_sha256 {m.drift_sha256} != {measured.drift_sha256} ({n} lines): the lake lines "
                    "this row binds changed since it was reviewed"))
        for m in counted:                                                    # items 3 and 4
            check.roster(m, measures[m.member_row_id], lines[m.member_row_id][1])
        for m in counted:
            check.unmapped(m, lines[m.member_row_id][1])
        check.guard(guard_rows)                                              # item 5
        for m in counted:                                                    # item 6
            if m.code in lake_sets["collision"]:
                check.member_failure(m, f"code {m.code!r} is a collision code (dim_programs reuses it across "
                                        "accounts or organizations); the series key cannot tell its members "
                                        "apart")
        check.doc_refs(site_dir, members)                                    # item 7
        if fam.status == "published":                                        # item 8
            check.sources(registry, site_dir, members)
        if fam.on_partial_corpus == "raise":                                 # item 9
            try:
                corpus_gate(fam, lake_sets["series"])
            except RegistryError as exc:
                report.failures.extend(exc.errors)
        for m in members:                                                    # item 10
            if m.verdict == "deferred":
                report.undecided.append(_line(MEMBERS_CSV.name, m.member_row_id,
                                              "deferred row, recorded and never exported"))
    return report


# ---------------------------------------------------------------------------
# The one-off F-15 seed (spec §7.6, §8.1; D13, R6)
# ---------------------------------------------------------------------------

F15_SEED_SPEC = ("docs/superpowers/specs/2026-10-20-families-piece2-registry-design.md"
                 "#81-r-dec-fam-f15-seed-class-ruling-pending-owner-approval-with-this-spec")
F15_SEED_RULING_TEXT = (
    "The F-15 family's registry entry records the eight budget lines the shipped F-15 history counts,"
    " as they ship at main 88538c6a. It records: their order, identities and display titles; the"
    " coverage rule (which lines are expected in which editions); the 31-line era roster; the F015E0"
    " note and its receipt; the family's scope and coverage notes. It changes no published figure,"
    " text or fact ID. Its evidence basis is the shipped family history (the F-15 family funding"
    " correction recorded in the ROADMAP on 2026-09-24), the piece-1 era decisions and the page's"
    " record-link narratives. Piece 3 must add an independent signal for each line before any F-15"
    " row changes.")
# §7.6's values the Python module does not hold.
F15_SEED_ROWS = {  # code: (valid_from, valid_to, era_decision_ids, record-link fact, purpose, variant_ids, variant_scope, program_tag)
    "0207134F": (2017, 2026, (), "0026c924483c6624", "develop", "C;D;E;EX", "shared", "F-15"),
    "0207146F": (2021, 2026, (), "d9982379a24ebac6", "develop", "EX", "variant", "F-15EX"),
    "0207171F": (2017, 2026, (), "b30fae1fd23999e4", "develop", "E;EX", "shared", "EPAWSS"),
    "F01500": (2017, 2026, ("F01500|3010F||2017-2023",), "c9e5c28dfebc879e", "upgrade", "C;D;E", "shared", "F-15"),
    "F015EX": (2021, 2026, ("F015EX|3010F||2021-2023",), "e7d5bcfb4a30f458", "buy", "EX", "variant", "F-15EX"),
    "F15EWS": (2019, 2026, ("F15EWS|3010F||2019-2023",), "b3dbc6ed30814fdf", "upgrade", "E", "variant", "EPAWSS"),
    "F0150P": (2017, 2020, ("F0150P|3010F||2017-2020",), None, "", "", "", "F-15"),
    "F015E0": (2020, 2022, ("F015E0|3010F||2020-2022",), None, "", "", "", "F-15EX"),
}
# RECORD_CONFIG[code].scopeNote, verbatim from site/src/lib/f15-family-data.ts:10-15
# (Task 7e's mirror test proves them against the TypeScript). The two
# history-only rows have no site record, so no scope note.
F15_SEED_SCOPE_NOTES: Mapping[str, str] = {
    "0207134F": ("Shared across C, D, E and EX: the software narrative explicitly names all four. The program "
                 "total is not allocated among variants."),
    "0207146F": ("F-15EX research, development, integration and test funding. Separate from aircraft "
                 "procurement."),
    "0207171F": ("Related EPAWSS system development. E and EX use the system; this record does not allocate "
                 "development funding between variants. No FY2025 or FY2026 TOA figure is published here."),
    "F01500": ("Shared C, D and E modifications and support equipment. The BLI total is not a "
               "variant-specific price or allocation."),
    "F015EX": ("This procurement BLI combines aircraft, modifications and support across budget activities. "
               "It is broader than the price of new aircraft alone."),
    "F15EWS": ("The procurement narrative describes EPAWSS installation and related support for F-15E "
               "aircraft. This amount is not assigned to F-15EX."),
}
_F15 = "f-15"
_F15_HISTORY_FILE = "f15_funding_history.json"
_ABSENT = object()


@dataclass(frozen=True)
class F15Constants:
    """The shipped F-15 constants the seed reads: f15_funding_history.py's
    (88538c6a) plus the hardcoded expected set as expected_from (code -> the
    first edition its absence is reported, None for never)."""

    programs: Sequence[Mapping]
    modern_members: Mapping[str, tuple[str, str, str]]
    era_members: Mapping[int, Sequence[tuple[int, str, str]]]
    era_program_codes: Mapping[int, Mapping[int, str]]
    program_notes: Mapping[str, Mapping[str, str]]
    note_receipts: Mapping[str, Mapping[str, object]]
    scope_note: str
    coverage_notes: Sequence[str]
    expected_from: Mapping[str, int | None]


def _csv_data_rows(path: Path) -> int:
    text = path.read_bytes().decode("utf-8", errors="replace")
    return sum(1 for fields in list(csv.reader(io.StringIO(text, newline="")))[1:] if fields)


def _f15_seed_registry_refusals(paths: RegistryPaths) -> list[str]:
    """The seed is one-off: it writes only where the registry holds nothing
    (it replaces the three files whole)."""
    errors: list[str] = []
    name = paths.json.name
    if paths.json.exists():
        text, faults = _read_text(paths.json)
        doc, json_faults = _parse_json(text) if text is not None else (None, [])
        for fault in faults + json_faults:
            errors.append(_line(name, None, f"{fault}; the seed writes only into an empty registry"))
        top = doc if isinstance(doc, dict) else {}
        held = {key: [_entry_id(entry, id_key, f"{key}[{i}]") for i, entry in enumerate(top.get(key))]
                for key, id_key in (("families", "platform_family"), ("rulings", "ruling_id"),
                                    ("sources", "source_id"))
                if isinstance(top.get(key), list)}
        if held.get("families"):
            errors.append(_line(name, None, f"registry already holds {', '.join(held['families'])}; the seed "
                                            "writes only into an empty registry"))
        if _F15 in held.get("families", ()):
            errors.append(_line(name, _F15, f"{_F15} already in {name} (the seed is one-off)"))
        for key in ("rulings", "sources"):
            if held.get(key):
                errors.append(_line(name, None, f"registry already holds {key} {', '.join(held[key])}; the "
                                                "seed writes only into an empty registry"))
    for path in (paths.members, paths.era_lines):
        if path.exists() and (n := _csv_data_rows(path)):
            errors.append(_line(path.name, None, f"already holds {n} data row(s); the seed writes only into "
                                                 "an empty registry"))
    return errors


def _show(entry: Mapping, key: str) -> str:
    value = entry.get(key, _ABSENT)
    return "nothing" if value is _ABSENT else repr(value)


def _f15_shipped_refusals(constants: F15Constants, site_dir: Path) -> list[str]:
    """The constants against the shipped output (§7.6): fam2's
    json/f15_funding_history.json, whose sha256 must first equal the pin.
    The values are compared only against the pinned bytes."""
    from govbudget import verify_era_map

    errors: list[str] = []

    def add(reason: str) -> None:
        errors.append(_line(_F15_HISTORY_FILE, None, reason))

    path, pin_path = Path(site_dir) / "json" / _F15_HISTORY_FILE, verify_era_map.F15_HISTORY_PIN
    if not path.is_file():
        add(f"the shipped history is missing (no file at {path}); the seed compares the constants with it")
        return errors
    tokens = pin_path.read_text(encoding="utf-8").split() if pin_path.is_file() else []
    if not tokens or not _SHA256_RE.fullmatch(tokens[0]):
        add(f"the pin {pin_path} is missing or does not start with a sha256 hex digest")
        return errors
    data = path.read_bytes()
    actual = hashlib.sha256(data).hexdigest()
    if actual != tokens[0]:
        add(f"sha256 {actual} of {path} is not the pin {tokens[0]} ({pin_path}); the seed compares the "
            "constants only with the pinned shipped history")
        return errors
    payload = json.loads(data)
    mine, shipped = list(constants.programs), payload["programs"]
    if len(mine) != len(shipped):
        add(f"programs: the constants have {len(mine)} entries, the shipped history has {len(shipped)}")
    for i, (a, b) in enumerate(zip(mine, shipped)):
        for key in dict.fromkeys([*a, *b]):
            if a.get(key, _ABSENT) != b.get(key, _ABSENT):
                add(f"programs[{i}].{key}: the constants have {_show(a, key)}, the shipped history has "
                    f"{_show(b, key)}")
    if constants.scope_note != payload["scope_note"]:
        add(f"scope_note: the constants have {constants.scope_note!r}, the shipped history has "
            f"{payload['scope_note']!r}")
    notes, shipped_notes = list(constants.coverage_notes), payload["coverage_notes"]
    if len(notes) != len(shipped_notes):
        add(f"coverage_notes: the constants have {len(notes)} notes, the shipped history has "
            f"{len(shipped_notes)}")
    for i, (a, b) in enumerate(zip(notes, shipped_notes)):
        if a != b:
            add(f"coverage_notes[{i}]: the constants have {a!r}, the shipped history has {b!r}")
    return errors


def _f15_seed_doc(constants: F15Constants, *, decided_on: str, owner_reply: str) -> dict:
    """§7.2's F-15 column, the ruling (§8.1) and no sources."""
    notes = {}
    for program in constants.programs:
        note = constants.program_notes.get(program["code"])
        if note is not None:
            fid = note["note_fact_id"]
            notes[program["id"]] = {"note": note["note"], "note_fact_id": fid,
                                    "expected_receipt": dict(constants.note_receipts.get(fid, {}))}
    family = {
        "platform_family": _F15, "status": "published", "published_ruling_id": F15_SEED_RULING_ID,
        "short_name": "F-15", "history_dataset": "f15_funding_history", "on_partial_corpus": "skip",
        "headline": {"measure": "cumulative_actuals", "first_fy": 2015, "last_fy": 2024},
        "scope_note": constants.scope_note, "coverage_notes": list(constants.coverage_notes),
        "review_guard": {"exhibits": ["P-1", "R-1"], "title_substrings": ["F-15", "F15"]},
        "program_notes": notes,
        "variants": [{"id": v} for v in ("A", "B", "C", "D", "E", "EX")],
        "program_tags": [{"tag": t, "mdap_code": None} for t in ("F-15", "F-15EX", "EPAWSS")],
        "purpose_tags": [], "caveat_ids": ["classified", "construction", "nominal_dollars", "cr_years"],
        "exclusions": [], "aliases": [], "contract_routes": None, "official_estimates": None, "art": None,
    }
    ruling = {"ruling_id": F15_SEED_RULING_ID, "kind": "class", "platform_family": _F15,
              "decided_by": "owner", "decided_on": decided_on, "owner_reply": owner_reply,
              "spec": F15_SEED_SPEC, "text": F15_SEED_RULING_TEXT, "seed_basis": True}
    return {"schema_version": 1, "rulings": [ruling], "sources": [], "families": [family]}


def _f15_seed_member_rows(constants: F15Constants, *, decided_on: str) -> list[dict[str, str]]:
    """One counted row per program, in order; the lake columns are blank
    until measured."""
    rows = []
    for position, program in enumerate(constants.programs, start=1):
        code = program["code"]
        frm, to, era_ids, fact, purpose, variants, scope, tag = F15_SEED_ROWS[code]
        expected_from = constants.expected_from[code]
        row = dict.fromkeys(MEMBER_COLUMNS, "")
        row.update(
            member_row_id=f"{_F15}|{program['id']}|{frm}-{to}", platform_family=_F15,
            member_id=program["id"], exhibit=program["exhibit"], account=program["account"],
            organization=program["organization"], code=code, role="counted", verdict="approved",
            appropriation="procurement" if program["exhibit"] == "P-1" else "rdte",
            display_title=program["title"], page_slug=program["program_slug"] or "",
            valid_from_edition=str(frm), valid_to_edition=str(to),
            coverage_expected_from_edition="" if expected_from is None else str(expected_from),
            era_decision_ids=";".join(era_ids), member_order=str(position), program_tag=tag,
            purpose=purpose, variant_ids=variants, variant_scope=scope,
            scope_note=F15_SEED_SCOPE_NOTES.get(code, ""), services_using="USAF", paying_org="AF",
            evidence=";".join([f"era:{d}" for d in era_ids] + ([f"fact:{fact}"] if fact else [])),
            decided_by="owner", ruling_id=F15_SEED_RULING_ID, decided_on=decided_on)
        rows.append(row)
    return rows


def _f15_seed_era_rows(constants: F15Constants) -> list[dict[str, str]]:
    """ERA_MEMBERS x ERA_PROGRAM_CODES, each line under its member's one era
    decision (§7.4)."""
    rows = []
    for edition, lines in constants.era_members.items():
        for line, budget_activity, title in lines:
            # A line without a code or decision gets a blank one, which the
            # roster check refuses (never a KeyError).
            code = constants.era_program_codes.get(edition, {}).get(line, "")
            decisions = F15_SEED_ROWS[code][2] if code in F15_SEED_ROWS else ()
            decision_id = decisions[0] if len(decisions) == 1 else ""
            rows.append({"platform_family": _F15, "member_id": f"P-1:3010F:AF:{code}", "edition": str(edition),
                         "era_key": f"3010F-AF-L{line}", "budget_activity": budget_activity,
                         "filed_title": title, "decision_id": decision_id})
    return rows


def _f15_roster_refusals(era_rows: Sequence[Mapping[str, str]], measured: set[tuple],
                         era_name: str) -> list[str]:
    """§7.6: the shipped roster against the p1_era_line_map lines under the
    members' era decisions, as (member_id, edition, era_key, budget_activity,
    filed_title, decision_id)."""
    roster = {(r["member_id"], int(r["edition"]), r["era_key"], r["budget_activity"], r["filed_title"],
               r["decision_id"]) for r in era_rows}
    errors = []
    for key in sorted(roster - measured, key=_none_as_empty):
        member_id, edition, era_key, ba, title, decision_id = key
        errors.append(_line(era_name, f"{_F15}|{edition}|{era_key}",
                            f"roster: the shipped roster line ({member_id}, BA {ba!r}, filed title {title!r}, "
                            f"decision {decision_id!r}) is not a p1_era_line_map line under that decision"))
    for key in sorted(measured - roster, key=_none_as_empty):
        member_id, edition, era_key, ba, title, decision_id = key
        errors.append(_line(era_name, f"{_F15}|{edition}|{era_key}",
                            f"roster: the p1_era_line_map line ({member_id}, BA {ba!r}, filed title {title!r}, "
                            f"decision {decision_id!r}) is not in the shipped roster (ERA_MEMBERS x "
                            "ERA_PROGRAM_CODES)"))
    return errors


def seed_shipped_f15(*, constants: F15Constants, duckdb_path: Path, site_dir: Path, decided_on: str,
                     owner_reply: str, paths: RegistryPaths = DEFAULT_PATHS) -> dict:
    """`propose --seed-shipped f-15` (§7.6): write the F-15 as registry entry
    #1, `published` under R-DEC-FAM-F15-SEED, from the shipped constants and
    §7.6's table, measured against the lake. The input refusals are listed
    together before the lake opens (R6); then the roster refusals; then stage
    1 and stage 2 on the result. A refusal writes nothing; the write replaces
    the three files atomically. `owner_reply` loses exactly one trailing
    newline (D13), so a caller passes the reply file's text as read."""
    owner_reply = owner_reply.removesuffix("\n")
    errors = _f15_seed_registry_refusals(paths)
    if not owner_reply.strip():
        errors.append(_line(paths.json.name, F15_SEED_RULING_ID,
                            f"owner_reply {owner_reply!r} is blank (the owner's verbatim approval is required)"))
    if not _iso_date(decided_on):
        errors.append(_line(paths.json.name, F15_SEED_RULING_ID,
                            f"decided_on {decided_on!r} is not an ISO date (YYYY-MM-DD)"))
    errors += _f15_shipped_refusals(constants, site_dir)
    codes = sorted(p["code"] for p in constants.programs)
    if sorted(constants.expected_from) != codes:
        errors.append(_line(paths.members.name, None,
                            f"expected_from has codes {sorted(constants.expected_from)}; the programs' codes are "
                            f"{codes} (coverage_expected_from_edition needs exactly one entry per program)"))
    if errors:
        raise RegistryError(errors)

    doc = _f15_seed_doc(constants, decided_on=decided_on, owner_reply=owner_reply)
    rows = _f15_seed_member_rows(constants, decided_on=decided_on)
    era_rows = _f15_seed_era_rows(constants)
    headline = Headline(**doc["families"][0]["headline"])
    lake = open_lake(duckdb_path)
    try:
        measured: set[tuple] = set()
        for row in rows:
            m = measure_member(lake, Member.from_row(row), headline=headline)
            row.update(titles_seen="; ".join(f"{e}: {t}" for e, t in m.titles_seen),
                       budget_activities=";".join(m.budget_activities), drift_sha256=m.drift_sha256)
            measured |= {(e.member_id, e.edition, e.era_key, e.budget_activity, e.filed_title, e.decision_id)
                         for e in m.era_lines}
        roster = _f15_roster_refusals(era_rows, measured, paths.era_lines.name)
        if roster:
            raise RegistryError(roster)
        decisions = {d["decision_id"]: d for d in era_map.read_seed(era_map.default_seed_path())}
        registry = parse_registry(doc, rows, era_rows, era_decisions=decisions, paths=paths)
        check_lake(registry, lake=lake, site_dir=site_dir, families=[_F15]).raise_if_failed()
    finally:
        lake.close()
    write_registry(paths, doc=doc, members=rows, era_lines=era_rows)
    return {"members": len(rows), "era_lines": len(era_rows), "rulings": 1}
